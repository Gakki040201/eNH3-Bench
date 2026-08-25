"""Deterministic, provenance-first ingestion for LiNRR Training Dataset v0.

This module deliberately separates candidate discovery from scientific
admission.  Bundle boundaries are hard ownership boundaries; no function in
this module links an asset or evidence locator to another bundle.
"""

from __future__ import annotations

import csv
import hashlib
import json
import math
import re
import zipfile
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Iterator


SCHEMA_VERSION = "linrr_experiment_record_v0"
SOURCE_GROUPS = {
    "Shaofeng_Li": "shaofeng_li_curated",
    "LiNRR_Electrolytes": "linrr_electrolyte_curated",
    "Battery_Additives": "battery_transfer_curated",
}
SUPPORTED_SUFFIXES = {".pdf", ".docx", ".xlsx", ".csv", ".zip", ".txt"}

UNIFIED_FIELDS = [
    "schema_version", "record_id", "paper_id", "provisional_bundle_id", "source_group",
    "reaction_family", "document_genre", "training_role", "source_asset_id", "source_filename",
    "source_file_sha256", "source_locator", "source_table", "source_figure", "source_page",
    "source_text_excerpt", "evidence_authority", "extraction_status", "review_status",
    "evidence_owner_id",
    "lithium_salt", "lithium_salt_concentration_mol_L", "lithium_salt_concentration_raw_value",
    "lithium_salt_concentration_raw_unit", "solvent", "cosolvent", "solvent_ratio", "proton_donor",
    "proton_donor_concentration_value", "proton_donor_concentration_unit",
    "proton_donor_concentration_raw_value", "proton_donor_concentration_raw_unit", "additive",
    "additive_concentration_value", "additive_concentration_unit", "additive_concentration_raw_value",
    "additive_concentration_raw_unit", "water_content_value", "water_content_unit",
    "water_content_raw_value", "water_content_raw_unit", "oxygen_condition", "electrolyte_architecture",
    "electrolyte_notes", "current_density_mA_cm2", "current_density_raw_value", "current_density_raw_unit",
    "current_A", "electrode_area_cm2", "total_charge_C", "duration_h", "duration_raw_value",
    "duration_raw_unit", "temperature_C", "pressure_bar", "pressure_raw_value", "pressure_raw_unit",
    "gas_flow_sccm", "liquid_flow_mL_min", "electrode_material", "substrate_material", "cell_type",
    "reactor_type", "fe_nh3_percent", "nh3_rate_value", "nh3_rate_unit", "nh3_rate_raw_value",
    "nh3_rate_raw_unit", "cell_voltage_V", "working_electrode_potential_V", "stability_duration_h",
    "replicate_n", "error_type", "error_value", "blank_control", "nitrogen_source_validation",
    "isotope_validation", "quantification_method", "failure_status", "failure_type",
    "missing_required_fields", "reported_not_measured_fields", "ambiguity_flags", "eligibility_tier",
    "model_eligible_fe", "interface_sei_characterization_present", "interface_lif_reported",
    "interface_n_containing_sei_reported", "source_bundle_id",
]


def sha256_file(path: str | Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def stable_id(prefix: str, *parts: object, length: int = 16) -> str:
    payload = "\x1f".join(str(part) for part in parts)
    return f"{prefix}_{hashlib.sha256(payload.encode('utf-8')).hexdigest()[:length]}"


def normalize_doi(value: object) -> str | None:
    text = str(value or "").strip().lower()
    text = re.sub(r"^(?:https?://(?:dx\.)?doi\.org/|doi\s*:\s*)", "", text)
    match = re.search(r"10\.\d{4,9}/[-._;()/:a-z0-9]+", text, re.I)
    if not match:
        return None
    return match.group(0).rstrip(".,;:)]}").lower()


def normalize_title(value: object) -> str:
    text = str(value or "").casefold()
    text = re.sub(r"[^\w]+", " ", text, flags=re.UNICODE)
    return " ".join(text.split())


def normalize_measurement(value: object, unit: object, kind: str) -> dict[str, Any]:
    """Normalize only explicitly supported, dimensionally unambiguous units."""
    try:
        number = float(str(value).replace(",", "").strip())
    except (TypeError, ValueError):
        return {"raw_value": value, "raw_unit": unit, "normalized_value": None, "normalized_unit": None}
    raw_unit = str(unit or "").strip()
    key = re.sub(r"\s+", "", raw_unit).replace("−", "-").lower()
    converted: tuple[float, str] | None = None
    if kind == "current_density" and key in {"a/cm2", "acm-2", "a·cm-2"}:
        converted = (number * 1000.0, "mA/cm2")
    elif kind == "current_density" and key in {"ma/cm2", "macm-2", "ma·cm-2"}:
        converted = (number, "mA/cm2")
    elif kind == "duration" and key in {"min", "mins", "minute", "minutes"}:
        converted = (number / 60.0, "h")
    elif kind == "duration" and key in {"h", "hr", "hrs", "hour", "hours"}:
        converted = (number, "h")
    elif kind == "pressure" and key == "kpa":
        converted = (number / 100.0, "bar")
    elif kind == "pressure" and key == "atm":
        converted = (number * 1.01325, "bar")
    elif kind == "pressure" and key == "bar":
        converted = (number, "bar")
    elif kind == "concentration" and key in {"m", "mol/l", "moll-1", "mol·l-1"}:
        converted = (number, "mol/L")
    return {
        "raw_value": number,
        "raw_unit": raw_unit or None,
        "normalized_value": converted[0] if converted else None,
        "normalized_unit": converted[1] if converted else None,
    }


def _pdf_blocks(path: Path) -> list[dict[str, Any]]:
    try:
        import fitz  # type: ignore
    except ImportError as exc:  # pragma: no cover - optional dependency path
        raise RuntimeError("PDF ingestion requires the document extra (PyMuPDF)") from exc
    blocks: list[dict[str, Any]] = []
    with fitz.open(path) as document:
        for index, page in enumerate(document, start=1):
            blocks.append({"locator": f"page:{index}", "page": index, "table": None, "text": page.get_text("text")})
    return blocks


def _docx_blocks(path: Path) -> list[dict[str, Any]]:
    try:
        from docx import Document  # type: ignore
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("DOCX ingestion requires python-docx") from exc
    doc = Document(path)
    blocks = [
        {"locator": f"paragraph:{i}", "page": None, "table": None, "text": p.text}
        for i, p in enumerate(doc.paragraphs, 1) if p.text.strip()
    ]
    for table_i, table in enumerate(doc.tables, 1):
        for row_i, row in enumerate(table.rows, 1):
            text = " | ".join(cell.text.strip() for cell in row.cells)
            if text.strip(" |"):
                blocks.append({"locator": f"table:{table_i}:row:{row_i}", "page": None, "table": f"Table {table_i}", "text": text})
    return blocks


def _xlsx_blocks(path: Path) -> list[dict[str, Any]]:
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    blocks: list[dict[str, Any]] = []
    try:
        for sheet in workbook.worksheets:
            for row_i, row in enumerate(sheet.iter_rows(values_only=True), 1):
                values = [str(value) for value in row if value is not None]
                if values:
                    blocks.append({"locator": f"sheet:{sheet.title}:row:{row_i}", "page": None, "table": sheet.title, "text": " | ".join(values)})
    finally:
        workbook.close()
    return blocks


def _csv_blocks(path: Path) -> list[dict[str, Any]]:
    blocks: list[dict[str, Any]] = []
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
        for row_i, row in enumerate(csv.reader(handle), 1):
            if any(str(cell).strip() for cell in row):
                blocks.append({"locator": f"row:{row_i}", "page": None, "table": path.stem, "text": " | ".join(row)})
    return blocks


def read_asset_blocks(path: str | Path) -> list[dict[str, Any]]:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _pdf_blocks(path)
    if suffix == ".docx":
        return _docx_blocks(path)
    if suffix == ".xlsx":
        return _xlsx_blocks(path)
    if suffix == ".csv":
        return _csv_blocks(path)
    if suffix == ".txt":
        return [{"locator": "text:1", "page": None, "table": None, "text": path.read_text(encoding="utf-8", errors="replace")}]
    return []


def _archive_members(path: Path) -> list[dict[str, Any]]:
    if path.suffix.lower() != ".zip":
        return []
    try:
        with zipfile.ZipFile(path) as archive:
            return [
                {"name": item.filename, "size_bytes": item.file_size, "compressed_size_bytes": item.compress_size, "crc32": f"{item.CRC:08x}"}
                for item in sorted(archive.infolist(), key=lambda value: value.filename)
            ]
    except zipfile.BadZipFile:
        return [{"error": "BAD_ZIP"}]


def _role_for_asset(path: Path, preview: str, sibling_count: int) -> str:
    suffix = path.suffix.lower()
    lower_name = path.name.casefold()
    lower_text = preview[:12000].casefold()
    if suffix == ".zip":
        return "archive"
    if suffix == ".txt":
        return "metadata_text"
    if suffix in {".csv", ".xlsx"}:
        return "source_data" if any(token in lower_name for token in ("source", "train", "test", "candidate", "data")) else "supplementary_data"
    si_signals = sum(token in lower_text for token in ("supporting information", "supplementary information", "experimental procedures"))
    si_signals += sum(token in lower_name for token in ("_si", "suppl", "supp", "mmc", "moesm", "sm."))
    article_signals = sum(token in lower_text for token in ("abstract", "introduction", "article recommendations", "received:"))
    if si_signals > article_signals:
        return "supporting_information"
    if article_signals or sibling_count == 1:
        return "main_article"
    return "other"


def scan_bundle(bundle_path: str | Path, source_group: str) -> dict[str, Any]:
    bundle_path = Path(bundle_path).resolve()
    bundle_id = stable_id("BUNDLE", source_group, bundle_path.name, length=12)
    provisional = "NEWBUNDLE_" + hashlib.sha256(f"{source_group}\x1f{bundle_path.name}".encode()).hexdigest()[:12]
    paths = sorted((p for p in bundle_path.rglob("*") if p.is_file()), key=lambda p: p.relative_to(bundle_path).as_posix().casefold())
    assets: list[dict[str, Any]] = []
    for path in paths:
        digest = sha256_file(path)
        preview = ""
        parse_error = None
        if path.suffix.lower() in SUPPORTED_SUFFIXES - {".zip"}:
            try:
                preview = "\n".join(block["text"] for block in read_asset_blocks(path)[:3])
            except Exception as exc:  # retain manifest even for a malformed asset
                parse_error = f"{type(exc).__name__}: {exc}"
        role = _role_for_asset(path, preview, len(paths))
        assets.append({
            "source_asset_id": stable_id("ASSET", bundle_id, path.relative_to(bundle_path).as_posix(), digest),
            "relative_path": path.relative_to(bundle_path).as_posix(),
            "absolute_path": str(path),
            "size_bytes": path.stat().st_size,
            "sha256": digest,
            "format": path.suffix.lower().lstrip(".") or "unknown",
            "probable_role": role,
            "role_basis": "bundle-local content and deterministic file-structure signals",
            "parse_error": parse_error,
            "archive_members": _archive_members(path),
        })
    return {
        "schema_version": "training_bundle_manifest_v0",
        "bundle_id": bundle_id,
        "provisional_bundle_id": provisional,
        "source_group": source_group,
        "bundle_path": str(bundle_path),
        "assets": assets,
    }


def scan_bundle_roots(roots: dict[str, str | Path]) -> list[dict[str, Any]]:
    manifests: list[dict[str, Any]] = []
    for source_group, root_value in sorted(roots.items()):
        root = Path(root_value)
        if not root.exists():
            continue
        for child in sorted((p for p in root.iterdir() if p.is_dir()), key=lambda p: p.name.casefold()):
            manifests.append(scan_bundle(child, source_group))
    return manifests


def _asset_text(asset: dict[str, Any], limit_blocks: int | None = None) -> str:
    path = Path(asset["absolute_path"])
    if path.suffix.lower() == ".zip":
        return ""
    try:
        blocks = read_asset_blocks(path)
    except Exception:
        return ""
    if limit_blocks is not None:
        blocks = blocks[:limit_blocks]
    return "\n".join(str(block.get("text") or "") for block in blocks)


def derive_paper_identity(manifest: dict[str, Any]) -> dict[str, Any]:
    assets = sorted(manifest["assets"], key=lambda a: (a["probable_role"] != "main_article", a["relative_path"].casefold()))
    texts = [(asset, _asset_text(asset, 4)) for asset in assets if asset["format"] in {"pdf", "docx", "txt"}]
    combined = "\n".join(text for _, text in texts)
    doi_candidates = [normalize_doi(value) for value in re.findall(r"(?:https?://(?:dx\.)?doi\.org/|doi\s*:?\s*)?(10\.\d{4,9}/[-._;()/:A-Z0-9]+)", combined, re.I)]
    doi = next((value for value in doi_candidates if value), None)
    title = None
    for asset, text in texts:
        if asset["probable_role"] != "main_article":
            continue
        lines = [line.strip() for line in text.splitlines() if 20 <= len(line.strip()) <= 300]
        excluded = ("download", "journal", "abstract", "supporting information", "www.", "http")
        title = next((line for line in lines[:40] if not any(token in line.casefold() for token in excluded)), None)
        if title:
            break
    year_matches = re.findall(r"\b(?:19|20)\d{2}\b", combined[:15000])
    year = int(year_matches[0]) if year_matches else None
    journal = None
    journal_patterns = [r"Cite This:\s*([^\n,]+)", r"([A-Z][A-Za-z &]+)\s+(?:19|20)\d{2},"]
    for pattern in journal_patterns:
        match = re.search(pattern, combined[:10000])
        if match:
            journal = " ".join(match.group(1).split())
            break
    authors = None
    if title:
        pos = combined.find(title)
        following = [line.strip() for line in combined[pos + len(title):pos + len(title) + 1500].splitlines() if line.strip()]
        authors = next((line for line in following if "," in line and len(line) < 500), None)
    return {"title": title, "doi": doi, "year": year, "journal": journal, "authors": authors, "first_author": _first_author(authors)}


def _first_author(authors: str | None) -> str | None:
    if not authors:
        return None
    return re.split(r",|;|\band\b", authors, maxsplit=1, flags=re.I)[0].strip() or None


def load_paper_registry(path: str | Path) -> list[dict[str, str]]:
    with Path(path).open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def match_paper_identity(identity: dict[str, Any], registry: list[dict[str, str]]) -> dict[str, Any]:
    doi = normalize_doi(identity.get("doi"))
    if doi:
        exact = [row for row in registry if normalize_doi(row.get("doi")) == doi]
        if exact:
            return {"existing_paper_id": exact[0]["paper_id"], "probable_existing_paper_id": None, "match_type": "DOI_EXACT", "match_confidence": 1.0, "review_status": "MATCHED"}
    title = normalize_title(identity.get("title"))
    if title:
        candidates: list[tuple[float, dict[str, str]]] = []
        title_words = set(title.split())
        for row in registry:
            other = normalize_title(row.get("title"))
            if not other:
                continue
            other_words = set(other.split())
            score = len(title_words & other_words) / max(len(title_words | other_words), 1)
            if score >= 0.82 and not (doi and normalize_doi(row.get("doi")) and doi != normalize_doi(row.get("doi"))):
                candidates.append((score, row))
        if candidates:
            score, row = max(candidates, key=lambda value: value[0])
            return {"existing_paper_id": None, "probable_existing_paper_id": row["paper_id"], "match_type": "TITLE_PROBABLE", "match_confidence": round(score, 4), "review_status": "REVIEW_REQUIRED"}
    return {"existing_paper_id": None, "probable_existing_paper_id": None, "match_type": "NEW_PROVISIONAL", "match_confidence": 0.0, "review_status": "REVIEW_REQUIRED"}


def build_identity_map(manifests: list[dict[str, Any]], registry: list[dict[str, str]]) -> list[dict[str, Any]]:
    rows = []
    for manifest in manifests:
        identity = derive_paper_identity(manifest)
        match = match_paper_identity(identity, registry)
        manifest["identity"] = identity
        manifest["identity_match"] = match
        rows.append({
            "bundle_id": manifest["bundle_id"], "source_group": manifest["source_group"],
            "title": identity.get("title"), "doi": identity.get("doi"),
            "existing_paper_id": match.get("existing_paper_id"), "probable_existing_paper_id": match.get("probable_existing_paper_id"),
            "match_type": match["match_type"], "match_confidence": match["match_confidence"], "review_status": match["review_status"],
        })
    return rows


def _empty_record() -> dict[str, Any]:
    record = {field: None for field in UNIFIED_FIELDS}
    record.update({
        "schema_version": SCHEMA_VERSION, "missing_required_fields": [], "reported_not_measured_fields": [],
        "ambiguity_flags": [], "model_eligible_fe": False,
        "interface_sei_characterization_present": False, "interface_lif_reported": False,
        "interface_n_containing_sei_reported": False,
    })
    return record


SALTS = ["LiBF4", "LiClO4", "LiTFSI", "LiFSI", "LiOTf", "LiCF3SO3", "LiNO3", "LiPF6", "LiBOB", "LiDFOB"]
SOLVENTS = [("tetrahydrofuran", "THF"), ("THF", "THF"), ("1,2-dimethoxyethane", "DME"), ("DME", "DME"), ("diglyme", "diglyme"), ("tetraglyme", "tetraglyme"), ("diethyl ether", "DEE"), ("DMSO", "DMSO")]
DONORS = ["ethanol", "EtOH", "phenol", "methanol", "isopropanol", "water", "H2O", "tert-butanol", "1-butanol"]


def _first_term(text: str, terms: Iterable[str]) -> str | None:
    for term in terms:
        if re.search(rf"(?<![A-Za-z0-9]){re.escape(term)}(?![A-Za-z0-9])", text, re.I):
            return term
    return None


def _extract_number(text: str, pattern: str) -> tuple[float | None, str | None]:
    match = re.search(pattern, text, re.I)
    if not match:
        return None, None
    try:
        return float(match.group(1)), match.group(2)
    except (ValueError, IndexError):
        return None, None


def _extract_conditions(text: str, record: dict[str, Any]) -> None:
    record["lithium_salt"] = _first_term(text, SALTS)
    solvent_hits = [canonical for term, canonical in SOLVENTS if re.search(rf"(?<!\w){re.escape(term)}(?!\w)", text, re.I)]
    record["solvent"] = solvent_hits[0] if solvent_hits else None
    record["cosolvent"] = solvent_hits[1] if len(solvent_hits) > 1 else None
    record["proton_donor"] = _first_term(text, DONORS)
    salt = record["lithium_salt"]
    if salt:
        match = re.search(rf"(\d+(?:\.\d+)?)\s*(M|mol\s*/\s*L|mM)\s+{re.escape(salt)}|{re.escape(salt)}[^.;\n]{{0,30}}?(\d+(?:\.\d+)?)\s*(M|mol\s*/\s*L|mM)", text, re.I)
        if match:
            value, unit = (match.group(1), match.group(2)) if match.group(1) else (match.group(3), match.group(4))
            norm = normalize_measurement(value, unit, "concentration")
            record["lithium_salt_concentration_raw_value"] = norm["raw_value"]
            record["lithium_salt_concentration_raw_unit"] = norm["raw_unit"]
            record["lithium_salt_concentration_mol_L"] = norm["normalized_value"]
    current, current_unit = _extract_number(text, r"(\d+(?:\.\d+)?)\s*(mA\s*(?:cm[−-]?2|/\s*cm2)|A\s*(?:cm[−-]?2|/\s*cm2))")
    if current is not None:
        norm = normalize_measurement(current, current_unit, "current_density")
        record["current_density_raw_value"], record["current_density_raw_unit"] = norm["raw_value"], norm["raw_unit"]
        record["current_density_mA_cm2"] = norm["normalized_value"]
    duration, duration_unit = _extract_number(text, r"(\d+(?:\.\d+)?)\s*(h(?:ours?)?|min(?:utes?)?)\b")
    if duration is not None:
        norm = normalize_measurement(duration, duration_unit, "duration")
        record["duration_raw_value"], record["duration_raw_unit"] = norm["raw_value"], norm["raw_unit"]
        record["duration_h"] = norm["normalized_value"]
    pressure, pressure_unit = _extract_number(text, r"(\d+(?:\.\d+)?)\s*(bar|kPa|atm)\b")
    if pressure is not None:
        norm = normalize_measurement(pressure, pressure_unit, "pressure")
        record["pressure_raw_value"], record["pressure_raw_unit"] = norm["raw_value"], norm["raw_unit"]
        record["pressure_bar"] = norm["normalized_value"]
    temperature, _ = _extract_number(text, r"(−?-?\d+(?:\.\d+)?)\s*(°?C)\b")
    record["temperature_C"] = temperature
    gas_flow, _ = _extract_number(text, r"(\d+(?:\.\d+)?)\s*(sccm)\b")
    record["gas_flow_sccm"] = gas_flow
    water, water_unit = _extract_number(text, r"(?:water|H2O)[^.;\n]{0,35}?(\d+(?:\.\d+)?)\s*(ppm|vol%|wt%|mM|M)\b")
    if water is not None:
        record["water_content_raw_value"], record["water_content_raw_unit"] = water, water_unit
        record["water_content_value"], record["water_content_unit"] = water, water_unit
    donor = record["proton_donor"]
    if donor:
        match = re.search(rf"(\d+(?:\.\d+)?)\s*(M|mM|vol%|wt%)\s+{re.escape(donor)}|{re.escape(donor)}[^.;\n]{{0,25}}?(\d+(?:\.\d+)?)\s*(M|mM|vol%|wt%)", text, re.I)
        if match:
            value, unit = (match.group(1), match.group(2)) if match.group(1) else (match.group(3), match.group(4))
            record["proton_donor_concentration_raw_value"] = float(value)
            record["proton_donor_concentration_raw_unit"] = unit
            record["proton_donor_concentration_value"] = float(value)
            record["proton_donor_concentration_unit"] = "mol/L" if unit.casefold() == "m" else unit
    lower = text.casefold()
    record["oxygen_condition"] = "oxygen_present" if re.search(r"\b(?:oxygen|o2)\b", lower) else None
    record["electrolyte_architecture"] = "LHCE" if "localized high-concentration" in lower or "lhce" in lower else ("high_concentration" if "high-concentration" in lower else None)
    record["interface_sei_characterization_present"] = bool(re.search(r"\b(?:xps|cryo-?tem|cryo-?em|sei characterization|solid electrolyte interphase)\b", lower))
    record["interface_lif_reported"] = bool(re.search(r"\blif\b|lithium fluoride", lower))
    record["interface_n_containing_sei_reported"] = bool(re.search(r"n-containing sei|nitrogen-containing sei|\bli3n\b", lower))
    record["blank_control"] = "reported" if re.search(r"blank|control experiment", lower) else None
    record["isotope_validation"] = "reported" if re.search(r"15n|isotop", lower) else None
    record["nitrogen_source_validation"] = "reported" if re.search(r"gas purification|n2 purification|nitrogen source", lower) else None
    for method in ("NMR", "indophenol", "ion chromatography", "spectrophotometry"):
        if method.casefold() in lower:
            record["quantification_method"] = method
            break


def classify_document(identity: dict[str, Any], full_text: str, source_group: str) -> tuple[str, str, str]:
    lower = f"{identity.get('title') or ''}\n{full_text[:30000]}".casefold()
    if re.search(r"\breview\b|\bperspective\b|\bcommentary\b", lower):
        genre = "review_or_perspective"
    elif re.search(r"\babstract\b|\bexperimental\b|\bresults\b", lower):
        genre = "primary_experimental"
    else:
        genre = "unknown"
    if source_group == "battery_transfer_curated":
        return "battery", genre, "descriptor_prior_only" if genre != "review_or_perspective" else "evidence_only"
    if re.search(r"calcium-mediated|ca-mediated", lower):
        return "CaNRR", genre, "excluded"
    if re.search(r"lithium[- ]mediated|lithium redox[- ]mediated|linrr|li-nrr", lower) and re.search(r"\bn2\b|nitrogen", lower):
        reaction, role = "LiNRR", "target_label_candidate"
    elif re.search(r"nitrate|nitrite|nitric oxide|nox", lower):
        reaction, role = "other_fixed_nitrogen", "excluded"
    else:
        reaction, role = "unclear", "excluded"
    if genre == "review_or_perspective":
        role = "evidence_only"
    return reaction, genre, role


FE_PATTERN = re.compile(r"(?:faradaic\s+efficien(?:cy|cies)|faradaic\s+selectivity|\bFE(?:\s*NH3)?\b)[^%\n]{0,70}?(\d+(?:\.\d+)?)\s*%|(\d+(?:\.\d+)?)\s*%[^.\n]{0,45}(?:faradaic\s+efficien|\bFE\b)", re.I)


def extract_curated_records(manifests: list[dict[str, Any]]) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    for manifest in manifests:
        identity = manifest.get("identity") or derive_paper_identity(manifest)
        all_text = "\n".join(_asset_text(asset, 8) for asset in manifest["assets"] if asset["format"] != "zip")
        reaction, genre, role = classify_document(identity, all_text, manifest["source_group"])
        if manifest["source_group"] == "battery_transfer_curated":
            continue
        bundle_record_count = 0
        for asset in manifest["assets"]:
            if asset["format"] not in {"pdf", "docx", "xlsx", "csv", "txt"}:
                continue
            try:
                blocks = read_asset_blocks(asset["absolute_path"])
            except Exception:
                continue
            for block in blocks:
                text = str(block.get("text") or "")
                matches = list(FE_PATTERN.finditer(text))
                for match_i, match in enumerate(matches, 1):
                    value = float(match.group(1) or match.group(2))
                    if not 0 <= value <= 100:
                        continue
                    start, end = max(0, match.start() - 700), min(len(text), match.end() + 700)
                    excerpt = " ".join(text[start:end].split())
                    record = _empty_record()
                    match_info = manifest.get("identity_match") or {}
                    resolved_paper_id = match_info.get("existing_paper_id")
                    evidence_owner_id = resolved_paper_id or manifest["provisional_bundle_id"]
                    record.update({
                        "record_id": stable_id("REC", manifest["bundle_id"], asset["source_asset_id"], block["locator"], match_i, value),
                        "paper_id": resolved_paper_id, "provisional_bundle_id": manifest["provisional_bundle_id"],
                        "source_group": manifest["source_group"], "reaction_family": reaction, "document_genre": genre,
                        "training_role": role, "source_asset_id": asset["source_asset_id"], "source_filename": asset["relative_path"],
                        "source_file_sha256": asset["sha256"], "source_locator": block["locator"], "source_table": block.get("table"),
                        "source_page": block.get("page"), "source_text_excerpt": excerpt, "fe_nh3_percent": value,
                        "evidence_authority": "primary_structured_author_experiment" if asset["probable_role"] in {"supporting_information", "supplementary_data", "source_data"} and genre == "primary_experimental" else "primary_text_candidate",
                        "extraction_status": "DETERMINISTIC_CANDIDATE", "review_status": "PENDING_REVIEW",
                        "source_bundle_id": manifest["bundle_id"], "evidence_owner_id": evidence_owner_id,
                    })
                    context = " ".join(text[max(0, start - 1200):min(len(text), end + 1200)].split())
                    _extract_conditions(context, record)
                    lower = excerpt.casefold()
                    if asset["probable_role"] == "main_article" and re.search(r"reported|previous|literature|according to", lower) and not re.search(r"we |our |this work|herein|achiev|demonstrat", lower):
                        record["ambiguity_flags"].append("TARGET_OWNERSHIP_UNCLEAR")
                    records.append(record)
                    bundle_record_count += 1
        if bundle_record_count == 0:
            main = next((a for a in manifest["assets"] if a["probable_role"] == "main_article"), manifest["assets"][0] if manifest["assets"] else None)
            record = _empty_record()
            resolved_paper_id = (manifest.get("identity_match") or {}).get("existing_paper_id")
            record.update({
                "record_id": stable_id("REC", manifest["bundle_id"], "NO_EXPLICIT_FE"),
                "paper_id": resolved_paper_id,
                "provisional_bundle_id": manifest["provisional_bundle_id"], "source_group": manifest["source_group"],
                "reaction_family": reaction, "document_genre": genre, "training_role": role,
                "source_asset_id": main.get("source_asset_id") if main else None, "source_filename": main.get("relative_path") if main else None,
                "source_file_sha256": main.get("sha256") if main else None, "source_locator": "document-level:no-explicit-fe-found",
                "source_text_excerpt": None, "evidence_authority": "document_triage_only", "extraction_status": "NO_EXPLICIT_FE_FOUND",
                "review_status": "REVIEW_REQUIRED", "source_bundle_id": manifest["bundle_id"],
                "evidence_owner_id": resolved_paper_id or manifest["provisional_bundle_id"],
            })
            records.append(record)
    mark_duplicates_and_conflicts(records)
    return records


LEGACY_MAPPING_SPEC = [
    ("paper_id", "paper_id", "identity", "grouping/provenance only"),
    ("span_id", "source_locator", "identity", "exact legacy evidence anchor"),
    ("reaction_family", "reaction_family", "identity", "LiNRR family required"),
    ("evidence_type", "document_genre", "primary_claim -> primary experimental candidate", "primary ownership gate"),
    ("electrolyte", "electrolyte_notes", "verbatim", "electrolyte still must be sufficiently defined"),
    ("current_density_mA_cm2", "current_density_mA_cm2", "none", "condition identity"),
    ("faradaic_efficiency_percent", "fe_nh3_percent", "none", "FE availability gate"),
    ("nh3_yield_value", "nh3_rate_value", "none", "secondary target candidate only"),
    ("nh3_yield_unit", "nh3_rate_unit", "verbatim; no incompatible-unit conversion", "secondary target unit coherence"),
    ("stability_hours", "stability_duration_h", "none", "QC/context only"),
    ("isotope_validation", "isotope_validation", "none", "eligibility tier, not row-existence gate"),
    ("blank_control", "blank_control", "none", "eligibility tier"),
    ("detection_method", "quantification_method", "none", "eligibility tier"),
    ("source_text", "source_text_excerpt", "verbatim exact span", "source attribution"),
]


def inspect_legacy_mapping(path: str | Path) -> list[dict[str, str]]:
    with Path(path).open("r", encoding="utf-8") as handle:
        first = json.loads(next(line for line in handle if line.strip()))
    return [
        {"legacy_field": old, "unified_field": new, "conversion": conversion, "admission_impact": impact}
        for old, new, conversion, impact in LEGACY_MAPPING_SPEC if old in first
    ]


def ingest_legacy_candidates(path: str | Path) -> list[dict[str, Any]]:
    path = Path(path)
    digest = sha256_file(path)
    records: list[dict[str, Any]] = []
    with path.open("r", encoding="utf-8") as handle:
        for line_i, line in enumerate(handle, 1):
            if not line.strip():
                continue
            source = json.loads(line)
            record = _empty_record()
            reaction = source.get("reaction_family")
            if reaction in {None, "", "unclear"}:
                reaction = source.get("paper_level_reaction_family") or "unclear"
            primary = source.get("evidence_type") == "primary_claim" and source.get("allow_field_extraction") is not False
            role = "target_label_candidate" if reaction == "LiNRR" and primary else ("evidence_only" if primary else "excluded")
            source_text = source.get("source_text") or source.get("text")
            legacy_paper_id = source.get("paper_id")
            registry_match = re.match(r"^(P\d{4})(?:_|$)", str(legacy_paper_id or ""))
            canonical_paper_id = registry_match.group(1) if registry_match else legacy_paper_id
            record.update({
                "record_id": stable_id("LEGACYREC", source.get("span_id"), line_i), "paper_id": canonical_paper_id,
                "provisional_bundle_id": None, "source_group": "legacy_oa", "reaction_family": reaction,
                "document_genre": "primary_experimental_candidate" if primary else "unknown", "training_role": role,
                "source_asset_id": stable_id("ASSET", str(path.resolve()), digest), "source_filename": path.name,
                "source_file_sha256": digest, "source_locator": source.get("span_id") or f"line:{line_i}",
                "source_text_excerpt": source_text, "evidence_authority": "legacy_candidate_evidence_not_gold_training",
                "extraction_status": "LEGACY_CANDIDATE_MAPPED", "review_status": "PENDING_REVIEW",
                "electrolyte_notes": source.get("electrolyte"), "current_density_mA_cm2": source.get("current_density_mA_cm2"),
                "fe_nh3_percent": source.get("faradaic_efficiency_percent"), "nh3_rate_value": source.get("nh3_yield_value"),
                "nh3_rate_unit": source.get("nh3_yield_unit"), "nh3_rate_raw_value": source.get("nh3_yield_value"),
                "nh3_rate_raw_unit": source.get("nh3_yield_unit"), "stability_duration_h": source.get("stability_hours"),
                "isotope_validation": source.get("isotope_validation"), "blank_control": source.get("blank_control"),
                "quantification_method": source.get("detection_method"), "source_bundle_id": canonical_paper_id,
                "evidence_owner_id": canonical_paper_id,
            })
            _extract_conditions(str(source_text or "") + " " + str(source.get("electrolyte") or ""), record)
            records.append(record)
    mark_duplicates_and_conflicts(records)
    return records


def mark_duplicates_and_conflicts(records: list[dict[str, Any]]) -> None:
    for record in records:
        record["ambiguity_flags"] = [flag for flag in (record.get("ambiguity_flags") or []) if flag not in {"POSSIBLE_DUPLICATE_ROW", "SOURCE_CONFLICT"}]
    duplicates: dict[tuple[Any, ...], list[dict[str, Any]]] = {}
    condition_targets: dict[tuple[Any, ...], set[Any]] = {}
    for record in records:
        owner = record.get("paper_id") or record.get("source_bundle_id") or record.get("provisional_bundle_id")
        key = (owner, record.get("lithium_salt"), record.get("solvent"), record.get("current_density_mA_cm2"), record.get("duration_h"), record.get("fe_nh3_percent"))
        if key[-1] is not None:
            duplicates.setdefault(key, []).append(record)
        condition_key = key[:-1]
        if any(value is not None for value in condition_key[1:]) and record.get("fe_nh3_percent") is not None:
            condition_targets.setdefault(condition_key, set()).add(record["fe_nh3_percent"])
    authority_rank = {"primary_structured_author_experiment": 0, "primary_text_candidate": 1, "legacy_candidate_evidence_not_gold_training": 2}
    for same_rows in duplicates.values():
        if len(same_rows) < 2:
            continue
        preferred = min(same_rows, key=lambda row: (authority_rank.get(str(row.get("evidence_authority")), 9), str(row.get("source_filename") or ""), str(row.get("source_locator") or ""), str(row.get("record_id"))))
        for row in same_rows:
            if row is not preferred:
                row["ambiguity_flags"].append("POSSIBLE_DUPLICATE_ROW")
    conflicts = {key for key, targets in condition_targets.items() if len(targets) > 1}
    for record in records:
        owner = record.get("paper_id") or record.get("source_bundle_id") or record.get("provisional_bundle_id")
        key = (owner, record.get("lithium_salt"), record.get("solvent"), record.get("current_density_mA_cm2"), record.get("duration_h"))
        if key in conflicts:
            record["ambiguity_flags"].append("SOURCE_CONFLICT")


def apply_admission(record: dict[str, Any]) -> dict[str, Any]:
    flags = set(record.get("ambiguity_flags") or [])
    primary = record.get("document_genre") in {"primary_experimental", "primary_experimental_candidate"}
    record_owner = record.get("paper_id") or record.get("provisional_bundle_id")
    evidence_owner = record.get("evidence_owner_id")
    gates = {
        "primary_experimental": primary,
        "linrr_n2_to_nh3": record.get("reaction_family") == "LiNRR",
        "author_owned_target": "TARGET_OWNERSHIP_UNCLEAR" not in flags,
        "fe_available": record.get("fe_nh3_percent") is not None,
        "electrolyte_defined": bool(record.get("lithium_salt") and record.get("solvent")),
        "condition_defined": any(record.get(field) is not None for field in ("current_density_mA_cm2", "duration_h", "lithium_salt_concentration_mol_L", "proton_donor", "pressure_bar", "gas_flow_sccm")),
        "exact_source": bool(record.get("source_asset_id") and record.get("source_locator") and record.get("source_file_sha256")),
        "ownership_clear": "SOURCE_CONFLICT" not in flags and bool(record_owner) and evidence_owner == record_owner,
        "duplicate_clear": "POSSIBLE_DUPLICATE_ROW" not in flags,
        "signal_valid": record.get("failure_type") not in {"CONTAMINATION_INVALIDATED", "QC_INVALIDATED"},
    }
    required = [name for name, passed in gates.items() if not passed]
    record["missing_required_fields"] = required
    eligible = all(gates.values()) and record.get("training_role") == "target_label_candidate"
    strong_qc = sum(bool(record.get(field)) for field in ("isotope_validation", "blank_control", "nitrogen_source_validation", "quantification_method")) >= 3
    if eligible:
        tier = "TIER_A" if strong_qc and record.get("evidence_authority") == "primary_structured_author_experiment" else "TIER_B"
    elif record.get("training_role") == "target_label_candidate" and primary and gates["linrr_n2_to_nh3"] and gates["fe_available"]:
        tier = "TIER_C"
    else:
        tier = "EXCLUDED"
    record["eligibility_tier"] = tier
    record["model_eligible_fe"] = eligible
    if eligible:
        record["review_status"] = "MODEL_ELIGIBLE_FE"
    return {"record_id": record["record_id"], "eligibility_tier": tier, "model_eligible_fe": eligible, "gate_results": gates, "reasons": required}


def build_battery_additive_prior(manifests: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    chemical_patterns = [
        ("LiNO3", "lithium nitrate", "inorganic salt"), ("FEC", "fluoroethylene carbonate", "carbonate"),
        ("VC", "vinylene carbonate", "carbonate"), ("LiDFOB", "lithium difluoro(oxalato)borate", "borate salt"),
        ("TMSPi", "tris(trimethylsilyl) phosphite", "phosphite"), ("pentaerythritol", "pentaerythritol", "polyol"),
        ("PTT", "pentaerythritol", "polyol"), ("tert-butanol", "tert-butanol", "alcohol"),
    ]
    for manifest in manifests:
        if manifest["source_group"] != "battery_transfer_curated":
            continue
        identity = manifest.get("identity") or {}
        for asset in manifest["assets"]:
            if asset["format"] not in {"pdf", "docx", "xlsx", "csv", "txt"}:
                continue
            try:
                blocks = read_asset_blocks(asset["absolute_path"])
            except Exception:
                continue
            found: set[tuple[str, str]] = set()
            for block in blocks:
                text = str(block.get("text") or "")
                for token, canonical, chemical_class in chemical_patterns:
                    if (token.casefold(), block["locator"]) in found or not re.search(rf"(?<!\w){re.escape(token)}(?!\w)", text, re.I):
                        continue
                    found.add((token.casefold(), block["locator"]))
                    lower = text.casefold()
                    fixed_n = canonical == "lithium nitrate" or bool(re.search(r"(?:nitrate|nitrite|nitrogen-containing)", canonical, re.I))
                    rows.append({
                        "paper_identity": identity.get("doi") or identity.get("title") or manifest["provisional_bundle_id"],
                        "chemical_name": token, "canonical_name": canonical, "chemical_class": chemical_class,
                        "functional_class": "interphase/solvation additive", "base_salt": _first_term(text, ["ZnSO4", "LiPF6", "LiFSI", "LiTFSI"]),
                        "base_solvent": _first_term(text, ["water", "DME", "EC", "DMC", "ether"]),
                        "additive_concentration": _nearby_concentration(text, token), "battery_system": "aqueous_zinc" if "zinc" in lower or "zn" in lower else "lithium_battery",
                        "electrode_system": "Zn" if "zinc" in lower or "zn" in lower else "Li",
                        "reported_interphase_effect": _sentence_fragment(text, token, ("interphase", "SEI")),
                        "reported_solvation_effect": _sentence_fragment(text, token, ("solvation",)),
                        "reported_li_deposition_effect": _sentence_fragment(text, token, ("deposition", "plating", "stripping")),
                        "reported_transport_effect": _sentence_fragment(text, token, ("transport", "conductivity")),
                        "target_metric_name": "battery metric; never FE_NH3", "target_metric_value": None, "target_metric_unit": None,
                        "fixed_nitrogen_risk": fixed_n, "linrr_transfer_hypothesis": "EXCLUDED_FROM_AUTOMATIC_LINRR_RECOMMENDATION" if fixed_n else "descriptor prior only; requires future source-attribution controls",
                        "source_asset": asset["relative_path"], "source_locator": block["locator"], "source_bundle_id": manifest["bundle_id"],
                    })
    unique: dict[tuple[Any, ...], dict[str, Any]] = {}
    for row in rows:
        key = (row["paper_identity"], row["canonical_name"], row["source_asset"], row["source_locator"])
        unique[key] = row
    return [unique[key] for key in sorted(unique, key=lambda value: tuple(str(x) for x in value))]


def _nearby_concentration(text: str, token: str) -> str | None:
    match = re.search(rf"(?:\d+(?:\.\d+)?\s*(?:M|mM|mol\s*%|wt%|vol%)[^.;]{{0,30}}{re.escape(token)}|{re.escape(token)}[^.;]{{0,30}}\d+(?:\.\d+)?\s*(?:M|mM|mol\s*%|wt%|vol%))", text, re.I)
    return " ".join(match.group(0).split()) if match else None


def _sentence_fragment(text: str, token: str, concepts: tuple[str, ...]) -> str | None:
    for sentence in re.split(r"(?<=[.!?])\s+", text):
        if token.casefold() in sentence.casefold() and any(concept.casefold() in sentence.casefold() for concept in concepts):
            return " ".join(sentence.split())[:500]
    return None


def source_inventory(paths: Iterable[str | Path]) -> list[dict[str, Any]]:
    inventory = []
    for value in paths:
        path = Path(value)
        row: dict[str, Any] = {"path": str(path.resolve()), "exists": path.exists(), "size": None, "sha256": None, "record_count": None, "schema_fields": [], "scientific_role": _scientific_role(path)}
        if path.is_file():
            row["size"] = path.stat().st_size
            row["sha256"] = sha256_file(path)
            suffix = path.suffix.lower()
            if suffix == ".jsonl":
                fields: set[str] = set()
                count = 0
                with path.open("r", encoding="utf-8") as handle:
                    for line in handle:
                        if not line.strip():
                            continue
                        count += 1
                        if count <= 25:
                            item = json.loads(line)
                            if isinstance(item, dict): fields.update(item)
                row["record_count"], row["schema_fields"] = count, sorted(fields)
            elif suffix == ".csv":
                with path.open("r", encoding="utf-8-sig", newline="") as handle:
                    reader = csv.reader(handle)
                    header = next(reader, [])
                    row["record_count"] = sum(1 for _ in reader)
                    row["schema_fields"] = header
        inventory.append(row)
    return inventory


def _scientific_role(path: Path) -> str:
    name = path.name
    if "experiment_triage" in name:
        return "legacy experiment-triage candidates; not Gold training rows"
    if "performance_ledger" in name:
        return "legacy performance evidence candidates"
    if "semantic_spans" in name:
        return "cleanroom exact evidence anchors"
    if "evidence_links" in name:
        return "within-paper evidence links; cross-paper links forbidden"
    if "human_gold_claim_rights" in name:
        return "paper/span claim-right review; not experiment-row Gold"
    if "reviewed_audit_records" in name:
        return "human audit evidence; not automatic training admission"
    return "legacy context asset"


def write_jsonl(rows: Iterable[dict[str, Any]], path: str | Path) -> None:
    path = Path(path); path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="\n") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False, sort_keys=True, separators=(",", ":")) + "\n")


def _csv_value(value: Any) -> Any:
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return value


def write_csv(rows: list[dict[str, Any]], path: str | Path, fields: list[str] | None = None) -> None:
    path = Path(path); path.parent.mkdir(parents=True, exist_ok=True)
    if fields is None:
        fields = sorted({key for row in rows for key in row})
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({field: _csv_value(row.get(field)) for field in fields})


def summarize_dataset(records: list[dict[str, Any]], manifests: list[dict[str, Any]], battery_rows: list[dict[str, Any]], legacy_count: int) -> dict[str, Any]:
    curated = [r for r in records if r.get("source_group") != "legacy_oa"]
    return {
        "schema_version": "linrr_training_dataset_summary_v0", "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "number_of_bundles": len(manifests), "bundles_by_source_group": dict(sorted(Counter(m["source_group"] for m in manifests).items())),
        "number_with_doi": sum(bool((m.get("identity") or {}).get("doi")) for m in manifests),
        "number_matched_existing_paper_id": sum(bool((m.get("identity_match") or {}).get("existing_paper_id")) for m in manifests),
        "legacy_candidate_rows": legacy_count, "new_curated_candidate_rows": len(curated), "condition_level_candidate_records": len(records),
        "linrr_target_rows": sum(r.get("training_role") == "target_label_candidate" for r in records), "battery_prior_rows": len(battery_rows),
        "tier_counts": dict(sorted(Counter(r.get("eligibility_tier") or "UNASSESSED" for r in records).items())),
        "model_eligible_fe_rows": sum(bool(r.get("model_eligible_fe")) for r in records),
        "unique_paper_groups": len({_group(r) for r in records if r.get("model_eligible_fe")}),
        "field_coverage": {field: sum(r.get(field) is not None for r in records) for field in ("fe_nh3_percent", "nh3_rate_value", "current_density_mA_cm2", "lithium_salt_concentration_mol_L", "proton_donor_concentration_value", "water_content_value", "pressure_bar", "gas_flow_sccm")},
        "unique_papers": len({_group(r) for r in records}), "unique_laboratories": None,
        "unique_salts": sorted({str(r["lithium_salt"]) for r in records if r.get("lithium_salt")}),
        "unique_solvents": sorted({str(r["solvent"]) for r in records if r.get("solvent")}),
        "unique_proton_donors": sorted({str(r["proton_donor"]) for r in records if r.get("proton_donor")}),
        "unique_additives": sorted({str(r["additive"]) for r in records if r.get("additive")}),
    }


def _group(record: dict[str, Any]) -> str:
    return str(record.get("paper_id") or record.get("provisional_bundle_id") or record.get("source_bundle_id") or "UNKNOWN")


def write_workbook(path: str | Path, records: list[dict[str, Any]], manifests: list[dict[str, Any]], identity_rows: list[dict[str, Any]], battery_rows: list[dict[str, Any]], mapping_rows: list[dict[str, Any]], summary: dict[str, Any], master_workbook: dict[str, Any] | None = None) -> None:
    from openpyxl import Workbook
    workbook = Workbook()
    workbook.remove(workbook.active)
    readme = [{"key": "dataset", "value": "LiNRR Training Dataset v0"}, {"key": "scientific_boundary", "value": "Candidate discovery is separate from admission; this workbook is not Gold."}, {"key": "model_boundary", "value": "Any fitted predictor is a literature shadow model."}]
    if master_workbook:
        readme.extend({"key": key, "value": value} for key, value in master_workbook.items())
    asset_rows = [{**{k: m.get(k) for k in ("bundle_id", "provisional_bundle_id", "source_group", "bundle_path")}, **a} for m in manifests for a in m["assets"]]
    excluded = [r for r in records if r.get("eligibility_tier") == "EXCLUDED"]
    review = [r for r in records if r.get("review_status") in {"PENDING_REVIEW", "REVIEW_REQUIRED"} or r.get("eligibility_tier") == "TIER_C"]
    electrolyte = taxonomy(records, "lithium_salt", "electrolyte") + taxonomy(records, "solvent", "solvent") + taxonomy(records, "proton_donor", "proton_donor")
    sheets = [
        ("README", readme), ("PAPERS", identity_rows), ("ASSETS", asset_rows), ("ALL_RECORDS", records),
        ("MODEL_ELIGIBLE", [r for r in records if r.get("model_eligible_fe")]), ("EXCLUDED", excluded),
        ("REVIEW_QUEUE", review), ("ELECTROLYTE_TAXONOMY", electrolyte), ("BATTERY_ADDITIVE_PRIOR", battery_rows),
        ("LEGACY_MAPPING", mapping_rows), ("DATASET_SUMMARY", [{"metric": key, "value": _csv_value(value)} for key, value in summary.items()]),
    ]
    for title, rows in sheets:
        sheet = workbook.create_sheet(title)
        fields = list(rows[0]) if rows else ["EMPTY"]
        sheet.append(fields)
        for row in rows:
            sheet.append([_excel_value(row.get(field)) for field in fields])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    path = Path(path); path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def _excel_value(value: Any) -> Any:
    if isinstance(value, (list, dict)):
        value = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    if isinstance(value, str):
        value = re.sub(r"[\x00-\x08\x0B\x0C\x0E-\x1F]", "�", value)
        if len(value) > 32767:
            return value[:32750] + "…[TRUNCATED IN XLSX]"
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        return None
    return value


def taxonomy(records: list[dict[str, Any]], field: str, kind: str) -> list[dict[str, Any]]:
    counts = Counter(str(r[field]) for r in records if r.get(field))
    return [{"taxonomy_type": kind, "canonical_name": name, "record_count": count} for name, count in sorted(counts.items())]


def find_master_workbook(search_root: str | Path, excluded_path: str | Path) -> dict[str, Any] | None:
    excluded = Path(excluded_path).resolve()
    candidates = [p for p in Path(search_root).rglob("*.xlsx") if p.resolve() != excluded and "master" in p.name.casefold()]
    if not candidates:
        return None
    selected = sorted(candidates, key=lambda p: str(p).casefold())[0]
    return {"existing_master_workbook_path": str(selected.resolve()), "existing_master_workbook_sha256": sha256_file(selected)}


def write_source_inventory_reports(inventory: list[dict[str, Any]], report_dir: str | Path) -> None:
    report_dir = Path(report_dir); report_dir.mkdir(parents=True, exist_ok=True)
    (report_dir / "source_inventory.json").write_text(json.dumps(inventory, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    lines = ["# LiNRR Training v0 source inventory", "", "Required legacy assets were audited in place. Missing assets are recorded as `MISSING` and are never regenerated or silently substituted.", "", "| Path | Status | Bytes | SHA256 | Records | Scientific role |", "|---|---:|---:|---|---:|---|"]
    for row in inventory:
        lines.append(f"| `{row['path']}` | {'EXISTS' if row['exists'] else 'MISSING'} | {row['size'] or ''} | {row['sha256'] or ''} | {row['record_count'] if row['record_count'] is not None else ''} | {row['scientific_role']} |")
        if row.get("schema_fields"):
            lines.extend(["", f"Fields for `{row['path']}`: `{', '.join(row['schema_fields'])}`", ""])
    (report_dir / "source_inventory.md").write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")


def write_jacs_method_map(manifest: dict[str, Any], output_path: str | Path) -> None:
    csv_assets = [a for a in manifest["assets"] if a["format"] == "csv"]
    descriptions = []
    for asset in csv_assets:
        blocks = read_asset_blocks(asset["absolute_path"])
        header = blocks[0]["text"] if blocks else ""
        descriptions.append((asset, max(len(blocks) - 1, 0), header))
    training = next((item for item in descriptions if "CE" in item[2] or "LCE" in item[2]), None)
    candidates = next((item for item in descriptions if item is not training), None)
    lines = [
        "# JACS 10.1021/jacs.6c08057 methodological architecture map", "",
        "This map is derived from the bundle's article, SI, and CSV files. It transfers workflow architecture only. Zinc-battery labels are not LiNRR targets, and no prose or figures are copied.", "",
        "## Actual-file inventory", "",
    ]
    for asset in manifest["assets"]:
        lines.append(f"- `{asset['relative_path']}` — {asset['probable_role']}; SHA256 `{asset['sha256']}`")
    lines.extend(["", "## Data and validation hierarchy", ""])
    if training:
        lines.append(f"- Training/testing source: `{training[0]['relative_path']}` has {training[1]} formulations. Its columns span additive identity and concentration, interfacial descriptors (adsorption energy and charge transfer), molecular/electronic/physical descriptors, measured CE, and the derived LCE label.")
    if candidates:
        lines.append(f"- Candidate-screening source: `{candidates[0]['relative_path']}` has {candidates[1]} candidates with the descriptor columns but no measured CE/LCE label columns.")
    lines.extend([
        "- Descriptor hierarchy: formulation identity/concentration → interface descriptors → intrinsic molecular/electronic/physical descriptors → composite interface-energy descriptor → battery label.",
        "- Train/test organization: the article reports 84 measured electrolyte formulations and nested five-fold validation; composite-descriptor coefficients are fitted inside training folds so held-out labels do not enter descriptor fitting.",
        "- Candidate organization: a separate unlabeled descriptor table is scored after model development.",
        "- Experimental validation: the selected zinc candidate is tested with plating/stripping CE, asymmetric/symmetric cycling, full-cell behavior, and interface/solvation characterization.",
        "", "## Translation to LiNRR", "",
        "| JACS zinc concept | Possible LiNRR analogue | Available now? | Source | Future measurement/computation |",
        "|---|---|---:|---|---|",
        "| Additive/formulation hierarchy | Li salt–solvent–proton donor–additive recipe | Partial | Curated LiNRR condition records | Standardized formulation matrix |",
        "| Adsorption energy at Zn | Additive/species interaction with Li or reported SEI phases | No | Not reliably present in v0 | Surface-specific DFT under declared state |",
        "| Interfacial charge transfer | Interface electronic/ionic-transfer proxy | No | Not reliably present in v0 | Operando spectroscopy/electrochemistry or validated computation |",
        "| Intrinsic molecular descriptors | Small, preregistered donor/additive descriptors | No | Chemical identities only in v0 | Curated structures and reproducible descriptor calculation |",
        "| CE/LCE battery label | FE_NH3 under an exact LiNRR condition | Partial | Primary author experiment anchors | Harmonized FE protocol and replicate metadata |",
        "| Nested held-out validation | Paper-grouped outer validation with fold-local preprocessing | Yes | Shadow predictor v0 code | Prospective laboratory holdout later |",
        "| Unlabeled candidate table | Future controlled LiNRR candidate set | No | Screening prohibited in this milestone | Define OOD and fixed-N attribution controls first |",
        "| Battery cycling validation | FE, NH3 rate, stability, blanks, gas/source and isotope controls | Partial | Record-level QC fields | Prospective standardized validation campaign |",
        "", "The method is a blueprint for hierarchy and leakage control only; it does not authorize additive recommendation or zinc-to-LiNRR label transfer.",
    ])
    output_path = Path(output_path); output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def deterministic_dataset_hash(records: list[dict[str, Any]]) -> str:
    payload = "\n".join(json.dumps(row, ensure_ascii=False, sort_keys=True, separators=(",", ":")) for row in sorted(records, key=lambda row: str(row.get("record_id"))))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()
