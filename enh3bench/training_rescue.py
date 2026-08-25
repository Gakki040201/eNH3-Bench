"""Deterministic LiNRR v0.1 scientific-record rescue helpers.

The v0.1 profile is deliberately independent from the v0 extractor.  It adds
row-wise structured-table parsing, same-paper-series inheritance with dual
locators, conservative ownership adjudication, and human-review packets.
"""

from __future__ import annotations

import csv
import json
import re
from copy import deepcopy
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

from .training_ingest import (
    UNIFIED_FIELDS,
    _empty_record,
    _extract_conditions,
    apply_admission,
    normalize_measurement,
    read_asset_blocks,
    stable_id,
)


SCHEMA_VERSION = "linrr_experiment_record_v0_1"
NOT_PRESENT_IN_RECIPE = "NOT_PRESENT_IN_RECIPE"
NOT_REPORTED = "NOT_REPORTED"
UNKNOWN = "UNKNOWN"

PROVENANCE_FIELDS = [
    "condition_value_origin",
    "condition_inherited_from_asset",
    "condition_inherited_from_locator",
    "condition_direct_locator",
    "ownership_support_locator",
    "supporting_source_locators",
    "original_headers",
    "original_units",
    "component_reporting_status",
    "duplicate_classification",
    "source_conflict_classification",
    "human_review_required",
    "review_notes",
]
V01_FIELDS = [*UNIFIED_FIELDS, *PROVENANCE_FIELDS]

RESCUE_QUEUE_FIELDS = [
    "record_id", "paper_id", "bundle_id", "paper_title", "doi", "source_group",
    "source_asset", "source_locator", "source_page", "source_table", "fe_nh3_percent",
    "current_admission_failures", "lithium_salt", "salt_concentration", "solvent",
    "cosolvent", "proton_donor", "proton_donor_concentration", "additive",
    "additive_concentration", "water", "oxygen", "current_density", "charge",
    "duration", "temperature", "pressure", "gas_flow", "liquid_flow",
    "current_ambiguity_flags", "rescue_status", "rescued_from_asset",
    "rescued_from_locator", "rescued_fields", "remaining_failures",
    "human_review_required", "review_notes",
]


@dataclass(frozen=True)
class StructuredTable:
    number: int
    name: str
    caption: str | None
    rows: list[list[str]]
    source_asset: str


def _text(value: object) -> str:
    return " ".join(str(value or "").replace("\u2212", "-").split())


def _json_list(value: object) -> list[str]:
    if isinstance(value, list):
        return [str(item) for item in value]
    text = str(value or "").strip()
    if not text:
        return []
    try:
        parsed = json.loads(text)
    except json.JSONDecodeError:
        return [text]
    return [str(item) for item in parsed] if isinstance(parsed, list) else [str(parsed)]


def load_records_csv(path: str | Path) -> list[dict[str, Any]]:
    with Path(path).open("r", encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    list_fields = {"missing_required_fields", "reported_not_measured_fields", "ambiguity_flags"}
    float_fields = {
        "lithium_salt_concentration_mol_L", "lithium_salt_concentration_raw_value",
        "proton_donor_concentration_value", "proton_donor_concentration_raw_value",
        "additive_concentration_value", "additive_concentration_raw_value",
        "water_content_value", "water_content_raw_value", "current_density_mA_cm2",
        "current_density_raw_value", "current_A", "electrode_area_cm2", "total_charge_C",
        "duration_h", "duration_raw_value", "temperature_C", "pressure_bar",
        "pressure_raw_value", "gas_flow_sccm", "liquid_flow_mL_min", "fe_nh3_percent",
        "nh3_rate_value", "nh3_rate_raw_value", "replicate_n",
    }
    bool_fields = {"model_eligible_fe", "interface_sei_characterization_present", "interface_lif_reported", "interface_n_containing_sei_reported"}
    for row in rows:
        for field in list_fields:
            row[field] = _json_list(row.get(field))
        for field in float_fields:
            raw = str(row.get(field) or "").strip()
            row[field] = float(raw) if raw else None
        for field in bool_fields:
            row[field] = str(row.get(field) or "").casefold() in {"true", "1", "yes"}
        for field, value in list(row.items()):
            if value == "":
                row[field] = None
    return rows


def _docx_tables(path: Path) -> list[StructuredTable]:
    from docx import Document
    document = Document(path)
    captions = [_text(p.text) for p in document.paragraphs if re.match(r"^Table\s+S?\d+\b", _text(p.text), re.I)]
    output: list[StructuredTable] = []
    for table_i, table in enumerate(document.tables, 1):
        rows = [[_text(cell.text) for cell in row.cells] for row in table.rows]
        caption = captions[table_i - 1] if table_i <= len(captions) else None
        name_match = re.match(r"^(Table\s+S?\d+)", caption or "", re.I)
        output.append(StructuredTable(table_i, name_match.group(1) if name_match else f"Table {table_i}", caption, rows, path.name))
    return output


def _xlsx_tables(path: Path) -> list[StructuredTable]:
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    output: list[StructuredTable] = []
    try:
        for sheet_i, sheet in enumerate(workbook.worksheets, 1):
            rows = [[_text(value) for value in row] for row in sheet.iter_rows(values_only=True)]
            rows = [row for row in rows if any(row)]
            output.append(StructuredTable(sheet_i, sheet.title, None, rows, path.name))
    finally:
        workbook.close()
    return output


def _csv_tables(path: Path) -> list[StructuredTable]:
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
        rows = [[_text(cell) for cell in row] for row in csv.reader(handle)]
    return [StructuredTable(1, path.stem, None, [row for row in rows if any(row)], path.name)]


def _pdf_tables(path: Path) -> list[StructuredTable]:
    """Recover only PDF tables whose extracted text retains column separators."""
    try:
        import fitz  # type: ignore
    except ImportError:  # pragma: no cover - optional dependency path
        return []
    output: list[StructuredTable] = []
    with fitz.open(path) as document:
        for page_i, page in enumerate(document, 1):
            lines = [line.rstrip() for line in page.get_text("text").splitlines()]
            table_i = 0
            for line_i, line in enumerate(lines):
                if not re.search(r"\t|\s{2,}", line):
                    continue
                headers = [_text(cell) for cell in re.split(r"\t+|\s{2,}", line) if _text(cell)]
                aliases = [canonical_header(header) for header in headers]
                informative = {alias for alias in aliases if alias}
                if "fe_nh3_percent" not in informative or len(informative) < 2:
                    continue
                rows = [headers]
                for following in lines[line_i + 1:line_i + 31]:
                    if not following.strip():
                        break
                    cells = [_text(cell) for cell in re.split(r"\t+|\s{2,}", following) if _text(cell)]
                    if len(cells) < len(headers) or _number(cells[aliases.index("fe_nh3_percent")]) is None:
                        break
                    rows.append(cells[:len(headers)])
                if len(rows) > 1:
                    table_i += 1
                    output.append(StructuredTable(table_i, f"PDF page {page_i} table {table_i}", None, rows, path.name))
    return output


def read_structured_tables(path: str | Path) -> list[StructuredTable]:
    path = Path(path)
    if path.suffix.casefold() == ".docx":
        return _docx_tables(path)
    if path.suffix.casefold() == ".xlsx":
        return _xlsx_tables(path)
    if path.suffix.casefold() == ".csv":
        return _csv_tables(path)
    if path.suffix.casefold() == ".pdf":
        return _pdf_tables(path)
    return []


def split_header(header: str) -> tuple[str, str | None]:
    text = _text(header)
    match = re.search(r"\(([^()]*(?:%|bar|mA|A|h|C|M|mol|ppm|sccm|mL)[^()]*)\)\s*$", text, re.I)
    return (text[:match.start()].strip(), match.group(1).strip()) if match else (text, None)


def canonical_header(header: str) -> str | None:
    name, _ = split_header(header)
    key = re.sub(r"[^a-z0-9+#]+", " ", name.casefold()).strip()
    rules = [
        (r"^(?:fe|fe nh3|nh3 fe|faradaic efficien(?:cy|cies)|faradaic selectivity)$", "fe_nh3_percent"),
        (r"^average (?:fe|faradaic efficien(?:cy|cies))$", "average_fe_nh3_percent"),
        (r"^(?:j|current density)$", "current_density_mA_cm2"),
        (r"^(?:electric current|current)$", "current_A"),
        (r"^(?:time|duration|test time)$", "duration_h"),
        (r"^(?:charge|charge passed|charge quantity|total charge)$", "total_charge_C"),
        (r"^(?:pressure|n2 pressure|supplied n2)$", "pressure_bar"),
        (r"^(?:gas flow|n2 flow)$", "gas_flow_sccm"),
        (r"^(?:liquid flow|electrolyte flow)$", "liquid_flow_mL_min"),
        (r"^(?:temperature)$", "temperature_C"),
        (r"^(?:test area|electrode area|area)$", "electrode_area_cm2"),
        (r"^(?:nh3 rate|nh3 yield rate|yield rate|average nh3 yield rate|average yield rate)$", "nh3_rate_value"),
        (r"^(?:electrolyte|formulation|recipe)$", "electrolyte_notes"),
        (r"^(?:lif?bf4|lif?clo4|litfsi) concentration$", "lithium_salt_concentration_mol_L"),
        (r"^(?:salt concentration)$", "lithium_salt_concentration_mol_L"),
        (r"^(?:etoh|ethanol|proton donor)$", "proton_donor"),
        (r"^(?:h2o|water|water content)$", "water_content_value"),
        (r"^(?:replicate|run|#)$", "replicate_n"),
    ]
    for pattern, field in rules:
        if re.match(pattern, key, re.I):
            return field
    return None


def _number(value: object) -> float | None:
    match = re.search(r"[-+]?\d+(?:\.\d+)?", str(value or "").replace(",", ""))
    return float(match.group(0)) if match else None


def _set_measurement(record: dict[str, Any], field: str, raw: str, header: str) -> None:
    number = _number(raw)
    if number is None:
        return
    _, unit = split_header(header)
    record["original_headers"][field] = header
    record["original_units"][field] = unit
    if field == "fe_nh3_percent":
        record[field] = number
    elif field == "current_density_mA_cm2":
        normalized = normalize_measurement(number, unit or "mA/cm2", "current_density")
        record["current_density_raw_value"], record["current_density_raw_unit"] = number, unit
        record[field] = normalized["normalized_value"]
    elif field == "duration_h":
        normalized = normalize_measurement(number, unit or "h", "duration")
        record["duration_raw_value"], record["duration_raw_unit"] = number, unit
        record[field] = normalized["normalized_value"]
    elif field == "pressure_bar":
        normalized = normalize_measurement(number, unit or "bar", "pressure")
        record["pressure_raw_value"], record["pressure_raw_unit"] = number, unit
        record[field] = normalized["normalized_value"]
    elif field == "current_A":
        record[field] = number / 1000.0 if unit and unit.casefold() == "ma" else number
    elif field == "nh3_rate_value":
        record[field], record["nh3_rate_raw_value"], record["nh3_rate_unit"], record["nh3_rate_raw_unit"] = number, number, unit, unit
    else:
        record[field] = number


def _recipe_from_text(text: str, record: dict[str, Any]) -> None:
    _extract_conditions(text, record)
    if re.search(r"\bLiTFSI\b", text, re.I):
        record["lithium_salt"] = "LiTFSI"
    if re.search(r"\bDEE\b|1,2-diethoxyethane", text, re.I):
        record["solvent"] = "DEE"
    if re.search(r"\bTHF\b|\bVTHF\b", text, re.I):
        record["solvent"] = "THF"
    if re.search(r"\bTTE\b|\bVTTE\b", text, re.I):
        if record.get("solvent") == "THF":
            record["cosolvent"] = "TTE"
        elif not record.get("solvent"):
            record["solvent"] = "TTE"
    absent_donor = re.search(r"\b(?:without|no)\s+(n-pentanol|ethanol|EtOH|phenol|water|H2O)\b", text, re.I)
    donor = None if absent_donor else re.search(r"\b(n-pentanol|ethanol|EtOH|phenol|water|H2O)\b", text, re.I)
    if absent_donor:
        record.setdefault("component_reporting_status", {})["proton_donor"] = NOT_PRESENT_IN_RECIPE
        record["proton_donor"] = None
    if donor:
        canonical = {"etoh": "ethanol", "h2o": "water"}.get(donor.group(1).casefold(), donor.group(1))
        record["proton_donor"] = canonical
        concentration = re.search(rf"(\d+(?:\.\d+)?)\s*(vol%|wt%|mM|M)[^.;]{{0,45}}{re.escape(donor.group(1))}|{re.escape(donor.group(1))}[^.;]{{0,45}}?(\d+(?:\.\d+)?)\s*(vol%|wt%|mM|M)", text, re.I)
        if concentration:
            value, unit = (concentration.group(1), concentration.group(2)) if concentration.group(1) else (concentration.group(3), concentration.group(4))
            record["proton_donor_concentration_value"] = float(value)
            record["proton_donor_concentration_unit"] = unit
            record["proton_donor_concentration_raw_value"] = float(value)
            record["proton_donor_concentration_raw_unit"] = unit
    ratio = re.search(r"V?TTE\s*:\s*V?THF\s*=\s*([0-9.]+\s*:\s*[0-9.]+)", text, re.I)
    if ratio:
        record["solvent_ratio"] = ratio.group(1).replace(" ", "")


def ownership_resolution(text: str, structured_author_table: bool = False) -> tuple[str, str | None]:
    lower = _text(text).casefold()
    negative = r"\b(?:literature|previous(?:ly)?|reported by|reference|ref\.?\s*\d+|introduction)\b"
    positive = r"\b(?:our experiments?|we (?:measured|achieved|obtained|performed)|this work|under (?:the )?same .*conditions?|corresponding experiment)\b"
    if re.search(negative, lower) and not re.search(positive, lower):
        return "UNCLEAR_OR_EXTERNAL", None
    if re.search(positive, lower) or structured_author_table:
        return "AUTHOR_OWNED", _text(text)[:500]
    return "UNCLEAR_OR_EXTERNAL", None


def _base_structured_record(manifest: dict[str, Any], asset: dict[str, Any], table: StructuredTable, locator: str) -> dict[str, Any]:
    match = manifest.get("identity_match") or {}
    paper_id = match.get("existing_paper_id")
    owner = paper_id or manifest["provisional_bundle_id"]
    record = _empty_record()
    record.update({
        "schema_version": SCHEMA_VERSION,
        "paper_id": paper_id,
        "provisional_bundle_id": manifest["provisional_bundle_id"],
        "source_group": manifest["source_group"],
        "reaction_family": "LiNRR",
        "document_genre": "primary_experimental",
        "training_role": "target_label_candidate",
        "source_asset_id": asset["source_asset_id"],
        "source_filename": asset["relative_path"],
        "source_file_sha256": asset["sha256"],
        "source_locator": locator,
        "source_table": table.name,
        "source_text_excerpt": table.caption,
        "evidence_authority": "primary_structured_author_experiment",
        "extraction_status": "AUTO_RESOLVED_STRUCTURED_TABLE",
        "review_status": "AUTO_RESOLVED",
        "source_bundle_id": manifest["bundle_id"],
        "evidence_owner_id": owner,
        "condition_value_origin": "DIRECT",
        "condition_direct_locator": locator,
        "condition_inherited_from_asset": None,
        "condition_inherited_from_locator": None,
        "ownership_support_locator": locator,
        "supporting_source_locators": [locator],
        "original_headers": {},
        "original_units": {},
        "component_reporting_status": {},
        "human_review_required": False,
        "review_notes": None,
    })
    return record


def _row_oriented_records(manifest: dict[str, Any], asset: dict[str, Any], table: StructuredTable) -> list[dict[str, Any]]:
    if len(table.rows) < 2:
        return []
    headers = table.rows[0]
    canonical = [canonical_header(header) for header in headers]
    if "fe_nh3_percent" not in canonical:
        return []
    reference_column = next((index for index, header in enumerate(headers) if re.search(r"\breference\b", header, re.I)), None)
    output = []
    for row_i, values in enumerate(table.rows[1:], 2):
        if reference_column is not None:
            reference_value = values[reference_column] if reference_column < len(values) else ""
            if not re.search(r"\bthis work\b", reference_value, re.I):
                continue
        mapping = {field: (values[col_i] if col_i < len(values) else "", headers[col_i]) for col_i, field in enumerate(canonical) if field}
        if _number(mapping.get("fe_nh3_percent", (None, None))[0]) is None:
            continue
        locator = f"{table.name}:row:{row_i}"
        record = _base_structured_record(manifest, asset, table, locator)
        record["record_id"] = stable_id("REC01", manifest["bundle_id"], asset["source_asset_id"], locator)
        for field, (raw, header) in mapping.items():
            if field == "electrolyte_notes":
                record[field] = raw
                _recipe_from_text(raw, record)
            elif field != "average_fe_nh3_percent":
                _set_measurement(record, field, raw, header)
        context = " ".join(filter(None, [table.caption, " | ".join(values)]))
        _recipe_from_text(context, record)
        status, support = ownership_resolution(context, structured_author_table=True)
        if status != "AUTHOR_OWNED":
            record["ambiguity_flags"].append("TARGET_OWNERSHIP_UNCLEAR")
            record["human_review_required"] = True
        elif support:
            record["ownership_support_locator"] = locator
        output.append(record)
    return output


def _transposed_records(manifest: dict[str, Any], asset: dict[str, Any], table: StructuredTable) -> list[dict[str, Any]]:
    if not table.rows or len(table.rows[0]) < 2:
        return []
    labels = [canonical_header(row[0]) if row else None for row in table.rows]
    if "fe_nh3_percent" not in labels:
        return []
    fe_index = labels.index("fe_nh3_percent")
    output = []
    for col_i in range(1, max(len(row) for row in table.rows)):
        fe_raw = table.rows[fe_index][col_i] if col_i < len(table.rows[fe_index]) else ""
        if _number(fe_raw) is None:
            continue
        locator = f"{table.name}:column:{col_i + 1}"
        record = _base_structured_record(manifest, asset, table, locator)
        record["record_id"] = stable_id("REC01", manifest["bundle_id"], asset["source_asset_id"], locator)
        record["replicate_n"] = col_i
        for row_i, field in enumerate(labels):
            if not field or field == "average_fe_nh3_percent":
                continue
            raw = table.rows[row_i][col_i] if col_i < len(table.rows[row_i]) else ""
            if field == "electrolyte_notes":
                record[field] = raw
                _recipe_from_text(raw, record)
            else:
                _set_measurement(record, field, raw, table.rows[row_i][0])
        if record.get("current_A") is not None and record.get("electrode_area_cm2"):
            record["current_density_mA_cm2"] = record["current_A"] * 1000.0 / record["electrode_area_cm2"]
            record["current_density_raw_value"] = record["current_A"] * 1000.0
            record["current_density_raw_unit"] = "mA"
        _recipe_from_text(table.caption or "", record)
        output.append(record)
    return output


def extract_structured_experiments(manifest: dict[str, Any]) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    for asset in manifest.get("assets", []):
        if asset.get("format") not in {"docx", "xlsx", "csv", "pdf"}:
            continue
        for table in read_structured_tables(asset["absolute_path"]):
            row_oriented = _row_oriented_records(manifest, asset, table)
            records.extend(row_oriented if row_oriented else _transposed_records(manifest, asset, table))
    return records


def _series_evidence(manifest: dict[str, Any], record: dict[str, Any]) -> list[dict[str, Any]]:
    solvent = record.get("solvent")
    evidence = []
    for asset in manifest.get("assets", []):
        if asset.get("format") not in {"pdf", "docx", "txt"}:
            continue
        try:
            blocks = read_asset_blocks(asset["absolute_path"])
        except Exception:
            continue
        for block in blocks:
            text = _text(block.get("text"))
            lower = text.casefold()
            if solvent and solvent.casefold() not in lower:
                continue
            if not re.search(r"\b(?:LiBF4|LiClO4|LiTFSI|LiFSI|LiOTf)\b", text, re.I):
                continue
            if not re.search(r"\b(?:in this work|for the experiments? of LMNRR|LMNRR performances?|same Li-NRR conditions|electrolyte preparation)\b", text, re.I):
                continue
            probe = _empty_record()
            _recipe_from_text(text, probe)
            if probe.get("lithium_salt"):
                evidence.append({"asset": asset["relative_path"], "locator": block["locator"], "text": text, "values": probe})
    return evidence


def inherit_same_paper_series(record: dict[str, Any], manifest: dict[str, Any]) -> bool:
    """Inherit only unanimous explicit same-paper experimental-series values."""
    evidence = _series_evidence(manifest, record)
    if not evidence:
        return False
    fields = ["lithium_salt", "lithium_salt_concentration_mol_L", "proton_donor", "proton_donor_concentration_value", "proton_donor_concentration_unit"]
    resolved: dict[str, Any] = {}
    for field in fields:
        values = {item["values"].get(field) for item in evidence if item["values"].get(field) is not None}
        if len(values) > 1:
            return False
        if len(values) == 1:
            resolved[field] = next(iter(values))
    if not resolved.get("lithium_salt"):
        return False
    chosen = evidence[0]
    changed = []
    for field, value in resolved.items():
        if field.startswith("proton_donor") and record.get("component_reporting_status", {}).get("proton_donor") == NOT_PRESENT_IN_RECIPE:
            continue
        if record.get(field) is None:
            record[field] = value
            changed.append(field)
    if not changed:
        return False
    record["condition_value_origin"] = "INHERITED_SAME_PAPER_SERIES"
    record["condition_inherited_from_asset"] = chosen["asset"]
    record["condition_inherited_from_locator"] = chosen["locator"]
    record["supporting_source_locators"].append(f"{chosen['asset']}#{chosen['locator']}")
    record["review_notes"] = "Inherited fields: " + ", ".join(changed)
    return True


CONDITION_IDENTITY_FIELDS = [
    "lithium_salt", "lithium_salt_concentration_mol_L", "solvent", "cosolvent", "solvent_ratio",
    "proton_donor", "proton_donor_concentration_value", "additive", "additive_concentration_value",
    "water_content_value", "oxygen_condition", "current_density_mA_cm2", "total_charge_C", "duration_h",
    "temperature_C", "pressure_bar", "gas_flow_sccm", "liquid_flow_mL_min", "electrode_material", "replicate_n",
]


def condition_identity(record: dict[str, Any]) -> tuple[Any, ...]:
    owner = record.get("paper_id") or record.get("provisional_bundle_id") or record.get("source_bundle_id")
    return (owner, *(record.get(field) for field in CONDITION_IDENTITY_FIELDS))


def resolve_duplicates(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    groups: dict[tuple[Any, ...], list[dict[str, Any]]] = {}
    for row in records:
        if row.get("fe_nh3_percent") is None:
            continue
        groups.setdefault((*condition_identity(row), round(float(row["fe_nh3_percent"]), 6)), []).append(row)
    rank = {"source_data": 0, "primary_structured_author_experiment": 1, "SI prose": 2, "primary_text_candidate": 3, "legacy_candidate_evidence_not_gold_training": 4}
    output = []
    for same in groups.values():
        if len(same) < 2:
            continue
        canonical = min(same, key=lambda row: (not bool(row.get("model_eligible_fe")), rank.get(str(row.get("evidence_authority")), 9), str(row.get("record_id"))))
        support = canonical.setdefault("supporting_source_locators", [])
        for row in same:
            if row is canonical:
                continue
            row["duplicate_classification"] = "TRUE_DUPLICATE"
            row["model_eligible_fe"] = False
            row["review_status"] = "MERGED_DUPLICATE"
            support.append(f"{row.get('source_filename')}#{row.get('source_locator')}")
            output.append({"record_id": row.get("record_id"), "canonical_record_id": canonical.get("record_id"), "classification": "TRUE_DUPLICATE", "supporting_locator": row.get("source_locator")})
        canonical["duplicate_classification"] = "CANONICAL"
    return output


def classify_source_conflicts(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    by_owner: dict[object, list[dict[str, Any]]] = {}
    for row in records:
        owner = row.get("paper_id") or row.get("provisional_bundle_id") or row.get("source_bundle_id")
        by_owner.setdefault(owner, []).append(row)
    output = []
    for owner, rows in by_owner.items():
        flagged = [row for row in rows if "SOURCE_CONFLICT" in (row.get("ambiguity_flags") or [])]
        for row in flagged:
            same_condition = [other for other in rows if other is not row and condition_identity(other) == condition_identity(row)]
            if not same_condition:
                classification = "RESOLVED_DIFFERENT_CONDITIONS"
                row["ambiguity_flags"] = [flag for flag in row["ambiguity_flags"] if flag != "SOURCE_CONFLICT"]
            elif any(abs(float(other.get("fe_nh3_percent") or 0) - float(row.get("fe_nh3_percent") or 0)) <= 0.15 for other in same_condition):
                classification = "RESOLVED_ROUNDING_DIFFERENCE"
            elif any(other.get("duplicate_classification") == "TRUE_DUPLICATE" for other in same_condition):
                classification = "RESOLVED_MAIN_SI_DUPLICATE"
            else:
                classification = "TRUE_UNRESOLVED_CONFLICT"
            row["source_conflict_classification"] = classification
            output.append({"record_id": row.get("record_id"), "paper_group": owner, "classification": classification, "source_locator": row.get("source_locator")})
    return output


def component_reporting(record: dict[str, Any]) -> dict[str, str]:
    status = dict(record.get("component_reporting_status") or {})
    for field in ("lithium_salt", "solvent", "cosolvent", "proton_donor", "additive", "water_content_value"):
        if status.get(field) == NOT_PRESENT_IN_RECIPE:
            continue
        if record.get(field) is not None:
            status[field] = "REPORTED"
        elif record.get("electrolyte_notes") and (re.search(rf"\bno\s+{field}\b|without\s+{field}", str(record["electrolyte_notes"]), re.I) or (field == "proton_donor" and re.search(r"without\s+(?:EtOH|ethanol|proton donor)", str(record["electrolyte_notes"]), re.I))):
            status[field] = NOT_PRESENT_IN_RECIPE
        else:
            status[field] = NOT_REPORTED
    return status


def recalculate_v01(record: dict[str, Any]) -> dict[str, Any]:
    record["schema_version"] = SCHEMA_VERSION
    record["component_reporting_status"] = component_reporting(record)
    result = apply_admission(record)
    if record.get("human_review_required"):
        record["model_eligible_fe"] = False
        record["eligibility_tier"] = "TIER_C"
        record["review_status"] = "HUMAN_REVIEW_REQUIRED"
        result["model_eligible_fe"] = False
        result["eligibility_tier"] = "TIER_C"
    return result


def rescue_priority(record: dict[str, Any]) -> str:
    failures = _json_list(record.get("missing_required_fields"))
    flags = set(_json_list(record.get("ambiguity_flags")))
    if "TARGET_OWNERSHIP_UNCLEAR" in flags:
        return "P4_OWNERSHIP_REVIEW"
    if "SOURCE_CONFLICT" in flags:
        return "P2_SOURCE_CONFLICT"
    if "POSSIBLE_DUPLICATE_ROW" in flags:
        return "P3_DUPLICATE_RESOLUTION"
    if record.get("source_group") == "legacy_oa":
        return "P5_LOW_VALUE"
    if len(failures) == 1:
        return "P0_ONE_FIELD_MISSING"
    if len(failures) == 2:
        return "P1_TWO_FIELDS_MISSING"
    return "P5_LOW_VALUE"


def build_rescue_queue(tier_c: list[dict[str, Any]], manifests: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
    output = []
    for record in tier_c:
        manifest = manifests.get(str(record.get("source_bundle_id")), {})
        identity = manifest.get("identity") or {}
        output.append({
            "record_id": record.get("record_id"), "paper_id": record.get("paper_id"), "bundle_id": record.get("source_bundle_id"),
            "paper_title": identity.get("title"), "doi": identity.get("doi"), "source_group": record.get("source_group"),
            "source_asset": record.get("source_filename"), "source_locator": record.get("source_locator"), "source_page": record.get("source_page"),
            "source_table": record.get("source_table"), "fe_nh3_percent": record.get("fe_nh3_percent"),
            "current_admission_failures": record.get("missing_required_fields"), "lithium_salt": record.get("lithium_salt"),
            "salt_concentration": record.get("lithium_salt_concentration_mol_L"), "solvent": record.get("solvent"), "cosolvent": record.get("cosolvent"),
            "proton_donor": record.get("proton_donor"), "proton_donor_concentration": record.get("proton_donor_concentration_value"),
            "additive": record.get("additive"), "additive_concentration": record.get("additive_concentration_value"),
            "water": record.get("water_content_value"), "oxygen": record.get("oxygen_condition"), "current_density": record.get("current_density_mA_cm2"),
            "charge": record.get("total_charge_C"), "duration": record.get("duration_h"), "temperature": record.get("temperature_C"),
            "pressure": record.get("pressure_bar"), "gas_flow": record.get("gas_flow_sccm"), "liquid_flow": record.get("liquid_flow_mL_min"),
            "current_ambiguity_flags": record.get("ambiguity_flags"), "rescue_status": "HUMAN_REVIEW_REQUIRED",
            "rescued_from_asset": None, "rescued_from_locator": None, "rescued_fields": [],
            "remaining_failures": record.get("missing_required_fields"), "human_review_required": True,
            "review_notes": "No deterministic automatic promotion was applied to this v0 Tier-C candidate.",
        })
    return output


def review_packet_markdown(paper: dict[str, Any], candidates: list[dict[str, Any]]) -> str:
    identity = paper.get("identity") or {}
    lines = [f"# Review packet: {identity.get('title') or candidates[0].get('paper_id') or paper.get('bundle_id')}", "", f"- DOI: {identity.get('doi') or ''}", f"- Paper ID: {candidates[0].get('paper_id') or ''}", ""]
    for record in candidates:
        flags = set(_json_list(record.get("ambiguity_flags")))
        suggestion = "SOURCE_CONFLICT" if "SOURCE_CONFLICT" in flags else ("MERGE_DUPLICATE" if "POSSIBLE_DUPLICATE_ROW" in flags else ("KEEP_TIER_C" if record.get("training_role") == "target_label_candidate" else "REJECT"))
        lines.extend([
            f"## Candidate {record.get('record_id')}", "", f"- Admission failures: {', '.join(_json_list(record.get('missing_required_fields')))}",
            f"- Locator: `{record.get('source_filename')}#{record.get('source_locator')}`", f"- Suggested resolution: **{suggestion}**", "",
            "### Exact source excerpt", "", str(record.get("source_text_excerpt") or "[No excerpt available in v0 candidate]"), "",
            "### Current parsed values", "", "```json", json.dumps({field: record.get(field) for field in ("fe_nh3_percent", "lithium_salt", "lithium_salt_concentration_mol_L", "solvent", "cosolvent", "proton_donor", "proton_donor_concentration_value", "current_density_mA_cm2", "total_charge_C", "duration_h", "temperature_C", "pressure_bar", "gas_flow_sccm", "liquid_flow_mL_min")}, indent=2, ensure_ascii=False), "```", "",
            "HUMAN_DECISION:", "", "HUMAN_NOTES:", "",
        ])
    return "\n".join(lines)


def write_review_packets(tier_c: list[dict[str, Any]], manifests: dict[str, dict[str, Any]], output_dir: str | Path) -> list[Path]:
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    by_bundle: dict[str, list[dict[str, Any]]] = {}
    for row in tier_c:
        by_bundle.setdefault(str(row.get("source_bundle_id")), []).append(row)
    paths = []
    for bundle, rows in sorted(by_bundle.items()):
        paper = manifests.get(bundle, {"bundle_id": bundle})
        name = str(rows[0].get("paper_id") or bundle).replace("/", "_") + ".md"
        path = output_dir / name
        path.write_text(review_packet_markdown(paper, rows), encoding="utf-8")
        paths.append(path)
    return paths


def validate_no_cross_paper_inheritance(records: Iterable[dict[str, Any]], manifests: dict[str, dict[str, Any]]) -> None:
    for row in records:
        asset = row.get("condition_inherited_from_asset")
        if not asset:
            continue
        manifest = manifests.get(str(row.get("source_bundle_id")))
        if not manifest or asset not in {item.get("relative_path") for item in manifest.get("assets", [])}:
            raise AssertionError(f"Cross-paper or missing inheritance asset for {row.get('record_id')}: {asset}")


def copy_for_v01(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    output = deepcopy(records)
    for row in output:
        row["schema_version"] = SCHEMA_VERSION
        for field in PROVENANCE_FIELDS:
            row.setdefault(field, [] if field in {"supporting_source_locators"} else ({} if field in {"original_headers", "original_units", "component_reporting_status"} else None))
    return output
