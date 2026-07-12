#!/usr/bin/env python3
"""Sequentially download legal OA PDFs and write results back to the master XLSX.

Only public URLs exposed by the workbook, DOI resolution, Unpaywall, OpenAlex, or
an OA landing page are considered.  No proxy, credentials, browser cookies, or
access-control bypass is used.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import shutil
import sys
import time
from collections import Counter, defaultdict
from copy import copy
from dataclasses import dataclass, field
from datetime import datetime, timezone
from difflib import SequenceMatcher
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence
from urllib.parse import quote, urljoin, urlparse

import requests
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


PAPERS_WRITE_FIELDS = [
    "OA_Verified", "OA_Status", "Download_Status", "Download_Source",
    "Resolved_PDF_URL", "Landing_Page_URL", "PDF_Version", "OA_License",
    "Host_Type", "HTTP_Status", "PDF_Validation_Status", "PDF_File_Name",
    "Local_PDF_Path", "PDF_Size_Bytes", "PDF_SHA256", "Downloaded_At_UTC",
    "Download_Attempts", "Download_Error", "Needs_Manual_Followup",
    "Last_Checked_At_UTC",
]
MANIFEST_FIELDS = [
    "sequence", "paper_id", "year", "title", "journal", "doi",
    "workbook_access_model", "workbook_institutional_access", "oa_verified",
    "oa_status", "status", "chosen_source", "chosen_url", "landing_url",
    "host_type", "version", "license", "http_status", "validation_status",
    "file_name", "relative_path", "file_size_bytes", "sha256",
    "download_attempts", "attempted_urls", "needs_manual_followup", "error",
    "timestamp_utc",
]
RAG_FIELDS = [
    "Paper_ID", "DOI", "Title", "Year", "Journal", "Reaction_Family",
    "PDF_File_Name", "Local_PDF_Path", "PDF_Exists",
    "PDF_Validation_Status", "PDF_Size_Bytes", "PDF_SHA256", "OA_Verified",
    "OA_License", "PDF_Version", "Download_Source", "Download_Status",
    "Ready_For_RAG", "RAG_Exclusion_Reason", "Last_Checked_At_UTC",
]
FOLLOWUP_FIELDS = [
    "Priority", "Paper_ID", "DOI", "Title", "Journal", "Year",
    "OA_Verified", "OA_Status", "Download_Status", "Publisher_URL",
    "Resolved_PDF_URL", "Download_Error", "Suggested_Action",
    "Last_Checked_At_UTC",
]
FAILURE_STATUSES = {
    "missing_doi", "no_pdf_url", "oa_metadata_error", "download_http_error",
    "invalid_pdf", "supplementary_rejected", "manual_followup_required",
    "interrupted",
}
FINAL_STATUSES = {"downloaded", "skipped_existing", "not_oa"}
MIN_PDF_BYTES = 10 * 1024
TITLE_FILL = "17365D"
HEADER_FILL = "1F4E78"
SECTION_FILL = "D9EAF7"
STATUS_FILLS = {
    "downloaded": "C6EFCE", "skipped_existing": "C6EFCE", "not_oa": "D9E1F2",
    "manual_followup_required": "FFF2CC", "no_pdf_url": "FFF2CC",
    "pending": "DDEBF7", "download_http_error": "F4CCCC", "invalid_pdf": "F4CCCC",
    "oa_metadata_error": "F4CCCC", "supplementary_rejected": "F4CCCC",
}


@dataclass(frozen=True)
class Candidate:
    url: str
    source: str
    landing_url: str = ""
    host_type: str = ""
    version: str = ""
    license: str = ""
    order: int = 0


@dataclass
class Resolution:
    oa_verified: bool | None = None
    oa_status: str = "unknown"
    candidates: list[Candidate] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    metadata_titles: list[str] = field(default_factory=list)
    unpaywall_is_oa: bool | None = None


@dataclass
class Result:
    status: str = "pending"
    oa_verified: bool | None = None
    oa_status: str = "unknown"
    chosen_source: str = ""
    chosen_url: str = ""
    landing_url: str = ""
    host_type: str = ""
    version: str = ""
    license: str = ""
    http_status: int | str = ""
    validation_status: str = ""
    file_name: str = ""
    relative_path: str = ""
    file_size_bytes: int | str = ""
    sha256: str = ""
    downloaded_at_utc: str = ""
    attempts: int = 0
    attempted_urls: list[str] = field(default_factory=list)
    error: str = ""
    needs_followup: bool = False


class PDFMetaParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.urls: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        values = {str(k).lower(): (v or "") for k, v in attrs}
        tag = tag.lower()
        if tag == "link" and values.get("type", "").lower().split(";", 1)[0] == "application/pdf":
            if values.get("href", "").strip():
                self.urls.append(values["href"].strip())
        if tag == "meta":
            name = (values.get("name") or values.get("property") or "").lower()
            if name == "citation_pdf_url" and values.get("content", "").strip():
                self.urls.append(values["content"].strip())


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def paper_id_number(value: Any) -> int:
    match = re.fullmatch(r"P0*(\d+)", str(value or "").strip(), re.IGNORECASE)
    if not match:
        raise ValueError(f"invalid Paper_ID: {value!r}")
    return int(match.group(1))


def normalize_doi(value: Any) -> str:
    doi = str(value or "").strip()
    doi = re.sub(r"^(?:doi\s*:\s*|https?://(?:dx\.)?doi\.org/)", "", doi, flags=re.I)
    return re.sub(r"[\s.,;:]+$", "", doi.strip().strip("<>[]{}" )).lower()


def is_supplementary_reference(value: str) -> bool:
    text = str(value or "").lower()
    path = urlparse(text).path if "://" in text else text
    base = Path(path).name
    explicit = r"(?:supplement(?:ary)?|supporting[-_ ]?(?:information|info|file)?|suppinfo|suppl|electronic[-_ ]?supplement|[._-]esi[._-])"
    suffix = r"(?:[._/-](?:s|si)\d{1,4})(?:\.[a-z0-9]+)?$"
    return bool(re.search(explicit, path, re.I) or re.search(suffix, base, re.I))


def sanitize_component(value: Any, max_length: int = 90) -> str:
    text = re.sub(r"[<>:\"/\\|?*\x00-\x1f]", " ", str(value or ""))
    text = re.sub(r"[^\w. -]+", " ", text, flags=re.UNICODE)
    text = re.sub(r"[\s_]+", "_", text).strip(" ._")
    return (text or "Unknown")[:max_length].rstrip(" ._")


def first_author(authors: Any) -> str:
    first = str(authors or "").split(";", 1)[0].strip()
    surname = first.split(",", 1)[0] if "," in first else (first.split()[-1] if first else "Unknown")
    return sanitize_component(surname, 40)


def generated_filename(record: Mapping[str, Any]) -> str:
    year_match = re.search(r"\d{4}", str(record.get("Year") or ""))
    year = year_match.group(0) if year_match else "UnknownYear"
    stem = f"{record.get('Paper_ID')}_{first_author(record.get('Authors'))}_{year}_{sanitize_component(record.get('Title'), 105)}"
    return sanitize_component(stem, 176) + ".pdf"


def allocate_filename(record: Mapping[str, Any], used: set[str]) -> str:
    existing = str(record.get("PDF_File_Name") or "").strip()
    name = sanitize_component(Path(existing).stem, 176) + ".pdf" if existing else generated_filename(record)
    key = name.casefold()
    if key in used:
        token = hashlib.sha1((normalize_doi(record.get("DOI")) or str(record.get("Paper_ID"))).encode()).hexdigest()[:8]
        name = sanitize_component(Path(name).stem, 167) + f"_{token}.pdf"
        key = name.casefold()
    used.add(key)
    return name


def locate_workbook(input_dir: Path, explicit: Path | None = None) -> Path:
    if explicit:
        path = explicit if explicit.is_absolute() else Path.cwd() / explicit
        if not path.is_file():
            raise FileNotFoundError(path)
        return path.resolve()
    preferred = input_dir / "eNRR_Master_Database.xlsx"
    if preferred.is_file():
        return preferred.resolve()
    excluded = ("backup", "bak", "temporary", "downloaded_copy")
    choices = [p for p in input_dir.glob("*.xlsx") if not p.name.startswith("~$") and not any(x in p.name.lower() for x in excluded)]
    if len(choices) == 1:
        return choices[0].resolve()
    if not choices:
        raise FileNotFoundError(f"no valid .xlsx workbook in {input_dir}")
    raise ValueError("multiple workbooks found; use --workbook: " + ", ".join(p.name for p in choices))


def find_header_row(ws, required: set[str], limit: int = 10) -> int:
    for row in range(1, min(ws.max_row, limit) + 1):
        values = {str(ws.cell(row, col).value or "").strip() for col in range(1, ws.max_column + 1)}
        if required.issubset(values):
            return row
    raise ValueError(f"header containing {sorted(required)} not found in first {limit} rows of {ws.title}")


def header_map(ws, row: int) -> dict[str, int]:
    return {str(ws.cell(row, col).value).strip(): col for col in range(1, ws.max_column + 1) if ws.cell(row, col).value is not None}


def extend_title_merge(ws, last_col: int) -> None:
    for merged in list(ws.merged_cells.ranges):
        if merged.min_row == merged.max_row == 1 and merged.min_col == 1:
            ws.unmerge_cells(str(merged))
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
            return


def ensure_columns(ws, header_row: int, fields: Sequence[str]) -> dict[str, int]:
    headers = header_map(ws, header_row)
    template_col = max(headers.values())
    for field_name in fields:
        if field_name in headers:
            continue
        col = ws.max_column + 1
        source = ws.cell(header_row, template_col)
        target = ws.cell(header_row, col, field_name)
        target._style = copy(source._style)
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        ws.column_dimensions[get_column_letter(col)].width = min(max(len(field_name) + 2, 14), 30)
        headers[field_name] = col
    extend_title_merge(ws, ws.max_column)
    return headers


def records_from_sheet(ws, header_row: int) -> list[dict[str, Any]]:
    headers = header_map(ws, header_row)
    records: list[dict[str, Any]] = []
    for row in range(header_row + 1, ws.max_row + 1):
        paper_id = ws.cell(row, headers["Paper_ID"]).value
        if not str(paper_id or "").strip():
            continue
        record = {name: ws.cell(row, col).value for name, col in headers.items()}
        record["_row"] = row
        records.append(record)
    records.sort(key=lambda item: paper_id_number(item["Paper_ID"]))
    return records


def selected_records(records: Sequence[Mapping[str, Any]], args: argparse.Namespace) -> list[Mapping[str, Any]]:
    start = paper_id_number(args.start_id) if args.start_id else None
    end = paper_id_number(args.end_id) if args.end_id else None
    output = []
    for record in records:
        number = paper_id_number(record["Paper_ID"])
        if start is not None and number < start:
            continue
        if end is not None and number > end:
            continue
        output.append(record)
        if args.max_records and len(output) >= args.max_records:
            break
    return output


def _valid_url(value: Any) -> str:
    text = str(value or "").strip()
    return text if text.lower().startswith(("http://", "https://")) else ""


def _candidate(location: Mapping[str, Any], source: str, order: int) -> Candidate | None:
    url = _valid_url(location.get("url_for_pdf") or location.get("pdf_url"))
    if not url or is_supplementary_reference(url):
        return None
    source_info = location.get("source") if isinstance(location.get("source"), Mapping) else {}
    return Candidate(url, source, _valid_url(location.get("url_for_landing_page") or location.get("landing_page_url")),
                     str(location.get("host_type") or source_info.get("type") or ""),
                     str(location.get("version") or ""), str(location.get("license") or ""), order)


def _payload_locations(payload: Mapping[str, Any], source: str) -> tuple[list[Candidate], list[str]]:
    locations = payload.get("oa_locations") if source == "unpaywall" else payload.get("locations")
    best = payload.get("best_oa_location")
    sequence = ([best] if isinstance(best, Mapping) else []) + (locations if isinstance(locations, list) else [])
    candidates, landings = [], []
    for item in sequence:
        if not isinstance(item, Mapping):
            continue
        candidate = _candidate(item, source, len(candidates))
        if candidate:
            candidates.append(candidate)
        landing = _valid_url(item.get("url_for_landing_page") or item.get("landing_page_url"))
        if landing:
            landings.append(landing)
    return candidates, landings


def landing_pdf_candidates(session: requests.Session, landing: str, source: str, timeout: float) -> list[Candidate]:
    response = None
    try:
        response = session.get(landing, timeout=timeout, stream=True, allow_redirects=True,
                               headers={"Accept": "text/html,application/xhtml+xml,application/pdf"})
        response.raise_for_status()
        final_url = response.url or landing
        content_type = response.headers.get("Content-Type", "").lower()
        if "application/pdf" in content_type and not is_supplementary_reference(final_url):
            return [Candidate(final_url, source, landing, order=0)]
        if "html" not in content_type:
            return []
        chunks, size = [], 0
        for chunk in response.iter_content(64 * 1024):
            if not chunk:
                continue
            remaining = 2_000_000 - size
            chunks.append(chunk[:remaining])
            size += min(len(chunk), remaining)
            if size >= 2_000_000:
                break
        parser = PDFMetaParser()
        parser.feed(b"".join(chunks).decode(response.encoding or "utf-8", errors="replace"))
        return [Candidate(urljoin(final_url, url), source, final_url, order=i) for i, url in enumerate(parser.urls)
                if not is_supplementary_reference(urljoin(final_url, url))]
    except (requests.RequestException, UnicodeError):
        return []
    finally:
        if response is not None:
            response.close()


def resolve_candidates(session: requests.Session, record: Mapping[str, Any], email: str,
                       api_key: str, timeout: float) -> Resolution:
    resolution = Resolution()
    doi = normalize_doi(record.get("DOI"))
    landings: list[str] = []
    direct: list[Candidate] = []
    for name, value in record.items():
        url = _valid_url(value)
        if not url or name.startswith("_"):
            continue
        if is_supplementary_reference(url):
            continue
        if ".pdf" in urlparse(url).path.lower():
            direct.append(Candidate(url, "excel_direct", host_type="publisher", version="publishedVersion", order=len(direct)))
        elif name == "Publisher_URL":
            landings.append(url)
    resolution.candidates.extend(direct)
    service_answers: list[bool] = []
    oa_landings: list[str] = []
    if doi:
        try:
            response = session.get(f"https://api.unpaywall.org/v2/{quote(doi, safe='')}", params={"email": email}, timeout=timeout)
            if response.status_code == 404:
                resolution.errors.append("Unpaywall: DOI not found")
            else:
                response.raise_for_status()
                payload = response.json()
                is_oa = bool(payload.get("is_oa"))
                resolution.unpaywall_is_oa = is_oa
                service_answers.append(is_oa)
                resolution.oa_status = str(payload.get("oa_status") or "unknown")
                if payload.get("title"):
                    resolution.metadata_titles.append(str(payload["title"]))
                if is_oa:
                    found, found_landings = _payload_locations(payload, "unpaywall")
                    resolution.candidates.extend(found)
                    oa_landings.extend(found_landings)
        except (requests.RequestException, ValueError, json.JSONDecodeError) as exc:
            resolution.errors.append(f"Unpaywall: {exc}")
        try:
            params = {"mailto": email}
            if api_key:
                params["api_key"] = api_key
            response = session.get(f"https://api.openalex.org/works/https://doi.org/{quote(doi, safe='')}", params=params, timeout=timeout)
            if response.status_code != 404:
                response.raise_for_status()
                payload = response.json()
                oa = payload.get("open_access") if isinstance(payload.get("open_access"), Mapping) else {}
                is_oa = bool(oa.get("is_oa"))
                service_answers.append(is_oa)
                if resolution.oa_status == "unknown":
                    resolution.oa_status = str(oa.get("oa_status") or "unknown")
                if payload.get("title"):
                    resolution.metadata_titles.append(str(payload["title"]))
                if is_oa:
                    found, found_landings = _payload_locations(payload, "openalex")
                    resolution.candidates.extend(found)
                    oa_landings.extend(found_landings)
            else:
                resolution.errors.append("OpenAlex: DOI not found")
        except (requests.RequestException, ValueError, json.JSONDecodeError) as exc:
            resolution.errors.append(f"OpenAlex: {exc}")
    resolution.oa_verified = any(service_answers) if service_answers else (True if direct and not doi else None)
    if resolution.unpaywall_is_oa is False:
        resolution.oa_verified = False
    if resolution.oa_verified and not resolution.candidates:
        for landing in dict.fromkeys(oa_landings + landings):
            resolution.candidates.extend(landing_pdf_candidates(session, landing, "landing_metadata", timeout))
    if doi and resolution.oa_verified and not resolution.candidates:
        resolution.candidates.extend(landing_pdf_candidates(session, f"https://doi.org/{quote(doi, safe='/')}", "doi_landing", timeout))
    return resolution


def title_conflict(workbook_title: Any, metadata_titles: Sequence[str]) -> bool:
    left = re.sub(r"\W+", " ", str(workbook_title or "").lower()).strip()
    if not left or not metadata_titles:
        return False
    scores = [SequenceMatcher(None, left, re.sub(r"\W+", " ", title.lower()).strip()).ratio() for title in metadata_titles]
    return max(scores, default=1.0) < 0.45


def candidate_key(candidate: Candidate) -> tuple[int, int, int, int]:
    host = candidate.host_type.lower()
    version = candidate.version.lower()
    if host == "publisher" and version == "publishedversion": rank = 0
    elif host == "publisher" and version == "acceptedversion": rank = 1
    elif host == "repository" and version == "publishedversion": rank = 2
    elif host == "repository" and version == "acceptedversion": rank = 3
    elif host == "repository" and version == "submittedversion": rank = 4
    else: rank = 5
    source_rank = {"excel_direct": 0, "doi_landing": 1, "unpaywall": 2, "openalex": 3, "landing_metadata": 4}.get(candidate.source, 5)
    return rank, source_rank, candidate.order, len(candidate.url)


def ordered_candidates(candidates: Iterable[Candidate]) -> list[Candidate]:
    unique, seen = [], set()
    for candidate in candidates:
        if candidate.url not in seen and not is_supplementary_reference(candidate.url):
            unique.append(candidate); seen.add(candidate.url)
    return sorted(unique, key=candidate_key)


def inspect_pdf(path: Path, min_bytes: int = MIN_PDF_BYTES, max_bytes: int | None = None) -> tuple[bool, str, int, str]:
    try:
        size = path.stat().st_size
        if max_bytes is not None and size > max_bytes:
            return False, "file_too_large", size, ""
        with path.open("rb") as handle:
            prefix = handle.read(512)
        lower = prefix.lstrip().lower()
        if lower.startswith((b"<!doctype html", b"<html", b"<form")):
            return False, "html_instead_of_pdf", size, ""
        if lower.startswith((b"<?xml", b"<error", b"{")):
            return False, "invalid_signature", size, ""
        if not prefix.startswith(b"%PDF-"):
            return False, "invalid_signature", size, ""
        if size < min_bytes:
            return False, "file_too_small", size, ""
        hasher = hashlib.sha256()
        with path.open("rb") as handle:
            for chunk in iter(lambda: handle.read(1024 * 1024), b""):
                hasher.update(chunk)
        digest = hasher.hexdigest()
        return True, "valid_pdf", size, digest
    except OSError:
        return False, "incomplete_download", 0, ""


def download_pdf(session: requests.Session, candidate: Candidate, destination: Path,
                 timeout: float, max_file_mb: float, retries: int = 3) -> tuple[bool, str, int, str, int, int, str]:
    part = destination.with_suffix(destination.suffix + ".part")
    max_bytes = int(max_file_mb * 1024 * 1024)
    last_error, last_status, validation = "", "", "incomplete_download"
    for attempt in range(1, retries + 1):
        try:
            with session.get(candidate.url, timeout=timeout, stream=True, allow_redirects=True,
                             headers={"Accept": "application/pdf"}) as response:
                last_status = response.status_code
                response.raise_for_status()
                if response.status_code != 200:
                    raise requests.HTTPError(f"expected HTTP 200, got {response.status_code}")
                if is_supplementary_reference(response.url or candidate.url):
                    return False, "supplementary_detected", 0, "", last_status, attempt, "redirected to supplementary material"
                declared = int(response.headers.get("Content-Length") or 0)
                if declared > max_bytes:
                    return False, "file_too_large", declared, "", last_status, attempt, "declared file too large"
                size = 0
                with part.open("wb") as handle:
                    for chunk in response.iter_content(64 * 1024):
                        if not chunk: continue
                        size += len(chunk)
                        if size > max_bytes:
                            raise ValueError("file exceeds --max-file-mb")
                        handle.write(chunk)
            if declared and size != declared:
                part.unlink(missing_ok=True)
                validation = "incomplete_download"
                last_error = f"Content-Length {declared}, received {size}"
                if attempt < retries:
                    time.sleep(2 ** (attempt - 1))
                    continue
                return False, validation, size, "", last_status, attempt, last_error
            valid, validation, size, digest = inspect_pdf(part, MIN_PDF_BYTES, max_bytes)
            if not valid:
                last_error = validation
                part.unlink(missing_ok=True)
                if validation in {"html_instead_of_pdf", "invalid_signature", "file_too_small"}:
                    return False, validation, size, "", last_status, attempt, last_error
            else:
                os.replace(part, destination)
                return True, "valid_pdf", size, digest, last_status, attempt, ""
        except (requests.RequestException, OSError, ValueError) as exc:
            last_error = str(exc); part.unlink(missing_ok=True)
            if attempt < retries:
                time.sleep(2 ** (attempt - 1))
    return False, validation, 0, "", last_status, retries, last_error


def local_path(record: Mapping[str, Any], repo_root: Path, input_dir: Path, filename: str) -> Path:
    value = str(record.get("Local_PDF_Path") or "").strip()
    if value:
        path = Path(value)
        return path if path.is_absolute() else repo_root / path
    existing_name = str(record.get("PDF_File_Name") or "").strip()
    return input_dir / (Path(existing_name).name if existing_name else filename)


def process_record(record: Mapping[str, Any], args: argparse.Namespace, session: requests.Session,
                   repo_root: Path, used_names: set[str]) -> Result:
    resume_matches = []
    if not str(record.get("PDF_File_Name") or "").strip() and not str(record.get("Local_PDF_Path") or "").strip():
        prefix = f"{record.get('Paper_ID')}_"
        resume_matches = [path for path in args.input_dir.glob(prefix + "*.pdf") if inspect_pdf(path)[0]]
    filename = resume_matches[0].name if len(resume_matches) == 1 else allocate_filename(record, used_names)
    used_names.add(filename.casefold())
    result = Result(file_name=filename)
    destination = local_path(record, repo_root, args.input_dir, filename)
    if destination.parent.resolve() != args.input_dir.resolve():
        destination = args.input_dir / filename
    if destination.exists() and not args.overwrite:
        valid, validation, size, digest = inspect_pdf(destination)
        if valid:
            result.status = "skipped_existing"; result.validation_status = "existing_valid_pdf"
            result.file_name = destination.name; result.relative_path = f"input_raw/{destination.name}"
            result.file_size_bytes = size; result.sha256 = digest; result.needs_followup = False
            return result
        result.validation_status = "existing_invalid_pdf"
    doi = normalize_doi(record.get("DOI"))
    explicit_urls = [_valid_url(v) for k, v in record.items() if not str(k).startswith("_") and _valid_url(v)]
    if is_supplementary_reference(doi) or (explicit_urls and all(is_supplementary_reference(url) for url in explicit_urls)):
        result.status = "supplementary_rejected"; result.validation_status = "supplementary_detected"
        result.error = "only supplementary-material reference found"; result.needs_followup = True; return result
    direct_pdf = any(".pdf" in urlparse(url).path.lower() and not is_supplementary_reference(url) for url in explicit_urls)
    if not doi and not direct_pdf:
        result.status = "missing_doi"; result.error = "DOI missing and no explicit public PDF URL"
        result.needs_followup = True; return result
    resolution = resolve_candidates(session, record, args.unpaywall_email, args.openalex_api_key, args.timeout)
    result.oa_verified = resolution.oa_verified; result.oa_status = resolution.oa_status
    if title_conflict(record.get("Title"), resolution.metadata_titles):
        result.status = "manual_followup_required"; result.error = "DOI/title metadata conflict"
        result.needs_followup = True; return result
    if resolution.unpaywall_is_oa is False:
        result.status = "not_oa"; result.needs_followup = False; return result
    if resolution.oa_verified is None and not resolution.candidates:
        result.status = "oa_metadata_error"; result.error = "; ".join(resolution.errors)
        result.needs_followup = True; return result
    candidates = ordered_candidates(resolution.candidates)
    if not candidates:
        result.status = "no_pdf_url"; result.error = "OA verified but no primary-article PDF URL"
        result.needs_followup = True; return result
    result.attempted_urls = [c.url for c in candidates]
    chosen = candidates[0]
    for attr, value in (("chosen_source", chosen.source), ("chosen_url", chosen.url),
                        ("landing_url", chosen.landing_url), ("host_type", chosen.host_type),
                        ("version", chosen.version), ("license", chosen.license)):
        setattr(result, attr, value)
    if args.dry_run:
        result.status = "pending"; result.error = "dry-run: candidate selected"; return result
    invalid_errors, http_errors = [], []
    for candidate in candidates:
        result.chosen_source = candidate.source; result.chosen_url = candidate.url
        result.landing_url = candidate.landing_url; result.host_type = candidate.host_type
        result.version = candidate.version; result.license = candidate.license
        ok, validation, size, digest, http_status, attempts, error = download_pdf(
            session, candidate, destination, args.timeout, args.max_file_mb)
        result.attempts += attempts; result.http_status = http_status; result.validation_status = validation
        if ok:
            result.status = "downloaded"; result.oa_verified = True
            result.file_name = destination.name; result.relative_path = f"input_raw/{destination.name}"
            result.file_size_bytes = size; result.sha256 = digest
            result.downloaded_at_utc = utc_now(); result.needs_followup = False; return result
        if validation in {"html_instead_of_pdf", "invalid_signature", "file_too_small", "file_too_large"}:
            invalid_errors.append(f"{candidate.url}: {validation}")
        else:
            http_errors.append(f"{candidate.url}: {error}")
    result.status = "invalid_pdf" if invalid_errors else "download_http_error"
    result.error = "; ".join(invalid_errors + http_errors); result.needs_followup = True
    return result


def apply_result(ws, headers: Mapping[str, int], row: int, result: Result, workbook_path: Path) -> None:
    values = {
        "OA_Verified": result.oa_verified if result.oa_verified is not None else "",
        "OA_Status": result.oa_status, "Download_Status": result.status,
        "Download_Source": result.chosen_source, "Resolved_PDF_URL": result.chosen_url,
        "Landing_Page_URL": result.landing_url, "PDF_Version": result.version,
        "OA_License": result.license, "Host_Type": result.host_type,
        "HTTP_Status": result.http_status, "PDF_Validation_Status": result.validation_status,
        "PDF_File_Name": result.file_name, "Local_PDF_Path": result.relative_path.replace("/", "\\"),
        "PDF_Size_Bytes": result.file_size_bytes, "PDF_SHA256": result.sha256,
        "Downloaded_At_UTC": result.downloaded_at_utc, "Download_Attempts": result.attempts,
        "Download_Error": result.error, "Needs_Manual_Followup": result.needs_followup,
        "Last_Checked_At_UTC": utc_now(),
    }
    for name, value in values.items():
        ws.cell(row, headers[name], value)
    path_cell = ws.cell(row, headers["Local_PDF_Path"])
    if result.relative_path:
        path_cell.hyperlink = (workbook_path.parent.parent / result.relative_path).resolve().as_uri()
        path_cell.style = "Hyperlink"
    status_cell = ws.cell(row, headers["Download_Status"])
    status_cell.fill = PatternFill("solid", fgColor=STATUS_FILLS.get(result.status, "FFFFFF"))


def atomic_save(workbook: Workbook, path: Path) -> None:
    saving = path.with_name(path.name + ".saving.xlsx")
    workbook.save(saving)
    # Re-open the temporary file before replacement so a bad checkpoint cannot replace the original.
    check = load_workbook(saving, read_only=True, data_only=False, keep_links=True)
    check.close()
    os.replace(saving, path)


def create_backup(path: Path) -> Path:
    backup_dir = path.parent / "excel_backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    backup = backup_dir / f"{path.stem}_before_oa_download_{stamp}.xlsx"
    counter = 1
    while backup.exists():
        backup = backup_dir / f"{path.stem}_before_oa_download_{stamp}_{counter}.xlsx"; counter += 1
    shutil.copy2(path, backup)
    return backup


def result_from_sheet(record: Mapping[str, Any]) -> Result:
    return Result(
        status=str(record.get("Download_Status") or "pending"),
        oa_verified=record.get("OA_Verified") if isinstance(record.get("OA_Verified"), bool) else None,
        oa_status=str(record.get("OA_Status") or "unknown"), chosen_source=str(record.get("Download_Source") or ""),
        chosen_url=str(record.get("Resolved_PDF_URL") or ""), landing_url=str(record.get("Landing_Page_URL") or ""),
        host_type=str(record.get("Host_Type") or ""), version=str(record.get("PDF_Version") or ""),
        license=str(record.get("OA_License") or ""), http_status=record.get("HTTP_Status") or "",
        validation_status=str(record.get("PDF_Validation_Status") or ""), file_name=str(record.get("PDF_File_Name") or ""),
        relative_path=str(record.get("Local_PDF_Path") or "").replace("\\", "/"),
        file_size_bytes=record.get("PDF_Size_Bytes") or "", sha256=str(record.get("PDF_SHA256") or ""),
        downloaded_at_utc=str(record.get("Downloaded_At_UTC") or ""), attempts=int(record.get("Download_Attempts") or 0),
        error=str(record.get("Download_Error") or ""), needs_followup=bool(record.get("Needs_Manual_Followup")),
    )


def rag_exclusion(record: Mapping[str, Any], result: Result, input_dir: Path, metadata_conflict: bool = False) -> str:
    path = input_dir / Path(result.file_name).name if result.file_name else None
    reasons = []
    if not path or not path.is_file(): reasons.append("PDF missing")
    if result.validation_status not in {"valid_pdf", "existing_valid_pdf"}: reasons.append("PDF invalid or unverified")
    if is_supplementary_reference(result.file_name): reasons.append("supplementary file")
    if not normalize_doi(record.get("DOI")): reasons.append("DOI missing")
    if metadata_conflict: reasons.append("duplicate DOI or title conflict")
    return "; ".join(reasons)


def conflicting_paper_ids(records: Sequence[Mapping[str, Any]]) -> set[str]:
    doi_groups: dict[str, list[str]] = defaultdict(list)
    title_groups: dict[str, list[str]] = defaultdict(list)
    for record in records:
        pid = str(record["Paper_ID"])
        doi = normalize_doi(record.get("DOI"))
        title = re.sub(r"\W+", " ", str(record.get("Title") or "").lower()).strip()
        if doi: doi_groups[doi].append(pid)
        if title: title_groups[title].append(pid)
    conflicts: set[str] = set()
    for group in list(doi_groups.values()) + list(title_groups.values()):
        if len(group) > 1: conflicts.update(group)
    return conflicts


def style_table_sheet(ws, header_row: int, last_row: int, last_col: int) -> None:
    for cell in ws[header_row]:
        if cell.column <= last_col:
            cell.fill = PatternFill("solid", fgColor=HEADER_FILL); cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(last_col)}{last_row}"
    ws.freeze_panes = f"A{header_row + 1}"
    ws.sheet_view.showGridLines = False


def format_papers_download_columns(ws, headers: Mapping[str, int]) -> None:
    widths = {
        "OA_Verified": 12, "OA_Status": 14, "Download_Source": 18,
        "Resolved_PDF_URL": 44, "Landing_Page_URL": 40, "PDF_Version": 18,
        "OA_License": 16, "Host_Type": 14, "HTTP_Status": 12,
        "PDF_Validation_Status": 22, "PDF_Size_Bytes": 16, "PDF_SHA256": 68,
        "Downloaded_At_UTC": 25, "Download_Attempts": 18, "Download_Error": 60,
        "Needs_Manual_Followup": 24, "Last_Checked_At_UTC": 25,
    }
    for name, width in widths.items():
        if name in headers:
            ws.column_dimensions[get_column_letter(headers[name])].width = width


def update_rag_manifest(wb: Workbook, records: Sequence[Mapping[str, Any]], results: Mapping[str, Result], input_dir: Path) -> int:
    ws = wb["PDF_RAG_MANIFEST"] if "PDF_RAG_MANIFEST" in wb.sheetnames else wb.create_sheet("PDF_RAG_MANIFEST")
    header_row = find_header_row(ws, {"Paper_ID"}) if ws.max_row else 1
    headers = ensure_columns(ws, header_row, RAG_FIELDS)
    existing = {}
    blank_rows = []
    for row in range(header_row + 1, ws.max_row + 1):
        pid = str(ws.cell(row, headers["Paper_ID"]).value or "").strip()
        if pid and pid not in existing: existing[pid] = row
        elif not pid: blank_rows.append(row)
    next_row = ws.max_row + 1
    used_rows = []
    conflicts = conflicting_paper_ids(records)
    for record in records:
        pid = str(record["Paper_ID"]); row = existing.get(pid)
        if row is None:
            row = blank_rows.pop(0) if blank_rows else next_row
            if row == next_row: next_row += 1
        used_rows.append(row)
        result = results[pid]; exclusion = rag_exclusion(record, result, input_dir, pid in conflicts)
        values = {
            "Paper_ID": pid, "DOI": normalize_doi(record.get("DOI")), "Title": record.get("Title") or "",
            "Year": record.get("Year") or "", "Journal": record.get("Journal") or "",
            "Reaction_Family": record.get("Primary_Domain") or "", "PDF_File_Name": result.file_name,
            "Local_PDF_Path": result.relative_path.replace("/", "\\"), "PDF_Exists": not bool(exclusion.startswith("PDF missing")),
            "PDF_Validation_Status": result.validation_status, "PDF_Size_Bytes": result.file_size_bytes,
            "PDF_SHA256": result.sha256, "OA_Verified": result.oa_verified if result.oa_verified is not None else "",
            "OA_License": result.license, "PDF_Version": result.version, "Download_Source": result.chosen_source,
            "Download_Status": result.status, "Ready_For_RAG": not bool(exclusion),
            "RAG_Exclusion_Reason": exclusion, "Last_Checked_At_UTC": utc_now(),
        }
        for name, value in values.items(): ws.cell(row, headers[name], value)
        if "Document_ID" in headers and not ws.cell(row, headers["Document_ID"]).value:
            ws.cell(row, headers["Document_ID"], f'=IF(A{row}="","","DOC-"&A{row}&"-"&TEXT(COUNTIF($A$3:A{row},A{row}),"00"))')
    style_table_sheet(ws, header_row, max(used_rows), ws.max_column)
    extend_title_merge(ws, ws.max_column)
    return len(records)


def _new_report_sheet(wb: Workbook, name: str):
    if name in wb.sheetnames: wb.remove(wb[name])
    ws = wb.create_sheet(name); ws.sheet_view.showGridLines = False
    return ws


def _section(ws, row: int, title: str, headers: Sequence[str], rows: Sequence[Sequence[Any]]) -> int:
    ws.cell(row, 1, title); ws.cell(row, 1).fill = PatternFill("solid", fgColor=SECTION_FILL)
    ws.cell(row, 1).font = Font(bold=True, color=TITLE_FILL); row += 1
    for col, value in enumerate(headers, 1):
        cell = ws.cell(row, col, value); cell.fill = PatternFill("solid", fgColor=HEADER_FILL); cell.font = Font(color="FFFFFF", bold=True)
    row += 1
    for values in rows:
        for col, value in enumerate(values, 1): ws.cell(row, col, value)
        row += 1
    return row + 1


def create_analysis(wb: Workbook, workbook_path: Path, records: Sequence[Mapping[str, Any]],
                    results: Mapping[str, Result], downloaded_this_run: int, input_dir: Path) -> None:
    ws = _new_report_sheet(wb, "DOWNLOAD_ANALYSIS")
    total = len(records); vals = list(results.values())
    conflicts = conflicting_paper_ids(records)
    valid = [r for r in vals if r.validation_status in {"valid_pdf", "existing_valid_pdf"}]
    ready = sum(not rag_exclusion(rec, results[str(rec["Paper_ID"])], input_dir, str(rec["Paper_ID"]) in conflicts) for rec in records)
    follow = sum(r.needs_followup for r in vals); oa_count = sum(r.oa_verified is True for r in vals)
    failures = sum(r.status in {"download_http_error", "invalid_pdf", "oa_metadata_error"} for r in vals)
    basic = [
        ("Workbook", workbook_path.name), ("Generated_At_UTC", utc_now()), ("Total_PAPERS", total),
        ("Records_With_DOI", sum(bool(normalize_doi(r.get("DOI"))) for r in records)),
        ("Records_Without_DOI", sum(not bool(normalize_doi(r.get("DOI"))) for r in records)),
        ("Records_With_Publisher_URL", sum(bool(_valid_url(r.get("Publisher_URL"))) for r in records)),
        ("Existing_Valid_PDFs", sum(r.validation_status == "existing_valid_pdf" for r in vals)),
        ("Downloaded_This_Run", downloaded_this_run), ("Total_Valid_PDFs", len(valid)),
        ("OA_Verified_Count", oa_count), ("Not_OA_Count", sum(r.status == "not_oa" for r in vals)),
        ("Manual_Followup_Count", follow), ("Invalid_PDF_Count", sum(r.status == "invalid_pdf" for r in vals)),
        ("Download_Failure_Count", failures), ("Ready_For_RAG_Count", ready),
        ("PDF_Coverage_Rate", len(valid) / total if total else 0),
        ("OA_Download_Success_Rate", len(valid) / oa_count if oa_count else 0),
        ("RAG_Readiness_Rate", ready / total if total else 0),
    ]
    row = 1
    row = _section(ws, row, "A. Basic Information", ["Metric", "Value"], basic)
    status_counts = Counter(r.status for r in vals)
    row = _section(ws, row, "B. Download_Status Distribution", ["Download_Status", "Count", "Percentage"],
                   [(k, v, v / total if total else 0) for k, v in sorted(status_counts.items())])
    oa_counts = Counter(r.oa_status for r in vals)
    row = _section(ws, row, "C. OA_Status Distribution", ["OA_Status", "Count", "Percentage"],
                   [(k, v, v / total if total else 0) for k, v in sorted(oa_counts.items())])
    def grouped(field_name):
        groups = defaultdict(list)
        for rec in records: groups[str(rec.get(field_name) or "Missing")].append(rec)
        return groups
    year_rows = []
    for key, group in grouped("Year").items():
        rs = [results[str(x["Paper_ID"])] for x in group]; vp = sum(x.validation_status in {"valid_pdf","existing_valid_pdf"} for x in rs)
        rr = sum(not rag_exclusion(x, results[str(x["Paper_ID"])], input_dir, str(x["Paper_ID"]) in conflicts) for x in group)
        year_rows.append((key, len(group), sum(x.oa_verified is True for x in rs), vp, rr, vp / len(group)))
    row = _section(ws, row, "D. Statistics by Year", ["Year","Total_Papers","OA_Verified","Valid_PDFs","Ready_For_RAG","Coverage_Rate"],
                   sorted(year_rows, key=lambda x: str(x[0]), reverse=True))
    journal_rows = []
    for key, group in grouped("Journal").items():
        rs=[results[str(x["Paper_ID"])] for x in group]; vp=sum(x.validation_status in {"valid_pdf","existing_valid_pdf"} for x in rs)
        rr=sum(not rag_exclusion(x, results[str(x["Paper_ID"])], input_dir, str(x["Paper_ID"]) in conflicts) for x in group)
        journal_rows.append((key,len(group),sum(x.oa_verified is True for x in rs),vp,rr,vp/len(group)))
    journal_rows.sort(key=lambda x:(-x[1],x[0]))
    row = _section(ws,row,"E. Statistics by Journal",["Journal","Total_Papers","OA_Verified","Valid_PDFs","Ready_For_RAG","Coverage_Rate"],journal_rows)
    access_rows=[]
    for key,group in grouped("Access_Model").items():
        rs=[results[str(x["Paper_ID"])] for x in group]
        access_rows.append((key,len(group),sum(x.oa_verified is True for x in rs),sum(x.validation_status in {"valid_pdf","existing_valid_pdf"} for x in rs),sum(x.status=="downloaded" for x in rs),sum(x.status=="not_oa" for x in rs),sum(x.needs_followup for x in rs)))
    row=_section(ws,row,"F. Statistics by Access Model",["Access_Model","Total","OA_Verified","Valid_PDFs","Downloaded","Not_OA","Followup"],access_rows)
    article_rows=[]
    for key,group in grouped("Article_Type").items():
        rs=[results[str(x["Paper_ID"])] for x in group]
        article_rows.append((key,len(group),sum(x.validation_status in {"valid_pdf","existing_valid_pdf"} for x in rs),sum(not rag_exclusion(x,results[str(x["Paper_ID"])],input_dir,str(x["Paper_ID"]) in conflicts) for x in group)))
    row=_section(ws,row,"G. Statistics by Article Type",["Article_Type","Total","Valid_PDFs","Ready_For_RAG"],article_rows)
    ids=[str(r["Paper_ID"]) for r in records]; dois=[normalize_doi(r.get("DOI")) for r in records if normalize_doi(r.get("DOI"))]
    titles=[str(r.get("Title") or "").strip().lower() for r in records if str(r.get("Title") or "").strip()]
    paths=[x.relative_path.casefold() for x in vals if x.relative_path]
    quality = [
        ("Duplicate_Paper_ID",sum(v-1 for v in Counter(ids).values() if v>1)),("Duplicate_DOI",sum(v-1 for v in Counter(dois).values() if v>1)),
        ("Duplicate_Title",sum(v-1 for v in Counter(titles).values() if v>1)),("Missing_Title",sum(not r.get("Title") for r in records)),
        ("Missing_Year",sum(not r.get("Year") for r in records)),("Missing_Journal",sum(not r.get("Journal") for r in records)),
        ("Missing_DOI",sum(not normalize_doi(r.get("DOI")) for r in records)),
        ("Downloaded_But_File_Missing",sum(r.status in {"downloaded","skipped_existing"} and not (input_dir/Path(r.file_name).name).is_file() for r in vals)),
        ("File_Exists_But_Status_Not_Downloaded",
         sum((input_dir/Path(r.file_name).name).is_file() and r.status not in {"downloaded","skipped_existing"} for r in vals if r.file_name)
         + len({p.name for p in input_dir.glob("*.pdf")} - {Path(r.file_name).name for r in vals if r.file_name})),
        ("DOI_File_Name_Conflict",len(conflicts)),("Multiple_Records_Point_To_Same_PDF",sum(v-1 for v in Counter(paths).values() if v>1)),
        ("Invalid_Local_Path",sum(bool(r.relative_path) and not r.relative_path.replace("\\","/").startswith("input_raw/") for r in vals)),
        ("OA_Verified_But_No_PDF",sum(r.oa_verified is True and r.validation_status not in {"valid_pdf","existing_valid_pdf"} for r in vals)),
        ("PDF_Not_Ready_For_RAG",total-ready),
    ]
    row=_section(ws,row,"H. Data Quality Checks",["Check","Count"],quality)
    numbers=sorted(paper_id_number(x) for x in ids if re.fullmatch(r"P0*\d+",x,re.I)); missing=[n for n in range(numbers[0],numbers[-1]+1) if n not in set(numbers)] if numbers else []
    sequence_rows=[("P0001 before P0002",ids.index("P0001")<ids.index("P0002") if all(x in ids for x in ("P0001","P0002")) else False),
                   ("P0009 before P0010",ids.index("P0009")<ids.index("P0010") if all(x in ids for x in ("P0009","P0010")) else False),
                   ("Missing Paper_ID count",len(missing)),("Invalid Paper_ID count",sum(not re.fullmatch(r"P0*\d+",x,re.I) for x in ids)),
                   ("Missing Paper_ID list",", ".join(f"P{n:04d}" for n in missing))]
    row=_section(ws,row,"I. Download Sequence Checks",["Check","Result"],sequence_rows)
    top=[(i+1,*x[:2],x[3],x[5]) for i,x in enumerate(journal_rows[:20])]
    _section(ws,row,"J. Top 20 Journals",["Rank","Journal","Total","Valid_PDFs","Coverage_Rate"],top)
    ws.freeze_panes="A3"; ws.auto_filter.ref=f"A2:G{ws.max_row}"
    for col,width in {"A":34,"B":30,"C":16,"D":16,"E":18,"F":16,"G":16}.items(): ws.column_dimensions[col].width=width
    for row_cells in ws.iter_rows():
        for cell in row_cells:
            if isinstance(cell.value,float): cell.number_format="0.0%"


def suggested_action(result: Result) -> tuple[int, str]:
    if result.status in {"download_http_error","invalid_pdf","manual_followup_required"}:
        return 1, "Replace invalid PDF" if result.status=="invalid_pdf" else "Verify DOI and title" if "conflict" in result.error else "Check publisher landing page manually"
    if result.status == "no_pdf_url": return 2, "Obtain legal repository copy"
    if result.status == "missing_doi": return 3, "Verify DOI and title"
    if result.status == "oa_metadata_error": return 3, "Check whether article is actually OA"
    return 3, "Confirm primary article versus supplementary file"


def create_followup(wb: Workbook, records: Sequence[Mapping[str, Any]], results: Mapping[str, Result]) -> int:
    ws=_new_report_sheet(wb,"DOWNLOAD_FOLLOWUP")
    for col,name in enumerate(FOLLOWUP_FIELDS,1): ws.cell(1,col,name)
    count=0
    for record in records:
        result=results[str(record["Paper_ID"])]
        if not result.needs_followup: continue
        priority,action=suggested_action(result); count+=1
        values=[priority,record["Paper_ID"],normalize_doi(record.get("DOI")),record.get("Title") or "",record.get("Journal") or "",record.get("Year") or "",result.oa_verified if result.oa_verified is not None else "",result.oa_status,result.status,record.get("Publisher_URL") or "",result.chosen_url,result.error,action,utc_now()]
        for col,value in enumerate(values,1): ws.cell(count+1,col,value)
    style_table_sheet(ws,1,max(1,count+1),len(FOLLOWUP_FIELDS))
    widths=[10,12,28,48,25,10,12,14,24,38,42,52,38,24]
    for col,width in enumerate(widths,1): ws.column_dimensions[get_column_letter(col)].width=width
    for row in range(2, count + 2):
        ws.row_dimensions[row].height = 45
        for col in range(1, len(FOLLOWUP_FIELDS) + 1):
            ws.cell(row, col).alignment = Alignment(vertical="top", wrap_text=col in {4, 10, 11, 12, 13})
    return count


def manifest_row(record: Mapping[str, Any], result: Result) -> dict[str, Any]:
    return {
        "sequence": paper_id_number(record["Paper_ID"]), "paper_id": record["Paper_ID"], "year": record.get("Year") or "",
        "title": record.get("Title") or "", "journal": record.get("Journal") or "", "doi": normalize_doi(record.get("DOI")),
        "workbook_access_model": record.get("Access_Model") or "", "workbook_institutional_access": record.get("Institutional_Access") or "",
        "oa_verified": result.oa_verified if result.oa_verified is not None else "", "oa_status": result.oa_status,
        "status": result.status, "chosen_source": result.chosen_source, "chosen_url": result.chosen_url,
        "landing_url": result.landing_url, "host_type": result.host_type, "version": result.version, "license": result.license,
        "http_status": result.http_status, "validation_status": result.validation_status, "file_name": result.file_name,
        "relative_path": result.relative_path, "file_size_bytes": result.file_size_bytes, "sha256": result.sha256,
        "download_attempts": result.attempts, "attempted_urls": result.attempted_urls,
        "needs_manual_followup": result.needs_followup, "error": result.error, "timestamp_utc": utc_now(),
    }


def write_reports(report_dir: Path, records: Sequence[Mapping[str, Any]], results: Mapping[str, Result], summary: Mapping[str, Any]) -> None:
    report_dir.mkdir(parents=True,exist_ok=True)
    rows=[manifest_row(r,results[str(r["Paper_ID"])]) for r in records]
    csv_path=report_dir/"oa_download_manifest.csv"; jsonl_path=report_dir/"oa_download_manifest.jsonl"
    with csv_path.open("w",newline="",encoding="utf-8-sig") as handle:
        writer=csv.DictWriter(handle,fieldnames=MANIFEST_FIELDS); writer.writeheader()
        for row in rows:
            out=dict(row); out["attempted_urls"]=json.dumps(out["attempted_urls"],ensure_ascii=False); writer.writerow(out)
    with jsonl_path.open("w",encoding="utf-8") as handle:
        for row in rows: handle.write(json.dumps(row,ensure_ascii=False)+"\n")
    with (report_dir/"oa_download_run_summary.json").open("w",encoding="utf-8") as handle: json.dump(dict(summary),handle,ensure_ascii=False,indent=2)
    with (report_dir/"oa_download_errors.csv").open("w",newline="",encoding="utf-8-sig") as handle:
        writer=csv.DictWriter(handle,fieldnames=MANIFEST_FIELDS); writer.writeheader()
        for row in rows:
            if row["error"] or row["needs_manual_followup"]:
                out=dict(row); out["attempted_urls"]=json.dumps(out["attempted_urls"],ensure_ascii=False); writer.writerow(out)


def should_process(record: Mapping[str, Any], args: argparse.Namespace, input_dir: Path) -> bool:
    status=str(record.get("Download_Status") or "").strip().lower()
    filename=str(record.get("PDF_File_Name") or "").strip()
    if filename and (input_dir/Path(filename).name).is_file(): return True
    if status=="not_oa" and not args.recheck_all: return False
    if status in FAILURE_STATUSES and not (args.retry_failed or args.recheck_all): return False
    return True


def build_parser() -> argparse.ArgumentParser:
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook",type=Path); parser.add_argument("--input-dir",type=Path,default=Path("input_raw"))
    parser.add_argument("--start-id"); parser.add_argument("--end-id"); parser.add_argument("--max-records",type=int)
    parser.add_argument("--overwrite",action="store_true"); parser.add_argument("--retry-failed",action="store_true")
    parser.add_argument("--recheck-all",action="store_true"); parser.add_argument("--timeout",type=float,default=60)
    parser.add_argument("--sleep",type=float,default=0.75); parser.add_argument("--max-file-mb",type=float,default=150)
    parser.add_argument("--checkpoint-every",type=int,default=10)
    parser.add_argument("--unpaywall-email",default=os.environ.get("UNPAYWALL_EMAIL",""))
    parser.add_argument("--openalex-api-key",default=os.environ.get("OPENALEX_API_KEY",""))
    parser.add_argument("--dry-run",action="store_true")
    return parser


def run(args: argparse.Namespace, session: requests.Session | None = None) -> int:
    if not re.fullmatch(r"[^@\s]+@[^@\s]+",args.unpaywall_email or ""):
        raise ValueError("set UNPAYWALL_EMAIL to a real contact email or pass --unpaywall-email")
    if args.timeout<=0 or args.sleep<0 or args.max_file_mb<=0 or args.checkpoint_every<1: raise ValueError("invalid numeric option")
    repo_root=Path.cwd().resolve(); args.input_dir=args.input_dir.resolve(); workbook_path=locate_workbook(args.input_dir,args.workbook)
    wb=load_workbook(workbook_path,data_only=False,keep_links=True)
    ws=wb["PAPERS"] if "PAPERS" in wb.sheetnames else None
    if ws is None: raise ValueError("PAPERS sheet missing")
    header_row=find_header_row(ws,{"Paper_ID","DOI"}); original_sheets=list(wb.sheetnames)
    headers=header_map(ws,header_row) if args.dry_run else ensure_columns(ws,header_row,PAPERS_WRITE_FIELDS)
    records=records_from_sheet(ws,header_row); selected=selected_records(records,args)
    print(f"Workbook: {workbook_path}"); print(f"PAPERS header row: {header_row}"); print(f"PAPERS records: {len(records)}")
    print(f"Paper_ID range: {records[0]['Paper_ID']}..{records[-1]['Paper_ID']}"); print(f"Selected: {len(selected)}")
    backup=None if args.dry_run else create_backup(workbook_path)
    http=session or requests.Session(); http.headers.update({"User-Agent":f"eNH3-Bench-OA-Downloader/1.0 {args.unpaywall_email}"})
    used_names={p.name.casefold() for p in args.input_dir.glob("*.pdf")}; run_results={}; processed=0; downloaded=0
    try:
        for index,record in enumerate(selected,1):
            pid=str(record["Paper_ID"])
            if not should_process(record,args,args.input_dir):
                result=result_from_sheet(record); run_results[pid]=result; print(f"{pid}: retained {result.status}"); continue
            result=process_record(record,args,http,repo_root,used_names); run_results[pid]=result; processed+=1
            downloaded+=result.status=="downloaded"; print(f"{pid}: {result.status}")
            if not args.dry_run:
                apply_result(ws,headers,int(record["_row"]),result,workbook_path)
                if processed%args.checkpoint_every==0: atomic_save(wb,workbook_path)
            if args.sleep and index<len(selected): time.sleep(args.sleep)
    except KeyboardInterrupt:
        if not args.dry_run: atomic_save(wb,workbook_path)
        raise
    # Refresh records from in-memory cells; dry-run keeps planned results separate.
    current_records=records_from_sheet(ws,header_row)
    all_results={str(r["Paper_ID"]):result_from_sheet(r) for r in current_records}
    all_results.update(run_results)
    followup=sum(r.needs_followup for r in all_results.values())
    rag_rows=0
    if not args.dry_run:
        style_table_sheet(ws,header_row,ws.max_row,ws.max_column)
        format_papers_download_columns(ws,headers)
        rag_rows=update_rag_manifest(wb,current_records,all_results,args.input_dir)
        create_analysis(wb,workbook_path,current_records,all_results,downloaded,args.input_dir)
        followup=create_followup(wb,current_records,all_results)
        if not set(original_sheets).issubset(wb.sheetnames): raise RuntimeError("an original sheet was lost")
        atomic_save(wb,workbook_path)
    summary={"workbook":workbook_path.name,"generated_at_utc":utc_now(),"dry_run":args.dry_run,"total_papers":len(records),
             "selected":len(selected),"processed":processed,"downloaded_this_run":downloaded,
             "existing_valid":sum(r.validation_status=="existing_valid_pdf" for r in all_results.values()),
             "not_oa":sum(r.status=="not_oa" for r in all_results.values()),"followup":followup,
             "pdf_rag_manifest_rows":rag_rows,"backup":str(backup) if backup else ""}
    write_reports(repo_root/"data"/"reports",current_records,all_results,summary)
    wb.close(); print(json.dumps(summary,ensure_ascii=False)); return 0


def main(argv: Sequence[str] | None = None) -> int:
    try: return run(build_parser().parse_args(argv))
    except (ValueError,FileNotFoundError,OSError,RuntimeError) as exc:
        print(f"error: {exc}",file=sys.stderr); return 2


if __name__=="__main__": raise SystemExit(main())
