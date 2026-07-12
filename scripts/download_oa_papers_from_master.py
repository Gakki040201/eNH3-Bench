#!/usr/bin/env python3
"""Resolve and download legally open-access papers from the master workbook.

The workbook is deliberately read-only in this pass.  Progress and provenance are
stored in CSV and JSONL manifests so interrupted runs can be resumed safely.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import sys
import tempfile
import time
from dataclasses import dataclass
from datetime import datetime, timezone
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence
from urllib.parse import quote, urljoin, urlparse

import requests
from openpyxl import load_workbook


MANIFEST_FIELDS = [
    "sequence", "paper_id", "year", "title", "journal", "doi",
    "workbook_access_model", "workbook_institutional_access", "oa_verified",
    "oa_status", "status", "chosen_source", "chosen_url", "landing_url",
    "host_type", "version", "license", "file_name", "relative_path",
    "file_size_bytes", "sha256", "attempted_urls", "error", "timestamp_utc",
]
FINAL_STATUSES = {"downloaded", "skipped_existing"}
UNRESOLVED_STATUSES = {
    "not_oa", "no_pdf_url", "invalid_pdf", "metadata_error", "missing_doi",
}
NONJOURNAL_TYPES = {"preprint", "thesis", "dissertation"}
PDF_META_KEYS = {"citation_pdf_url", "dc.identifier", "eprints.document_url"}
PART_SUFFIX = ".part"


@dataclass(frozen=True)
class Candidate:
    url: str
    source: str
    landing_url: str = ""
    host_type: str = ""
    version: str = ""
    license: str = ""
    order: int = 0


@dataclass
class Resolution:
    oa_verified: bool | None
    oa_status: str
    candidates: list[Candidate]
    errors: list[str]


class PDFMetaParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.urls: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        tag = tag.lower()
        if tag not in {"meta", "link"}:
            return
        values = {str(k).lower(): (v or "") for k, v in attrs}
        if tag == "link" and values.get("type", "").lower().split(";", 1)[0] == "application/pdf":
            href = values.get("href", "").strip()
            if href:
                self.urls.append(href)
            return
        key = (values.get("name") or values.get("property") or values.get("http-equiv", "")).lower()
        content = values.get("content", "").strip()
        if key in PDF_META_KEYS and content:
            self.urls.append(content)
        elif key == "content-type" and "application/pdf" in content.lower():
            self.urls.append(content.split("url=", 1)[-1].strip())


def paper_id_number(value: Any) -> int:
    match = re.fullmatch(r"P0*(\d+)", str(value or "").strip(), flags=re.IGNORECASE)
    if not match:
        raise ValueError(f"invalid Paper_ID: {value!r}")
    return int(match.group(1))


def normalize_doi(value: Any) -> str:
    doi = str(value or "").strip()
    doi = re.sub(r"^(?:doi\s*:\s*|https?://(?:dx\.)?doi\.org/)", "", doi, flags=re.IGNORECASE)
    doi = doi.strip().strip("<>[]{}")
    doi = re.sub(r"[\s\.,;:]+$", "", doi)
    return doi.lower()


def is_supplementary_doi(doi: str) -> bool:
    value = normalize_doi(doi)
    suffix = value.split("/", 1)[-1]
    return bool(re.search(r"(?:[._/-](?:supp(?:lement(?:ary)?)?|suppl|si|s)\d*)$", suffix, re.IGNORECASE))


def is_supplementary_url(url: str) -> bool:
    path = urlparse(url).path.lower()
    return bool(re.search(r"(?:supplement(?:ary)?|supporting[-_]?information|suppinfo|[._-]si[._-])", path))


def is_journal_record(record: Mapping[str, Any]) -> bool:
    article_type = str(record.get("Article_Type") or "").strip().lower()
    # Whole-token matching is intentional: "synthesis" is not "thesis".
    if article_type in NONJOURNAL_TYPES or re.search(r"\b(?:thesis|dissertation|preprint)\b", article_type):
        return False
    journal = str(record.get("Journal") or "").strip().lower()
    return journal not in {"", "n/a", "none"}


def workbook_marks_oa(record: Mapping[str, Any]) -> bool:
    value = str(record.get("Access_Model") or "").lower()
    return bool(re.search(r"\b(?:oa|open\s+access|gold|green|diamond)\b", value))


def sanitize_component(value: Any, max_length: int = 80) -> str:
    text = re.sub(r"[<>:\"/\\|?*\x00-\x1f]", " ", str(value or ""))
    text = re.sub(r"[^\w\-.]+", "_", text, flags=re.UNICODE)
    text = re.sub(r"_+", "_", text).strip(" ._")
    return (text or "Unknown")[:max_length].rstrip(" ._")


def first_author_surname(authors: Any) -> str:
    first = str(authors or "").split(";", 1)[0].strip()
    if not first:
        return "Unknown"
    if "," in first:
        surname = first.split(",", 1)[0]
    else:
        surname = first.split()[-1]
    return sanitize_component(surname, 40)


def short_title(title: Any, max_length: int = 72) -> str:
    words = re.findall(r"[^\s]+", str(title or ""))
    chosen: list[str] = []
    for word in words:
        trial = " ".join(chosen + [word])
        if len(trial) > max_length and chosen:
            break
        chosen.append(word)
    return sanitize_component(" ".join(chosen), max_length)


def make_filename(record: Mapping[str, Any]) -> str:
    paper_id = sanitize_component(record.get("Paper_ID"), 20)
    year_match = re.search(r"\d{4}", str(record.get("Year") or ""))
    year = year_match.group(0) if year_match else "UnknownYear"
    return f"{paper_id}_{first_author_surname(record.get('Authors'))}_{year}_{short_title(record.get('Title'))}.pdf"


def locate_workbook(input_dir: Path, explicit: Path | None = None) -> Path:
    if explicit:
        path = explicit if explicit.is_absolute() else Path.cwd() / explicit
        if not path.is_file():
            raise FileNotFoundError(f"workbook not found: {path}")
        return path.resolve()
    preferred = input_dir / "eNRR_Master_Database.xlsx"
    if preferred.is_file():
        return preferred.resolve()
    choices = sorted(p for p in input_dir.glob("*.xlsx") if not p.name.startswith("~$"))
    if len(choices) == 1:
        return choices[0].resolve()
    if not choices:
        raise FileNotFoundError(f"no non-temporary .xlsx workbook in {input_dir}")
    raise ValueError(f"several .xlsx workbooks found in {input_dir}; specify --workbook")


def load_papers(workbook: Path) -> list[dict[str, Any]]:
    book = load_workbook(workbook, read_only=True, data_only=True)
    try:
        if "PAPERS" not in book.sheetnames:
            raise ValueError(f"PAPERS sheet not found in {workbook}")
        rows = book["PAPERS"].iter_rows(values_only=True)
        next(rows, None)
        headers = next(rows, None)
        if not headers or "Paper_ID" not in headers:
            raise ValueError("PAPERS row 2 does not contain the expected headers")
        records = [dict(zip(headers, row)) for row in rows if any(value is not None for value in row)]
        records = [record for record in records if str(record.get("Paper_ID") or "").strip()]
        records.sort(key=lambda record: paper_id_number(record["Paper_ID"]))
        return records
    finally:
        book.close()


def _location_candidate(location: Mapping[str, Any], source: str, order: int) -> Candidate | None:
    url = str(location.get("url_for_pdf") or location.get("pdf_url") or "").strip()
    if not url or not url.lower().startswith(("http://", "https://")) or is_supplementary_url(url):
        return None
    source_info = location.get("source") if isinstance(location.get("source"), Mapping) else {}
    host_type = str(location.get("host_type") or source_info.get("type") or "")
    return Candidate(
        url=url,
        source=source,
        landing_url=str(location.get("url_for_landing_page") or location.get("landing_page_url") or ""),
        host_type=host_type,
        version=str(location.get("version") or ""),
        license=str(location.get("license") or ""),
        order=order,
    )


def _add_locations(target: list[Candidate], payload: Mapping[str, Any], source: str) -> None:
    best = payload.get("best_oa_location")
    locations = payload.get("oa_locations") if source == "unpaywall" else payload.get("locations")
    sequence: list[Any] = ([best] if isinstance(best, Mapping) else []) + (locations if isinstance(locations, list) else [])
    for item in sequence:
        if isinstance(item, Mapping):
            candidate = _location_candidate(item, source, len(target))
            if candidate:
                target.append(candidate)


def _landing_urls(payload: Mapping[str, Any], source: str) -> list[str]:
    best = payload.get("best_oa_location")
    locations = payload.get("oa_locations") if source == "unpaywall" else payload.get("locations")
    sequence: list[Any] = ([best] if isinstance(best, Mapping) else []) + (locations if isinstance(locations, list) else [])
    urls: list[str] = []
    for item in sequence:
        if not isinstance(item, Mapping):
            continue
        url = str(item.get("url_for_landing_page") or item.get("landing_page_url") or "").strip()
        if url.startswith(("http://", "https://")):
            urls.append(url)
    return urls


def _response_json(response: requests.Response, service: str) -> Mapping[str, Any]:
    if response.status_code == 404:
        return {}
    response.raise_for_status()
    payload = response.json()
    if not isinstance(payload, Mapping):
        raise ValueError(f"{service} returned non-object JSON")
    return payload


def resolve_oa(
    session: requests.Session,
    doi: str,
    email: str,
    timeout: float,
    use_openalex: bool = True,
    inspect_landings: bool = True,
) -> Resolution:
    candidates: list[Candidate] = []
    errors: list[str] = []
    determinations: list[bool] = []
    statuses: list[str] = []
    landing_urls: list[str] = []
    encoded = quote(doi, safe="")
    try:
        response = session.get(f"https://api.unpaywall.org/v2/{encoded}", params={"email": email}, timeout=timeout)
        payload = _response_json(response, "Unpaywall")
        if payload:
            determinations.append(bool(payload.get("is_oa")))
            statuses.append(str(payload.get("oa_status") or "unknown"))
            if payload.get("is_oa"):
                _add_locations(candidates, payload, "unpaywall")
                landing_urls.extend(_landing_urls(payload, "unpaywall"))
    except (requests.RequestException, ValueError, json.JSONDecodeError) as exc:
        errors.append(f"Unpaywall: {exc}")

    if use_openalex:
        try:
            response = session.get(f"https://api.openalex.org/works/https://doi.org/{encoded}", params={"mailto": email}, timeout=timeout)
            payload = _response_json(response, "OpenAlex")
            if payload:
                oa = payload.get("open_access") if isinstance(payload.get("open_access"), Mapping) else {}
                determinations.append(bool(oa.get("is_oa")))
                statuses.append(str(oa.get("oa_status") or "unknown"))
                if oa.get("is_oa"):
                    _add_locations(candidates, payload, "openalex")
                    landing_urls.extend(_landing_urls(payload, "openalex"))
        except (requests.RequestException, ValueError, json.JSONDecodeError) as exc:
            errors.append(f"OpenAlex: {exc}")

    verified = any(determinations) if determinations else None
    if verified and inspect_landings:
        landing_urls.extend(c.landing_url for c in candidates if c.landing_url)
        candidates.extend(_landing_candidates(session, landing_urls, timeout))
    candidates = deduplicate_candidates(candidates)
    status = next((value for value in statuses if value and value != "unknown"), "unknown")
    return Resolution(verified, status, candidates, errors)


def _landing_candidates(session: requests.Session, landing_urls: Iterable[str], timeout: float) -> list[Candidate]:
    found: list[Candidate] = []
    landings = list(dict.fromkeys(url for url in landing_urls if url.startswith(("http://", "https://"))))
    for landing in landings:
        response = None
        try:
            response = session.get(
                landing, timeout=timeout, stream=True,
                headers={"Accept": "text/html,application/xhtml+xml"},
            )
            response.raise_for_status()
            content_type = response.headers.get("Content-Type", "").lower()
            if "html" not in content_type:
                continue
            chunks: list[bytes] = []
            size = 0
            for chunk in response.iter_content(chunk_size=64 * 1024):
                if not chunk:
                    continue
                remaining = 2_000_000 - size
                chunks.append(chunk[:remaining])
                size += min(len(chunk), remaining)
                if size >= 2_000_000:
                    break
            parser = PDFMetaParser()
            parser.feed(b"".join(chunks).decode(response.encoding or "utf-8", errors="replace"))
            for raw_url in parser.urls:
                url = urljoin(landing, raw_url)
                if url.startswith(("http://", "https://")) and not is_supplementary_url(url):
                    found.append(Candidate(url, "landing_metadata", landing_url=landing, order=len(found)))
        except (requests.RequestException, UnicodeError):
            continue
        finally:
            if response is not None:
                response.close()
    return found


def deduplicate_candidates(candidates: Iterable[Candidate]) -> list[Candidate]:
    unique: list[Candidate] = []
    seen: set[str] = set()
    for candidate in candidates:
        key = candidate.url.strip()
        if key not in seen:
            seen.add(key)
            unique.append(candidate)
    return unique


def candidate_sort_key(candidate: Candidate, prefer_publisher: bool = True) -> tuple[int, int, int, int]:
    host = candidate.host_type.lower()
    version = candidate.version.lower()
    host_rank = 0 if host == "publisher" else (2 if host == "repository" else 1)
    version_rank = {"publishedversion": 0, "acceptedversion": 1, "submittedversion": 2}.get(version, 3)
    source_rank = {"unpaywall": 0, "openalex": 1, "landing_metadata": 2}.get(candidate.source, 3)
    return (host_rank if prefer_publisher else 0, version_rank, source_rank, candidate.order)


def ordered_candidates(candidates: Iterable[Candidate], prefer_publisher: bool = True) -> list[Candidate]:
    return sorted(deduplicate_candidates(candidates), key=lambda item: candidate_sort_key(item, prefer_publisher))


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def valid_pdf(path: Path, min_size: int = 1024) -> bool:
    try:
        if not path.is_file() or path.stat().st_size < min_size:
            return False
        with path.open("rb") as handle:
            return handle.read(5) == b"%PDF-"
    except OSError:
        return False


def find_existing_pdf(record: Mapping[str, Any], destination: Path, output_dir: Path, repo_root: Path) -> Path | None:
    candidates = [destination]
    workbook_name = str(record.get("PDF_File_Name") or "").strip()
    if workbook_name:
        candidates.append(output_dir / Path(workbook_name).name)
    workbook_path = str(record.get("Local_PDF_Path") or "").strip()
    if workbook_path:
        path = Path(workbook_path)
        candidates.append(path if path.is_absolute() else repo_root / path)
    for path in candidates:
        if valid_pdf(path):
            return path
    return None


def download_pdf(
    session: requests.Session,
    candidate: Candidate,
    destination: Path,
    timeout: float,
    max_file_mb: float,
) -> tuple[int, str]:
    destination.parent.mkdir(parents=True, exist_ok=True)
    part = destination.with_suffix(destination.suffix + PART_SUFFIX)
    max_bytes = int(max_file_mb * 1024 * 1024)
    try:
        with session.get(candidate.url, timeout=timeout, stream=True, headers={"Accept": "application/pdf"}) as response:
            response.raise_for_status()
            declared = int(response.headers.get("Content-Length") or 0)
            if declared > max_bytes:
                raise ValueError(f"response exceeds --max-file-mb ({declared} bytes)")
            size = 0
            digest = hashlib.sha256()
            prefix = b""
            with part.open("wb") as handle:
                for chunk in response.iter_content(chunk_size=64 * 1024):
                    if not chunk:
                        continue
                    if len(prefix) < 5:
                        prefix = (prefix + chunk)[:5]
                    size += len(chunk)
                    if size > max_bytes:
                        raise ValueError(f"response exceeds --max-file-mb ({size} bytes)")
                    handle.write(chunk)
                    digest.update(chunk)
            if prefix != b"%PDF-":
                raise ValueError("response does not start with %PDF-")
            if size < 1024:
                raise ValueError(f"PDF is smaller than 1 KB ({size} bytes)")
        os.replace(part, destination)
        return size, digest.hexdigest()
    except Exception:
        try:
            part.unlink(missing_ok=True)
        except OSError:
            pass
        raise


def load_manifest(path: Path) -> dict[str, dict[str, Any]]:
    if not path.is_file():
        return {}
    with path.open("r", newline="", encoding="utf-8-sig") as handle:
        return {row["paper_id"]: dict(row) for row in csv.DictReader(handle) if row.get("paper_id")}


def write_manifests(rows_by_id: Mapping[str, Mapping[str, Any]], csv_path: Path, jsonl_path: Path) -> None:
    rows = sorted(rows_by_id.values(), key=lambda row: int(row["sequence"]))
    csv_path.parent.mkdir(parents=True, exist_ok=True)
    csv_fd, csv_tmp_name = tempfile.mkstemp(prefix=csv_path.name, suffix=".tmp", dir=csv_path.parent)
    json_fd, json_tmp_name = tempfile.mkstemp(prefix=jsonl_path.name, suffix=".tmp", dir=jsonl_path.parent)
    os.close(csv_fd)
    os.close(json_fd)
    csv_tmp = Path(csv_tmp_name)
    json_tmp = Path(json_tmp_name)
    try:
        with csv_tmp.open("w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.DictWriter(handle, fieldnames=MANIFEST_FIELDS, extrasaction="ignore")
            writer.writeheader()
            for row in rows:
                output = dict(row)
                output["attempted_urls"] = json.dumps(output.get("attempted_urls", []), ensure_ascii=False)
                writer.writerow(output)
        with json_tmp.open("w", encoding="utf-8") as handle:
            for row in rows:
                handle.write(json.dumps(dict(row), ensure_ascii=False) + "\n")
        os.replace(csv_tmp, csv_path)
        os.replace(json_tmp, jsonl_path)
    finally:
        csv_tmp.unlink(missing_ok=True)
        json_tmp.unlink(missing_ok=True)


def base_manifest_row(record: Mapping[str, Any], sequence: int, doi: str, filename: str) -> dict[str, Any]:
    return {
        "sequence": sequence, "paper_id": str(record.get("Paper_ID") or ""),
        "year": record.get("Year") or "", "title": record.get("Title") or "",
        "journal": record.get("Journal") or "", "doi": doi,
        "workbook_access_model": record.get("Access_Model") or "",
        "workbook_institutional_access": record.get("Institutional_Access") or "",
        "oa_verified": "", "oa_status": "", "status": "", "chosen_source": "",
        "chosen_url": "", "landing_url": "", "host_type": "", "version": "",
        "license": "", "file_name": filename, "relative_path": "",
        "file_size_bytes": "", "sha256": "", "attempted_urls": [], "error": "",
        "timestamp_utc": datetime.now(timezone.utc).isoformat(),
    }


def relative_to_repo(path: Path, repo_root: Path) -> str:
    try:
        return path.resolve().relative_to(repo_root.resolve()).as_posix()
    except ValueError:
        return path.resolve().as_posix()


def process_record(
    record: Mapping[str, Any], sequence: int, args: argparse.Namespace,
    session: requests.Session, repo_root: Path,
) -> dict[str, Any]:
    doi = normalize_doi(record.get("DOI"))
    filename = make_filename(record)
    destination = args.output_dir / filename
    row = base_manifest_row(record, sequence, doi, filename)

    existing = None if args.overwrite else find_existing_pdf(record, destination, args.output_dir, repo_root)
    if existing:
        row.update(status="skipped_existing", file_name=existing.name,
                   relative_path=relative_to_repo(existing, repo_root),
                   file_size_bytes=existing.stat().st_size, sha256=file_sha256(existing))
        return row
    if not doi:
        row.update(status="missing_doi", error="DOI is blank")
        return row
    if is_supplementary_doi(doi):
        row.update(status="no_pdf_url", error="supplementary-material DOI excluded")
        return row

    resolution = resolve_oa(session, doi, args.unpaywall_email, args.timeout, not args.no_openalex)
    row["oa_verified"] = resolution.oa_verified if resolution.oa_verified is not None else ""
    row["oa_status"] = resolution.oa_status
    if resolution.oa_verified is None:
        row.update(status="metadata_error", error="; ".join(resolution.errors) or "OA metadata unavailable")
        return row
    if not resolution.oa_verified:
        row.update(status="not_oa", error="; ".join(resolution.errors))
        return row
    candidates = ordered_candidates(resolution.candidates, not args.no_prefer_publisher)
    if not candidates:
        row.update(status="no_pdf_url", error="OA verified but no explicit PDF URL found")
        return row

    row["attempted_urls"] = [candidate.url for candidate in candidates]
    chosen = candidates[0]
    row.update(chosen_source=chosen.source, chosen_url=chosen.url, landing_url=chosen.landing_url,
               host_type=chosen.host_type, version=chosen.version, license=chosen.license)
    if args.dry_run:
        row["status"] = "dry_run_selected"
        return row

    errors: list[str] = []
    for candidate in candidates:
        row.update(chosen_source=candidate.source, chosen_url=candidate.url, landing_url=candidate.landing_url,
                   host_type=candidate.host_type, version=candidate.version, license=candidate.license)
        try:
            size, digest = download_pdf(session, candidate, destination, args.timeout, args.max_file_mb)
            row.update(status="downloaded", relative_path=relative_to_repo(destination, repo_root),
                       file_size_bytes=size, sha256=digest, error="")
            return row
        except (requests.RequestException, OSError, ValueError) as exc:
            errors.append(f"{candidate.url}: {exc}")
    row.update(status="invalid_pdf", error="; ".join(errors))
    return row


def select_records(records: Sequence[Mapping[str, Any]], args: argparse.Namespace) -> list[Mapping[str, Any]]:
    start = paper_id_number(args.start_id) if args.start_id else None
    end = paper_id_number(args.end_id) if args.end_id else None
    selected: list[Mapping[str, Any]] = []
    for record in records:
        number = paper_id_number(record.get("Paper_ID"))
        if start is not None and number < start:
            continue
        if end is not None and number > end:
            continue
        if not args.include_nonjournal and not is_journal_record(record):
            continue
        if args.only_workbook_marked_oa and not workbook_marks_oa(record):
            continue
        selected.append(record)
        if args.max_records is not None and len(selected) >= args.max_records:
            break
    return selected


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-dir", type=Path, default=Path("input_raw"))
    parser.add_argument("--output-dir", type=Path, default=Path("input_raw"))
    parser.add_argument("--manifest-dir", type=Path, default=Path("data/reports"))
    parser.add_argument("--workbook", type=Path)
    parser.add_argument("--unpaywall-email", required=True, help="real contact email required by Unpaywall")
    parser.add_argument("--start-id")
    parser.add_argument("--end-id")
    parser.add_argument("--max-records", type=int)
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--overwrite", action="store_true")
    parser.add_argument("--retry-unresolved", action="store_true")
    parser.add_argument("--only-workbook-marked-oa", action="store_true")
    parser.add_argument("--include-nonjournal", action="store_true")
    parser.add_argument("--no-prefer-publisher", action="store_true")
    parser.add_argument("--no-openalex", action="store_true")
    parser.add_argument("--timeout", type=float, default=30.0)
    parser.add_argument("--sleep", type=float, default=0.1)
    parser.add_argument("--max-file-mb", type=float, default=100.0)
    return parser


def run(args: argparse.Namespace, session: requests.Session | None = None) -> int:
    args.input_dir = args.input_dir.resolve()
    args.output_dir = args.output_dir.resolve()
    args.manifest_dir = args.manifest_dir.resolve()
    if args.max_records is not None and args.max_records < 1:
        raise ValueError("--max-records must be positive")
    if args.max_file_mb <= 0 or args.timeout <= 0 or args.sleep < 0:
        raise ValueError("--timeout and --max-file-mb must be positive; --sleep cannot be negative")
    workbook = locate_workbook(args.input_dir, args.workbook)
    records = load_papers(workbook)
    selected = select_records(records, args)
    repo_root = Path.cwd().resolve()
    csv_path = args.manifest_dir / "oa_download_manifest.csv"
    jsonl_path = args.manifest_dir / "oa_download_manifest.jsonl"
    manifest = load_manifest(csv_path)
    http = session or requests.Session()
    http.headers.update({"User-Agent": f"eNH3-Bench-OA-downloader/1.0 (mailto:{args.unpaywall_email})"})

    print(f"Workbook: {workbook}")
    print(f"PAPERS records: {len(records)}")
    print(f"Eligible selected records: {len(selected)}")
    processed = 0
    for record in selected:
        paper_id = str(record["Paper_ID"])
        prior = manifest.get(paper_id)
        if prior and prior.get("status") in UNRESOLVED_STATUSES and not args.retry_unresolved:
            print(f"{paper_id}: retained {prior['status']} (use --retry-unresolved)")
            continue
        sequence = paper_id_number(paper_id)
        row = process_record(record, sequence, args, http, repo_root)
        manifest[paper_id] = row
        write_manifests(manifest, csv_path, jsonl_path)
        processed += 1
        print(f"{paper_id}: {row['status']}")
        if args.sleep and processed < len(selected):
            time.sleep(args.sleep)
    if not processed:
        write_manifests(manifest, csv_path, jsonl_path)
    print(f"Processed this run: {processed}")
    print(f"Manifest: {csv_path}")
    return 0


def main(argv: Sequence[str] | None = None) -> int:
    try:
        return run(build_parser().parse_args(argv))
    except (FileNotFoundError, ValueError, OSError) as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
