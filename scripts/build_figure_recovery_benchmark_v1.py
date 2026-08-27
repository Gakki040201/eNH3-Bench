from __future__ import annotations

import argparse
import json
import re
import sys
import zipfile
from pathlib import Path, PurePosixPath

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from enh3bench.figure_recovery.locator import locate_figures  # noqa: E402
from enh3bench.figure_recovery.schema import file_sha256, stable_id, write_json  # noqa: E402


DEFAULT_REPORT_ROOT = Path(r"F:\eNH3_Bench_Work\04_Exports\Reports\LiNRR_Figure_Recovery_v1")
FIGURE_RE = re.compile(r"(?:Fig(?:ure)?)[ _-]*(S?\d+)", re.IGNORECASE)


def build(
    *, source_data_zip: Path, supplementary_pdf: Path, paper_id: str, bundle_id: str,
    source_asset_id: str, output_root: Path, limit: int,
) -> dict[str, object]:
    if not 10 <= limit <= 20:
        raise ValueError("benchmark pilot must contain 10-20 figures")
    output_root.mkdir(parents=True, exist_ok=True)
    truth_root = output_root / "ground_truth_source_data"
    truth_root.mkdir(parents=True, exist_ok=True)
    located = locate_figures(supplementary_pdf, paper_id=paper_id, bundle_id=bundle_id, source_asset_id=source_asset_id)
    by_label: dict[str, list[object]] = {}
    for figure in located:
        key = _label_key(figure.figure_label)
        by_label.setdefault(key, []).append(figure)
    candidates = []
    with zipfile.ZipFile(source_data_zip) as archive:
        for info in sorted(archive.infolist(), key=lambda item: item.filename.lower()):
            if info.is_dir() or not info.filename.lower().endswith(".xlsx"):
                continue
            match = FIGURE_RE.search(PurePosixPath(info.filename).name)
            if not match:
                continue
            label = match.group(1).upper().replace(" ", "")
            matching = by_label.get(label, [])
            if not matching:
                continue
            safe_name = f"{label}_{stable_id('GT', paper_id, info.filename)[3:]}.xlsx"
            output = truth_root / safe_name
            with archive.open(info) as source, output.open("wb") as target:
                target.write(source.read())
            profile = _workbook_profile(output)
            if profile["numeric_pair_rows"] < 2:
                output.unlink()
                continue
            figure = matching[0]
            candidates.append({
                "benchmark_id": stable_id("BMF", paper_id, figure.figure_id, info.filename),
                "paper_id": paper_id, "bundle_id": bundle_id,
                "figure_id": figure.figure_id, "panel_id": figure.panel_id,
                "figure_label": figure.figure_label, "page": figure.page,
                "figure_bbox": list(figure.figure_bbox), "panel_bbox": list(figure.panel_bbox) if figure.panel_bbox else None,
                "figure_content_type": figure.figure_content_type.value,
                "source_pdf": str(supplementary_pdf), "source_pdf_sha256": file_sha256(supplementary_pdf),
                "ground_truth_asset": str(output), "ground_truth_sha256": file_sha256(output),
                "ground_truth_origin": f"{source_data_zip}!/{info.filename}",
                "ground_truth_profile": profile, "caption": figure.caption,
            })
            if len(candidates) >= limit:
                break
    if len(candidates) < 10:
        raise RuntimeError(f"only {len(candidates)} source-backed benchmark figures could be matched; no benchmark was finalized")
    payload = {
        "benchmark_name": "LiNRR_Figure_Recovery_Benchmark_v1", "schema_version": "figure-recovery-benchmark/1.0",
        "threshold_policy": "No PASS threshold is frozen; report empirical distributions first.",
        "benchmark_figures": candidates,
        "counts": {"total": len(candidates), "vector": sum(item["figure_content_type"] == "VECTOR" for item in candidates), "raster": sum(item["figure_content_type"] == "RASTER" for item in candidates), "mixed": sum(item["figure_content_type"] == "MIXED" for item in candidates)},
    }
    write_json(output_root / "benchmark_manifest.json", payload)
    return payload


def _label_key(label: str) -> str:
    match = re.search(r"(?:S\s*)?(\d+)", label, re.IGNORECASE)
    if not match:
        return label.upper().replace(" ", "")
    supplementary = bool(re.search(r"Supplementary|S\s*\d+", label, re.IGNORECASE))
    return ("S" if supplementary else "") + match.group(1)


def _workbook_profile(path: Path) -> dict[str, object]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheets = []
    total_pairs = 0
    for worksheet in workbook.worksheets:
        rows, pairs = 0, 0
        for row in worksheet.iter_rows(values_only=True):
            values = [value for value in row if value is not None]
            if values:
                rows += 1
            if sum(isinstance(value, (int, float)) and not isinstance(value, bool) for value in values) >= 2:
                pairs += 1
        total_pairs += pairs
        sheets.append({"title": worksheet.title, "nonempty_rows": rows, "numeric_pair_rows": pairs})
    workbook.close()
    return {"sheets": sheets, "numeric_pair_rows": total_pairs}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build the source-data-backed LiNRR figure recovery benchmark v1.")
    parser.add_argument("--source-data-zip", type=Path, required=True)
    parser.add_argument("--supplementary-pdf", type=Path, required=True)
    parser.add_argument("--paper-id", required=True)
    parser.add_argument("--bundle-id", required=True)
    parser.add_argument("--source-asset-id", required=True)
    parser.add_argument("--output-root", type=Path, default=DEFAULT_REPORT_ROOT)
    parser.add_argument("--limit", type=int, default=15)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    result = build(source_data_zip=args.source_data_zip, supplementary_pdf=args.supplementary_pdf, paper_id=args.paper_id, bundle_id=args.bundle_id, source_asset_id=args.source_asset_id, output_root=args.output_root, limit=args.limit)
    print(json.dumps(result["counts"], indent=2, sort_keys=True))
    print(args.output_root / "benchmark_manifest.json")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
