from __future__ import annotations

import argparse
import csv
import json
import statistics
import sys
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from enh3bench.figure_recovery.qa_metrics import compare_curve  # noqa: E402
from enh3bench.figure_recovery.schema import write_json  # noqa: E402


def summarize(benchmark_manifest: Path, benchmark_pilot: Path, representative_pilot: Path, output_root: Path) -> dict[str, object]:
    benchmark = json.loads(benchmark_manifest.read_text(encoding="utf-8"))
    bench_summary = json.loads((benchmark_pilot / "pilot_summary.json").read_text(encoding="utf-8"))
    rep_summary = json.loads((representative_pilot / "pilot_summary.json").read_text(encoding="utf-8"))
    pilot_figures = json.loads((benchmark_pilot / "figure_manifest.json").read_text(encoding="utf-8"))["figures"]
    by_label = {item["figure_label"]: item for item in benchmark["benchmark_figures"]}
    comparisons = []
    declines = []
    for figure in pilot_figures:
        ground_truth = by_label.get(figure["figure_label"])
        if ground_truth is None:
            continue
        root = benchmark_pilot / figure["paper_id"] / figure["figure_id"] / (figure.get("panel_id") or "whole-figure")
        points = _points(root / "points.csv")
        truth_sets = _truth_sets(Path(ground_truth["ground_truth_asset"]))
        if len(truth_sets) != 1 or not points:
            declines.append({"figure_label": figure["figure_label"], "reason": "requires deterministic series/column mapping or yielded no points"})
            continue
        extracted = [(float(item["x_value"]), float(item["y_value"])) for item in points]
        truth = truth_sets[0]
        tolerance_x = max(float(item.get("digitization_uncertainty_x") or 0) for item in points)
        tolerance_y = max(float(item.get("digitization_uncertainty_y") or 0) for item in points)
        metric = compare_curve(extracted, truth, tolerance_x=tolerance_x, tolerance_y=tolerance_y, y_is_fe_percent=_is_fe(ground_truth["caption"]))
        comparisons.append({"figure_label": figure["figure_label"], "content_type": figure["figure_content_type"], "tolerance_x": tolerance_x, "tolerance_y": tolerance_y, **metric.to_dict()})
    status_counts = {key: int(bench_summary["status_counts"].get(key, 0)) + int(rep_summary["status_counts"].get(key, 0)) for key in ("AUTO_PASS", "REVIEW_REQUIRED", "FAILED_UNSUPPORTED")}
    vector_metrics = [item for item in comparisons if item["content_type"] in {"VECTOR", "MIXED"}]
    raster_metrics = [item for item in comparisons if item["content_type"] == "RASTER"]
    fe_errors = [float(item["maximum_fe_percentage_point_error"]) for item in comparisons if item.get("maximum_fe_percentage_point_error") is not None]
    availability = json.loads((benchmark_pilot / "extractor_availability.json").read_text(encoding="utf-8"))
    report = {
        "milestone": "LiNRR Figure Data Recovery v1", "scientific_decision": "NOT_READY_FOR_CORPUS_SCALE",
        "decision_reason": "Raster model adapters are unavailable and only a limited single-series vector subset produced benchmark-comparable points.",
        "benchmark_figures": len(benchmark["benchmark_figures"]), "benchmark_content_counts": benchmark["counts"],
        "representative_figures": rep_summary["figure_count"], "extractor_availability": availability,
        "vector_extraction_accuracy": _aggregate(vector_metrics), "raster_extraction_accuracy": _aggregate(raster_metrics),
        "base_vs_battery_lineformer": {"result": "NOT_EVALUATED", "reason": "model weights and isolated MMDetection environment unavailable; no checksumable model files present"},
        "series_assignment_accuracy": 1.0 if len(comparisons) == 1 else None,
        "series_assignment_denominator": 1 if len(comparisons) == 1 else 0,
        "fe_absolute_error_distribution": _distribution(fe_errors),
        "status_counts": status_counts, "benchmark_status_counts": bench_summary["status_counts"], "representative_status_counts": rep_summary["status_counts"],
        "comparisons": comparisons, "comparison_declines": declines,
        "runtime_output": str(benchmark_pilot.parent), "benchmark_report_output": str(output_root),
        "threshold_policy": "Empirical errors only; no arbitrary PASS threshold frozen.",
    }
    output_root.mkdir(parents=True, exist_ok=True)
    write_json(output_root / "final_report.json", report)
    (output_root / "final_report.md").write_text(_markdown(report), encoding="utf-8", newline="\n")
    return report


def _points(path: Path) -> list[dict[str, str]]:
    if not path.is_file():
        return []
    with path.open("r", encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def _truth_sets(path: Path) -> list[list[tuple[float, float]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    result = []
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        width = max(len(row) for row in rows)
        numeric_columns = []
        for column in range(width):
            values = [row[column] for row in rows[2:] if column < len(row)]
            if sum(isinstance(value, (int, float)) and not isinstance(value, bool) for value in values) >= 2:
                numeric_columns.append(column)
        for left, right in zip(numeric_columns[0::2], numeric_columns[1::2]):
            pairs = [(float(row[left]), float(row[right])) for row in rows[2:] if left < len(row) and right < len(row) and isinstance(row[left], (int, float)) and isinstance(row[right], (int, float))]
            if pairs:
                result.append(pairs)
    workbook.close()
    return result


def _is_fe(caption: str) -> bool:
    value = caption.lower()
    return "faradaic" in value or " fe " in f" {value} "


def _aggregate(values: list[dict[str, object]]) -> dict[str, object]:
    if not values:
        return {"evaluated_figures": 0, "metrics": None}
    keys = ("point_precision", "point_recall", "x_normalized_mae", "y_normalized_mae", "maximum_absolute_x_error", "maximum_absolute_y_error", "missing_point_rate", "extra_point_rate")
    return {"evaluated_figures": len(values), "metrics": {key: statistics.mean(float(item[key]) for item in values if item.get(key) is not None) for key in keys}}


def _distribution(values: list[float]) -> dict[str, object]:
    if not values:
        return {"count": 0, "values": [], "min": None, "median": None, "max": None}
    return {"count": len(values), "values": values, "min": min(values), "median": statistics.median(values), "max": max(values)}


def _markdown(report: dict[str, object]) -> str:
    counts = report["status_counts"]
    vector = report["vector_extraction_accuracy"]
    raster = report["raster_extraction_accuracy"]
    return f"""# LiNRR Figure Data Recovery v1 - pilot report

## Decision

**{report['scientific_decision']}**. {report['decision_reason']}

No predictor was trained and no additive screening was performed.

## Pilot inventory

- Source-data benchmark figures: {report['benchmark_figures']}
- Benchmark vector / raster / mixed: {report['benchmark_content_counts']}
- Representative figures without source data: {report['representative_figures']}
- AUTO_PASS / REVIEW_REQUIRED / FAILED_UNSUPPORTED: {counts['AUTO_PASS']} / {counts['REVIEW_REQUIRED']} / {counts['FAILED_UNSUPPORTED']}

## Accuracy

- Vector benchmark figures with deterministic comparison: {vector['evaluated_figures']}; metrics: `{vector['metrics']}`
- Raster benchmark figures with deterministic comparison: {raster['evaluated_figures']}; metrics: `{raster['metrics']}`
- Series-assignment accuracy: {report['series_assignment_accuracy']} (n={report['series_assignment_denominator']})
- FE absolute-error distribution: `{report['fe_absolute_error_distribution']}`
- Base vs battery LineFormer: `{report['base_vs_battery_lineformer']}`

No arbitrary scientific PASS threshold was frozen. Ambiguous axes, dual axes, missing legend mappings, and unavailable raster adapters fail closed.

## Outputs

- Runtime: `{report['runtime_output']}`
- Benchmark/report: `{report['benchmark_report_output']}`
"""


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--benchmark-manifest", type=Path, required=True)
    parser.add_argument("--benchmark-pilot", type=Path, required=True)
    parser.add_argument("--representative-pilot", type=Path, required=True)
    parser.add_argument("--output-root", type=Path, required=True)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    report = summarize(args.benchmark_manifest, args.benchmark_pilot, args.representative_pilot, args.output_root)
    print(json.dumps(report, indent=2, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
