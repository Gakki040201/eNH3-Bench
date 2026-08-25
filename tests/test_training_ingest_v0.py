from __future__ import annotations

import csv
import json
import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import Workbook

from enh3bench.training_ingest import (
    _empty_record,
    apply_admission,
    build_battery_additive_prior,
    classify_document,
    deterministic_dataset_hash,
    inspect_legacy_mapping,
    mark_duplicates_and_conflicts,
    match_paper_identity,
    normalize_doi,
    normalize_measurement,
    read_asset_blocks,
    scan_bundle,
    sha256_file,
)


class TrainingIngestV0Tests(unittest.TestCase):
    def test_bundle_ownership_isolation_and_sha_manifest(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            one, two = root / "one", root / "two"
            one.mkdir(); two.mkdir()
            (one / "supplementary.txt").write_text("DOI: 10.1000/one", encoding="utf-8")
            (two / "supplementary.txt").write_text("DOI: 10.1000/two", encoding="utf-8")
            first = scan_bundle(one, "shaofeng_li_curated")
            second = scan_bundle(two, "shaofeng_li_curated")
            self.assertNotEqual(first["bundle_id"], second["bundle_id"])
            self.assertNotEqual(first["assets"][0]["source_asset_id"], second["assets"][0]["source_asset_id"])
            self.assertTrue(Path(first["assets"][0]["absolute_path"]).is_relative_to(one))
            self.assertEqual(first["assets"][0]["sha256"], sha256_file(one / "supplementary.txt"))

    def test_docx_si_parsing(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "supp.docx"
            doc = Document(); doc.add_paragraph("Supporting Information")
            table = doc.add_table(rows=1, cols=2); table.cell(0, 0).text = "FE"; table.cell(0, 1).text = "42%"
            doc.save(path)
            blocks = read_asset_blocks(path)
            self.assertTrue(any(block["locator"].startswith("table:") and "42%" in block["text"] for block in blocks))

    def test_xlsx_supplementary_parsing(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "source.xlsx"
            workbook = Workbook(); sheet = workbook.active; sheet.title = "Experiments"
            sheet.append(["FE", "LiBF4"]); sheet.append([42, "THF"]); workbook.save(path)
            blocks = read_asset_blocks(path)
            self.assertEqual(blocks[1]["locator"], "sheet:Experiments:row:2")
            self.assertIn("THF", blocks[1]["text"])

    def test_doi_normalization_and_exact_registry_match(self):
        self.assertEqual(normalize_doi("https://doi.org/10.1021/JACS.6C08057."), "10.1021/jacs.6c08057")
        match = match_paper_identity({"doi": "DOI:10.1000/ABC", "title": "Different"}, [{"paper_id": "P0042", "doi": "10.1000/abc", "title": "Registered"}])
        self.assertEqual(match["existing_paper_id"], "P0042")
        self.assertEqual(match["match_type"], "DOI_EXACT")

    def test_legacy_mapping_only_uses_present_fields(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "legacy.jsonl"
            path.write_text(json.dumps({"paper_id": "P1", "span_id": "S1", "faradaic_efficiency_percent": 50}) + "\n", encoding="utf-8")
            fields = {row["legacy_field"] for row in inspect_legacy_mapping(path)}
            self.assertEqual(fields, {"paper_id", "span_id", "faradaic_efficiency_percent"})

    def test_missing_values_remain_null(self):
        record = _empty_record()
        self.assertIsNone(record["water_content_value"])
        self.assertIsNone(record["proton_donor_concentration_value"])

    def test_unit_normalization_is_conservative(self):
        self.assertEqual(normalize_measurement(0.25, "A/cm2", "current_density")["normalized_value"], 250.0)
        self.assertAlmostEqual(normalize_measurement(30, "minutes", "duration")["normalized_value"], 0.5)
        self.assertAlmostEqual(normalize_measurement(101.325, "kPa", "pressure")["normalized_value"], 1.01325)
        self.assertEqual(normalize_measurement(1, "M", "concentration")["normalized_unit"], "mol/L")
        self.assertIsNone(normalize_measurement(5, "vol%", "concentration")["normalized_value"])
        self.assertIsNone(normalize_measurement(100, "ppm", "concentration")["normalized_value"])

    def test_battery_and_review_cannot_become_targets(self):
        reaction, genre, role = classify_document({"title": "Lithium battery additive"}, "ABSTRACT experimental zinc battery", "battery_transfer_curated")
        self.assertEqual((reaction, role), ("battery", "descriptor_prior_only"))
        reaction, genre, role = classify_document({"title": "Perspective on lithium-mediated nitrogen reduction"}, "N2 and ammonia", "shaofeng_li_curated")
        self.assertEqual(genre, "review_or_perspective")
        self.assertEqual(role, "evidence_only")

    def test_cross_paper_evidence_cannot_pass_exact_source_gate(self):
        record = _empty_record()
        record.update({"record_id": "R1", "paper_id": "P1", "source_bundle_id": "P2", "evidence_owner_id": "P2", "source_group": "legacy_oa", "reaction_family": "LiNRR", "document_genre": "primary_experimental", "training_role": "target_label_candidate", "source_asset_id": "A1", "source_file_sha256": "a" * 64, "source_locator": "S1", "fe_nh3_percent": 50, "lithium_salt": "LiBF4", "solvent": "THF", "duration_h": 1})
        admission = apply_admission(record)
        self.assertTrue(admission["gate_results"]["exact_source"])
        self.assertFalse(admission["gate_results"]["ownership_clear"])
        self.assertFalse(admission["model_eligible_fe"])

    def test_dataset_output_hash_is_deterministic(self):
        rows = [{"record_id": "B", "value": None}, {"record_id": "A", "value": 1}]
        self.assertEqual(deterministic_dataset_hash(rows), deterministic_dataset_hash(list(reversed(rows))))

    def test_structured_author_source_wins_duplicate_admission(self):
        common = {"paper_id": "P0001", "lithium_salt": "LiBF4", "solvent": "THF", "current_density_mA_cm2": 10, "duration_h": 1, "fe_nh3_percent": 50}
        prose = {**_empty_record(), **common, "record_id": "PROSE", "evidence_authority": "primary_text_candidate", "source_filename": "main.pdf", "source_locator": "page:2"}
        table = {**_empty_record(), **common, "record_id": "TABLE", "evidence_authority": "primary_structured_author_experiment", "source_filename": "si.docx", "source_locator": "table:1:row:2"}
        mark_duplicates_and_conflicts([prose, table])
        self.assertIn("POSSIBLE_DUPLICATE_ROW", prose["ambiguity_flags"])
        self.assertNotIn("POSSIBLE_DUPLICATE_ROW", table["ambiguity_flags"])


if __name__ == "__main__":
    unittest.main()
