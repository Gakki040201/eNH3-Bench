import argparse
import hashlib
import os
import tempfile
import unittest
from copy import copy
from pathlib import Path
from unittest.mock import Mock, patch

from openpyxl import Workbook, load_workbook

from scripts.download_oa_writeback_excel import (
    Candidate,
    Result,
    allocate_filename,
    apply_result,
    atomic_save,
    create_analysis,
    create_backup,
    create_followup,
    download_pdf,
    ensure_columns,
    find_header_row,
    header_map,
    inspect_pdf,
    is_supplementary_reference,
    normalize_doi,
    paper_id_number,
    records_from_sheet,
    run,
    update_rag_manifest,
)


class FakeResponse:
    def __init__(self, body=b"", status=200, content_type="application/pdf", url="https://example.org/a.pdf"):
        self.body = body
        self.status_code = status
        self.url = url
        self.encoding = "utf-8"
        self.headers = {"Content-Type": content_type, "Content-Length": str(len(body))}

    def __enter__(self): return self
    def __exit__(self, *args): return False
    def close(self): pass
    def raise_for_status(self):
        if self.status_code >= 400:
            import requests
            raise requests.HTTPError(str(self.status_code))
    def iter_content(self, chunk_size=65536):
        yield self.body


def base_record(pid="P0001"):
    return {"Paper_ID": pid, "DOI": "10.1000/test", "Title": "A useful synthesis study",
            "Authors": "Jane Q. Smith; John Roe", "Year": 2024, "Journal": "Journal",
            "Publisher_URL": "", "Access_Model": "Gold OA", "Institutional_Access": ""}


def make_book(path, header_row=2):
    wb = Workbook(); ws = wb.active; ws.title = "PAPERS"; ws.cell(1, 1, "PAPERS title")
    headers = ["Paper_ID", "DOI", "Title", "Authors", "Year", "Journal", "Publisher_URL",
               "Access_Model", "Institutional_Access", "Download_Status", "Local_PDF_Path", "PDF_File_Name"]
    for col, name in enumerate(headers, 1): ws.cell(header_row, col, name)
    for col, name in enumerate(headers, 1): ws.cell(header_row + 1, col, base_record().get(name, ""))
    other = wb.create_sheet("KEEP_ME"); other["A1"] = "preserve"
    rag = wb.create_sheet("PDF_RAG_MANIFEST"); rag["A1"] = "Manifest title"; rag["A2"] = "Paper_ID"
    wb.save(path); wb.close()


class DownloadOAWritebackTests(unittest.TestCase):
    def test_detects_second_row_header(self):
        wb = Workbook(); ws = wb.active; ws.append(["title"]); ws.append(["Paper_ID", "DOI"])
        self.assertEqual(find_header_row(ws, {"Paper_ID", "DOI"}), 2)

    def test_numeric_paper_id_ordering(self):
        self.assertEqual(sorted(["P0010", "P0002", "P0009"], key=paper_id_number), ["P0002", "P0009", "P0010"])

    def test_p0009_precedes_p0010(self):
        self.assertLess(paper_id_number("P0009"), paper_id_number("P0010"))

    def test_doi_normalization(self):
        self.assertEqual(normalize_doi(" https://doi.org/10.1038/ABC.1. "), "10.1038/abc.1")

    def test_supplementary_url_rejected(self):
        self.assertTrue(is_supplementary_reference("https://x.org/article_s001.pdf"))
        self.assertTrue(is_supplementary_reference("https://x.org/supporting-information/file.pdf"))
        self.assertFalse(is_supplementary_reference("https://x.org/article.pdf"))

    def test_valid_pdf_accepted(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "a.pdf"; path.write_bytes(b"%PDF-1.7\n" + b"x" * 11000)
            self.assertEqual(inspect_pdf(path)[:2], (True, "valid_pdf"))

    def test_html_login_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            session = Mock(); session.get.return_value = FakeResponse(b"<html>login</html>" * 1000, content_type="text/html")
            ok, validation, *_ = download_pdf(session, Candidate("https://x.org/login", "test"), Path(tmp) / "a.pdf", 1, 1, 1)
            self.assertFalse(ok); self.assertEqual(validation, "html_instead_of_pdf")

    def test_small_file_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "a.pdf"; path.write_bytes(b"%PDF-" + b"x" * 100)
            self.assertEqual(inspect_pdf(path)[1], "file_too_small")

    def test_partial_http_status_is_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            session = Mock(); session.get.return_value = FakeResponse(b"%PDF-" + b"x" * 11000, status=206)
            ok, *_ = download_pdf(session, Candidate("https://x.org/a.pdf", "test"), Path(tmp) / "a.pdf", 1, 1, 1)
            self.assertFalse(ok)

    def test_supplementary_redirect_is_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            session = Mock(); session.get.return_value = FakeResponse(b"%PDF-" + b"x" * 11000, url="https://x.org/article_s001.pdf")
            ok, validation, *_ = download_pdf(session, Candidate("https://x.org/a.pdf", "test"), Path(tmp) / "a.pdf", 1, 1, 1)
            self.assertFalse(ok); self.assertEqual(validation, "supplementary_detected")

    def test_filename_sanitization(self):
        record = base_record(); record["Title"] = 'Bad: title / with * chars?'
        name = allocate_filename(record, set())
        self.assertFalse(any(c in name for c in '<>:"/\\|?*')); self.assertTrue(name.startswith("P0001_Smith_2024"))

    def test_filename_collision_adds_hash(self):
        record = base_record(); first = allocate_filename(record, set()); used = {first.casefold()}
        second = allocate_filename(record, used)
        self.assertRegex(second, r"_[0-9a-f]{8}\.pdf$")

    def test_existing_pdf_metadata(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "existing.pdf"; path.write_bytes(b"%PDF-1.7\n" + b"x" * 11000)
            valid, status, size, digest = inspect_pdf(path)
            self.assertTrue(valid); self.assertEqual(status, "valid_pdf"); self.assertEqual(digest, hashlib.sha256(path.read_bytes()).hexdigest())

    def test_resume_claims_unique_valid_paper_id_pdf(self):
        from scripts.download_oa_writeback_excel import process_record
        with tempfile.TemporaryDirectory() as tmp:
            input_dir = Path(tmp); path = input_dir / "P0001_resume.pdf"
            path.write_bytes(b"%PDF-1.7\n" + b"x" * 11000)
            args = self.args(input_dir)
            result = process_record(base_record(), args, Mock(), input_dir.parent, {path.name.casefold()})
            self.assertEqual(result.status, "skipped_existing")
            self.assertEqual(result.file_name, path.name)

    def test_result_writes_back_to_papers(self):
        wb = Workbook(); ws = wb.active; ws.append(["Paper_ID", "DOI"]); headers = ensure_columns(ws, 1, ["Download_Status"] + [x for x in __import__('scripts.download_oa_writeback_excel', fromlist=['PAPERS_WRITE_FIELDS']).PAPERS_WRITE_FIELDS if x != "Download_Status"])
        ws.append(["P0001", "10.1/x"]); result = Result(status="downloaded", file_name="P0001.pdf", relative_path="input_raw/P0001.pdf", validation_status="valid_pdf")
        apply_result(ws, headers, 2, result, Path.cwd() / "book.xlsx")
        self.assertEqual(ws.cell(2, headers["Download_Status"]).value, "downloaded")

    def test_rag_manifest_upsert_by_paper_id(self):
        wb = Workbook(); ws = wb.active; ws.title = "PDF_RAG_MANIFEST"; ws.append(["Paper_ID"]); ws.append(["P0001"])
        records = [base_record()]; results = {"P0001": Result(status="not_oa")}
        update_rag_manifest(wb, records, results, Path(".")); update_rag_manifest(wb, records, results, Path("."))
        self.assertEqual(sum(ws.cell(r, 1).value == "P0001" for r in range(2, ws.max_row + 1)), 1)

    def test_rag_manifest_has_no_duplicate_rows(self):
        wb = Workbook(); ws = wb.active; ws.title = "PDF_RAG_MANIFEST"; ws.append(["Paper_ID"])
        records = [base_record("P0001"), {**base_record("P0002"), "DOI": "10.2/x"}]
        results = {x["Paper_ID"]: Result(status="not_oa") for x in records}
        self.assertEqual(update_rag_manifest(wb, records, results, Path(".")), 2)

    def test_download_analysis_statistics(self):
        wb = Workbook(); records = [base_record()]; results = {"P0001": Result(status="not_oa")}
        create_analysis(wb, Path("book.xlsx"), records, results, 0, Path("."))
        values = [wb["DOWNLOAD_ANALYSIS"].cell(r, 1).value for r in range(1, 30)]
        self.assertIn("Total_PAPERS", values)

    def test_followup_only_contains_anomalies(self):
        wb = Workbook(); records = [base_record("P0001"), base_record("P0002")]
        results = {"P0001": Result(status="not_oa"), "P0002": Result(status="no_pdf_url", needs_followup=True)}
        self.assertEqual(create_followup(wb, records, results), 1)
        self.assertEqual(wb["DOWNLOAD_FOLLOWUP"]["B2"].value, "P0002")

    def test_checkpoint_remains_openable(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "book.xlsx"; wb = Workbook(); wb.active["A1"] = "ok"; wb.save(path)
            wb.active["A2"] = "checkpoint"; atomic_save(wb, path); wb.close()
            check = load_workbook(path, read_only=True); self.assertEqual(check.active["A2"].value, "checkpoint"); check.close()

    def test_dry_run_does_not_modify_original_excel(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp); input_dir = root / "input_raw"; input_dir.mkdir(); path = input_dir / "eNRR_Master_Database.xlsx"; make_book(path)
            before = hashlib.sha256(path.read_bytes()).hexdigest(); args = self.args(input_dir); args.dry_run = True
            fake = Mock(); fake.get.side_effect = RuntimeError("network should be mocked at resolver")
            with patch("scripts.download_oa_writeback_excel.resolve_candidates") as resolver, patch("scripts.download_oa_writeback_excel.time.sleep"):
                from scripts.download_oa_writeback_excel import Resolution
                resolver.return_value = Resolution(False, "closed", unpaywall_is_oa=False)
                old = Path.cwd(); os.chdir(root)
                try: run(args, fake)
                finally: os.chdir(old)
            self.assertEqual(before, hashlib.sha256(path.read_bytes()).hexdigest())

    def test_original_sheet_and_format_preserved(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "book.xlsx"; make_book(path); wb = load_workbook(path)
            ws = wb["PAPERS"]; font = copy(ws["A2"].font); font.bold = True; ws["A2"].font = font; original = ws["A2"].style_id
            ensure_columns(ws, 2, ["OA_Verified"])
            self.assertIn("KEEP_ME", wb.sheetnames); self.assertEqual(ws["A2"].style_id, original)

    def test_status_file_inconsistency_can_recover(self):
        record = base_record(); record["Download_Status"] = "downloaded"; record["PDF_File_Name"] = "missing.pdf"
        from scripts.download_oa_writeback_excel import should_process
        self.assertTrue(should_process(record, self.args(Path(".")), Path(".")))

    def test_backup_created(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "book.xlsx"; path.write_bytes(b"original")
            backup = create_backup(path); self.assertTrue(backup.is_file()); self.assertEqual(backup.read_bytes(), b"original")

    def test_atomic_save_failure_leaves_original(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "book.xlsx"; path.write_bytes(b"original"); wb = Mock(); wb.save.side_effect = OSError("fail")
            with self.assertRaises(OSError): atomic_save(wb, path)
            self.assertEqual(path.read_bytes(), b"original")

    @staticmethod
    def args(input_dir):
        return argparse.Namespace(workbook=None, input_dir=Path(input_dir), start_id="P0001", end_id=None, max_records=1,
                                  overwrite=False, retry_failed=False, recheck_all=False, timeout=1.0, sleep=0.0,
                                  max_file_mb=1.0, checkpoint_every=1, unpaywall_email="real@example.org",
                                  openalex_api_key="", dry_run=False)


if __name__ == "__main__": unittest.main()
