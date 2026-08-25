from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import Workbook

from enh3bench.training_ingest import _empty_record, sha256_file
from enh3bench.training_rescue import (
    NOT_PRESENT_IN_RECIPE,
    NOT_REPORTED,
    canonical_header,
    classify_source_conflicts,
    component_reporting,
    extract_structured_experiments,
    inherit_same_paper_series,
    ownership_resolution,
    recalculate_v01,
    resolve_duplicates,
    review_packet_markdown,
    validate_no_cross_paper_inheritance,
)


def manifest(path: Path, bundle: str = "B1", paper: str = "P0001"):
    return {
        "bundle_id": bundle,
        "provisional_bundle_id": "NEW1",
        "source_group": "linrr_electrolyte_curated",
        "identity": {"title": "Test LiNRR paper", "doi": "10.1000/test"},
        "identity_match": {"existing_paper_id": paper},
        "assets": [{
            "source_asset_id": "A1", "relative_path": path.name, "absolute_path": str(path),
            "sha256": sha256_file(path), "format": path.suffix.lstrip("."), "probable_role": "supporting_information",
        }],
    }


class TrainingRescueV01Tests(unittest.TestCase):
    def test_docx_table_header_parsing_is_row_wise_and_preserves_header_unit(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "si.docx"
            doc = Document(); doc.add_paragraph("Table S1. Our experiments in 1 M LiBF4 in THF.")
            table = doc.add_table(rows=1, cols=4)
            for index, value in enumerate(["FE NH3 (%)", "j (mA cm-2)", "time (h)", "Electrolyte"]):
                table.cell(0, index).text = value
            cells = table.add_row().cells
            for index, value in enumerate(["42.1", "10", "1", "1 M LiBF4 in THF"]):
                cells[index].text = value
            doc.save(path)
            rows = extract_structured_experiments(manifest(path))
            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0]["fe_nh3_percent"], 42.1)
            self.assertEqual(rows[0]["current_density_mA_cm2"], 10.0)
            self.assertEqual(rows[0]["original_headers"]["fe_nh3_percent"], "FE NH3 (%)")
            self.assertEqual(rows[0]["original_units"]["duration_h"], "h")

    def test_xlsx_rows_are_distinct_experiments(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "source.xlsx"
            workbook = Workbook(); sheet = workbook.active; sheet.title = "Data"
            sheet.append(["Faradaic efficiency (%)", "current density (mA cm-2)", "LiBF4 concentration (M)", "Electrolyte"])
            sheet.append([40, 10, 1, "LiBF4 in THF"]); sheet.append([50, 20, 1, "LiBF4 in THF"])
            workbook.save(path)
            rows = extract_structured_experiments(manifest(path))
            self.assertEqual([row["fe_nh3_percent"] for row in rows], [40.0, 50.0])
            self.assertEqual([row["current_density_mA_cm2"] for row in rows], [10.0, 20.0])

    def test_common_header_aliases(self):
        self.assertEqual(canonical_header("Faradaic selectivity (%)"), "fe_nh3_percent")
        self.assertEqual(canonical_header("LiTFSI concentration (M)"), "lithium_salt_concentration_mol_L")
        self.assertEqual(canonical_header("H2O (ppm)"), "water_content_value")

    def test_same_paper_methods_inheritance_records_both_locators(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "main.txt"
            path.write_text("In this work, for the experiments of LMNRR, 1 M LiClO4 in DEE with 1.5 vol% n-pentanol was used.", encoding="utf-8")
            paper = manifest(path)
            row = _empty_record(); row.update({"record_id": "R", "paper_id": "P0001", "source_bundle_id": "B1", "solvent": "DEE", "condition_direct_locator": "Table S2:column:2", "supporting_source_locators": []})
            self.assertTrue(inherit_same_paper_series(row, paper))
            self.assertEqual(row["condition_value_origin"], "INHERITED_SAME_PAPER_SERIES")
            self.assertEqual(row["condition_inherited_from_asset"], "main.txt")
            self.assertEqual(row["lithium_salt"], "LiClO4")

    def test_cross_paper_inheritance_is_rejected(self):
        row = {"record_id": "R", "source_bundle_id": "B1", "condition_inherited_from_asset": "other.pdf"}
        with self.assertRaises(AssertionError):
            validate_no_cross_paper_inheritance([row], {"B1": {"assets": [{"relative_path": "own.pdf"}]}})

    def test_not_present_is_different_from_not_reported(self):
        row = {"electrolyte_notes": "LiBF4 in THF without additive", "lithium_salt": "LiBF4", "solvent": "THF"}
        status = component_reporting(row)
        self.assertEqual(status["additive"], NOT_PRESENT_IN_RECIPE)
        self.assertEqual(status["proton_donor"], NOT_REPORTED)

    def test_main_si_duplicate_merge_preserves_support(self):
        common = {"paper_id": "P1", "lithium_salt": "LiBF4", "solvent": "THF", "duration_h": 1, "fe_nh3_percent": 50, "supporting_source_locators": []}
        main = {**common, "record_id": "MAIN", "source_filename": "main.pdf", "source_locator": "page:2", "evidence_authority": "primary_text_candidate"}
        si = {**common, "record_id": "SI", "source_filename": "si.docx", "source_locator": "Table S1:row:2", "evidence_authority": "primary_structured_author_experiment", "supporting_source_locators": []}
        resolutions = resolve_duplicates([main, si])
        self.assertEqual(len(resolutions), 1)
        self.assertEqual(resolutions[0]["canonical_record_id"], "SI")
        self.assertIn("main.pdf#page:2", si["supporting_source_locators"])

    def test_distinct_conditions_are_preserved(self):
        first = {"record_id": "A", "paper_id": "P1", "lithium_salt": "LiBF4", "solvent": "THF", "duration_h": 1, "current_density_mA_cm2": 10, "fe_nh3_percent": 50}
        second = {**first, "record_id": "B", "current_density_mA_cm2": 20}
        self.assertEqual(resolve_duplicates([first, second]), [])

    def test_source_conflict_classification_does_not_average(self):
        first = {"record_id": "A", "paper_id": "P1", "lithium_salt": "LiBF4", "solvent": "THF", "duration_h": 1, "current_density_mA_cm2": 10, "fe_nh3_percent": 50, "ambiguity_flags": ["SOURCE_CONFLICT"]}
        second = {**first, "record_id": "B", "current_density_mA_cm2": 20, "fe_nh3_percent": 60, "ambiguity_flags": []}
        resolution = classify_source_conflicts([first, second])
        self.assertEqual(resolution[0]["classification"], "RESOLVED_DIFFERENT_CONDITIONS")
        self.assertEqual(first["fe_nh3_percent"], 50)

    def test_target_ownership_rules(self):
        self.assertEqual(ownership_resolution("We measured an FE of 55% in our experiments.")[0], "AUTHOR_OWNED")
        self.assertEqual(ownership_resolution("A previous study reported an FE of 55%.")[0], "UNCLEAR_OR_EXTERNAL")

    def test_review_packet_has_blank_human_fields(self):
        row = {"record_id": "R1", "paper_id": "P1", "missing_required_fields": ["electrolyte_defined"], "ambiguity_flags": [], "training_role": "target_label_candidate", "source_filename": "main.pdf", "source_locator": "page:2", "source_text_excerpt": "Exact excerpt"}
        packet = review_packet_markdown({"identity": {"title": "Paper", "doi": "10.1/x"}}, [row])
        self.assertIn("HUMAN_DECISION:\n", packet)
        self.assertIn("HUMAN_NOTES:\n", packet)
        self.assertIn("Exact excerpt", packet)

    def test_human_review_record_is_never_auto_promoted(self):
        row = _empty_record(); row.update({"record_id": "R", "paper_id": "P1", "evidence_owner_id": "P1", "source_bundle_id": "B1", "source_asset_id": "A", "source_file_sha256": "a" * 64, "source_locator": "Table S1:row:2", "reaction_family": "LiNRR", "document_genre": "primary_experimental", "training_role": "target_label_candidate", "lithium_salt": "LiBF4", "solvent": "THF", "duration_h": 1, "fe_nh3_percent": 50, "human_review_required": True})
        recalculate_v01(row)
        self.assertFalse(row["model_eligible_fe"])
        self.assertEqual(row["eligibility_tier"], "TIER_C")


if __name__ == "__main__":
    unittest.main()
